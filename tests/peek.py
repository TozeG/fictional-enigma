import sys
from openpyxl import load_workbook
f, sh = sys.argv[1], sys.argv[2]
rng = sys.argv[3] if len(sys.argv) > 3 else None
wv = load_workbook(f, data_only=True); wf = load_workbook(f)
ws, wsf = wv[sh], wf[sh]
cells = ws[rng] if rng else ws.iter_rows()
for row in cells:
    for c in row:
        if c.value is not None:
            fv = wsf[c.coordinate].value
            print(c.coordinate, repr(c.value)[:60], "|", str(fv)[:140] if isinstance(fv, str) and fv.startswith("=") else "")
