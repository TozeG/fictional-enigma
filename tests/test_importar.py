"""Teste de ponta a ponta do importador (execução em paralelo).

1. Gera o modelo de importação.
2. Preenche-o com o cenário de teste (src/data.py) + balancete de referência com UMA diferença deliberada.
3. Importa → matriz em modo produção; recalcula com LibreOffice; confirma que os números coincidem com a matriz de demonstração.
4. Confirma que a diferença deliberada é detectada no relatório.
5. Confirma que um lançamento desequilibrado bloqueia a importação.
"""
import json
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))
sys.path.insert(0, os.path.join(ROOT, "tests"))
os.environ["MATRIZ_MODO"] = "demo"
os.environ.setdefault("MATRIZ_LINHAS", "200")
import data as D  # noqa: E402
import sheets_base as SB  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from names import load  # noqa: E402

W = os.environ.get("MATRIZ_WORK", "/tmp/claude-0/w/imp")
os.makedirs(W, exist_ok=True)
RECALC = "/root/.claude/skills/synced/340ef685-a7ef-441c-a11a-f5069102cc9b_8be6ce2a-fc99-417e-ab54-d042138b7cfd/xlsx/scripts/recalc.py"
ok = []


def run(*args):
    env = dict(os.environ, MATRIZ_MODO="producao")
    return subprocess.run([sys.executable, os.path.join(ROOT, "src", "importar.py"), *args], capture_output=True, text=True, env=env)


def preencher(path, erro=False):
    wb = load_workbook(path)
    for nm in ("LANÇAMENTOS", "TERCEIROS", "PLANO_CONTAS", "BALANCETE_REFERÊNCIA"):
        ws = wb[nm]
        ws.delete_rows(2, ws.max_row)
    ent = wb["ENTIDADE"]
    for row in ent.iter_rows(min_row=2):
        cod = row[2].value
        row[1].value = {"CFG_Nome": "ENTIDADE PILOTO (teste)", "CFG_NIF": "5000000000", "CFG_Ano": D.ANO, "CFG_MesRep": 3, "CFG_SaldoMin": 1_500_000}.get(cod, row[1].value)
    lan = wb["LANÇAMENTOS"]
    for i, l in enumerate(D.JOURNAL):
        for j, (k, *_) in enumerate(SB.J_IN):
            lan.cell(row=2 + i, column=1 + j, value=l[k])
    if erro:
        lan.cell(row=3, column=12, value=99_999)  # desequilibra o lançamento 1
    t = wb["TERCEIROS"]
    for i, r in enumerate(D.TERCEIROS):
        for j, v in enumerate(r):
            t.cell(row=2 + i, column=1 + j, value=v)
    a = wb["ACTIVOS"]
    for i, r in enumerate(D.ACTIVOS):
        for j, v in enumerate(r):
            a.cell(row=2 + i, column=1 + j, value=v)
    # balancete de referência = saldos correctos, excepto o banco (extracto sem a despesa de 2 500)
    from collections import defaultdict
    s = defaultdict(float)
    for l in D.JOURNAL:
        s[l["Conta"]] += (l["Deb"] or 0) - (l["Cred"] or 0)
    b = wb["BALANCETE_REFERÊNCIA"]
    rows = [(c, "", v - (2_500 if c == "43.1.1" else 0)) for c, v in sorted(s.items()) if abs(v) > 0.005]
    rows += [("6", "Classe 6", sum(v for c, v in s.items() if c[0] == "6")), ("7", "Classe 7", sum(v for c, v in s.items() if c[0] == "7"))]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            b.cell(row=2 + i, column=1 + j, value=v)
    wb.save(path)


def check(desc, cond):
    ok.append((desc, cond))
    print(("✅ " if cond else "❌ ") + desc)


modelo = os.path.join(W, "modelo.xlsx")
r = run("modelo", modelo)
check("modelo de importação gerado", r.returncode == 0 and os.path.exists(modelo))

preencher(modelo)
saida = os.path.join(W, "MATRIZ_PILOTO_2026-03.xlsx")
r = run("carregar", modelo, "--saida", saida)
check("importação sem erros bloqueantes", r.returncode == 0)
rel = open(os.path.join(W, "MATRIZ_PILOTO_2026-03_RELATORIO.md"), encoding="utf-8").read()
check("relatório detecta a diferença deliberada no banco (43.1.1)", "| 43.1.1 |" in rel and "2,500.00 | 🔴" in rel)
check("relatório: apenas 1 conta com diferença", "Contas com diferença: 1 de" in rel)
rc = json.loads(subprocess.run([sys.executable, RECALC, saida, "600"], capture_output=True, text=True).stdout)
check(f"matriz importada recalcula sem erros ({rc.get('total_formulas')} fórmulas)", rc.get("total_errors") == 0)
wv, g = load(saida)
check("nome da entidade importado", g("CFG_Nome") == "ENTIDADE PILOTO (teste)")
check("resultado líquido = cenário de referência (1 206 460,50)", abs(g("DR_RL") - 1_206_460.5) < 0.01)
check("activo total = 13 799 614", abs(g("BS_AT") - 13_799_614) < 0.01)
check("balanço equilibrado", str(g("BS_Check")).startswith("🟢"))
check("IVA conciliado", str(g("IVA_Estado")).startswith("🟢"))
check("sistema não bloqueado", "BLOQUEADO" not in str(g("SYS_Estado")))

modelo2 = os.path.join(W, "modelo_erro.xlsx")
run("modelo", modelo2)
preencher(modelo2, erro=True)
saida2 = os.path.join(W, "MATRIZ_ERRO.xlsx")
if os.path.exists(saida2):
    os.remove(saida2)
r = run("carregar", modelo2, "--saida", saida2)
rel2 = open(os.path.join(W, "MATRIZ_ERRO_RELATORIO.md"), encoding="utf-8").read()
check("lançamento desequilibrado bloqueia a importação", r.returncode == 1 and not os.path.exists(saida2) and "não equilibrado" in rel2)

print(f"{sum(c for _, c in ok)}/{len(ok)} OK")
sys.exit(0 if all(c for _, c in ok) else 1)
