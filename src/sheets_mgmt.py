"""Governação: LEIA-ME, controlo interno, auditoria, fecho, encerramento, relatório de gestão, alertas e motor de consistência."""
from datetime import date
from openpyxl.utils import get_column_letter as CL
from core import *
from sheets_ledgers import MESES, listing
import sheets_base as SB
import data as D

S31, S32, S33, S37, S38, S39, S99 = ("31_CONTROLO_INTERNO", "32_AUDITORIA", "33_FECHO_MENSAL", "37_ENCERRAMENTO_EXERCÍCIO", "38_RELATÓRIO_GESTÃO",
                                     "39_ALERTAS", "99_CONTROLO_SISTEMA")
DIVIDA = "(BS_BPNC_EMP+BS_BPC_EMP)"
LIQ = "(BS_BA_DISP+BS_BA_APLIC)"

INDEX_SHEETS = [
    ("LEIA-ME", "—", "Guia, arquitectura, legenda de cores e índice."),
    ("00_CONFIGURAÇÃO", "INPUT", "Identificação, enquadramento, parâmetros financeiros e limites de alerta."),
    ("01_PLANO_CONTAS", "INPUT + PROC.", "PGC (Decreto n.º 82/01) com hierarquia, rubricas de balanço/DR e saldos automáticos."),
    ("01A_TABELAS", "INPUT", "Naturezas de operação (→ DFC), tipos de documento, rubricas, listas de validação."),
    ("01B_TERCEIROS", "INPUT + PROC.", "Clientes/fornecedores (NIF, prazos, limites, partes relacionadas) e saldos."),
    ("02A_OPERAÇÕES", "INPUT (modo aplicação)", "Registo simplificado: uma linha por operação (venda, compra, recebimento, pagamento, salários, impostos…). Gera os lançamentos automaticamente."),
    ("02B_MOTOR", "PROC.", "Motor que converte operações em partidas dobradas segundo os modelos de 01A_TABELAS."),
    ("02_DIÁRIO_LANÇAMENTOS", "INPUT (fonte única)", "Porta de entrada de todas as operações. Validações automáticas por linha."),
    ("03_DIÁRIO_GERAL", "OUTPUT", "Diário filtrável por período, conta, documento, terceiro, CC, projecto, natureza, utilizador."),
    ("04_RAZÃO", "OUTPUT", "Extracto de conta com saldo inicial, movimentos e saldo acumulado."),
    ("05_CAIXA", "OUTPUT", "Caixa mensal, contagem física, fundo fixo, alertas."),
    ("06_BANCOS", "OUTPUT", "Mapa de reconciliação bancária multi-conta."),
    ("07_CLIENTES", "OUTPUT", "Conta corrente por documento, aging, PMR, concentração, provisões."),
    ("08_FORNECEDORES", "OUTPUT", "Conta corrente por documento, aging, PMP, obrigações vencidas."),
    ("09_INVENTÁRIOS", "INPUT + OUTPUT", "Ficha de artigos, reposição, rotação, cobertura e kardex a custo médio."),
    ("10_ACTIVOS_FIXOS", "INPUT + PROC.", "Registo de imobilizado, depreciações mensais e reconciliação."),
    ("11_FISCALIDADE_AGT", "INPUT (legal)", "Parametrização fiscal rastreável (diploma, artigo, vigência, fonte, estado)."),
    ("12_IVA", "PROC. + OUTPUT", "Apuramento mensal do IVA e conciliação contabilidade × mapa fiscal."),
    ("13_FACTURAÇÃO_FISCAL", "OUTPUT", "Registo de documentos emitidos, controlos DP 71/25 e checklist AGT."),
    ("14_SAFT", "OUTPUT", "Mapeamento e validador SAF-T (AO)."),
    ("15_IMPOSTO_INDUSTRIAL", "PROC.", "Reconciliação contabilístico-fiscal e estimativa do imposto."),
    ("16_ORÇAMENTO_EMPRESARIAL", "INPUT", "Orçamentos parciais → resultados, caixa e patrimonial."),
    ("17_PLANEAMENTO_FINANCEIRO", "INPUT + OUTPUT", "Projecção a 5 anos, 4 cenários e simulador de impacto."),
    ("18_TESOURARIA", "OUTPUT", "Rolling de 13 semanas, necessidades de tesouraria, realizado × previsto."),
    ("19_FLUXO_DE_CAIXA", "OUTPUT", "Fluxo de caixa mensal por actividade (método directo)."),
    ("20_BALANCETE", "OUTPUT", "Balancete mensal/trimestral/semestral/anual/acumulado."),
    ("21_BALANÇO", "OUTPUT", "Balanço com teste Activo = CP + Passivo."),
    ("22_DRE", "OUTPUT", "Demonstração de resultados por natureza (mês, trimestre, semestre, acumulado)."),
    ("23_DFC", "OUTPUT", "Demonstração dos fluxos de caixa (método directo)."),
    ("24_GESTÃO_DE_PROJECTOS", "INPUT + OUTPUT", "Projectado × realizado, VAL, TIR, payback por projecto."),
    ("25_ANÁLISE_DE_INVESTIMENTOS", "INPUT + OUTPUT", "VAL, TIR, payback, IR, ROI, ROIC, WACC e sensibilidade."),
    ("26_RISCO_FINANCEIRO", "INPUT + OUTPUT", "Matriz de risco 5×5 com indicadores automáticos."),
    ("27_SUSTENTABILIDADE", "OUTPUT", "Equilíbrio financeiro, CCL, NFM, autofinanciamento."),
    ("28_KPI_FINANCEIROS", "OUTPUT", "KPI com limites parametrizáveis e semáforos."),
    ("29_DASHBOARD_EXECUTIVO", "OUTPUT", "Painel executivo com semáforos e gráficos."),
    ("30_BUDGET_VS_ACTUAL", "OUTPUT", "Orçamento × real com causa, responsável e acção."),
    ("31_CONTROLO_INTERNO", "OUTPUT", "Matriz de controlos e trilho de auditoria."),
    ("32_AUDITORIA", "OUTPUT", "Testes automáticos de auditoria."),
    ("33_FECHO_MENSAL", "INPUT + OUTPUT", "Checklist de fecho e encerramento de períodos."),
    ("34_CALENDÁRIO_FISCAL_AGT", "INPUT + OUTPUT", "Obrigações, prazos (parametrizados), estado e alertas."),
    ("35_BASE_LEGAL", "INPUT (referência)", "Matriz legal e checklist de conformidade."),
    ("36_BREAK_EVEN", "OUTPUT", "Ponto de equilíbrio, margem de segurança, alavancagem operacional."),
    ("37_ENCERRAMENTO_EXERCÍCIO", "OUTPUT", "Proposta de apuramento de resultados e abertura do exercício seguinte."),
    ("38_RELATÓRIO_GESTÃO", "OUTPUT", "Relatório mensal/trimestral/semestral/anual e relatório à Administração."),
    ("39_ALERTAS", "OUTPUT", "Sistema de alertas parametrizado."),
    ("99_CONTROLO_SISTEMA", "OUTPUT", "Motor de consistência e provas de reconciliação global."),
]


def build_leiame(wb):
    ws = wb.create_sheet("LEIA-ME")
    title(ws, "MATRIZ PRO MASTER — CONTABILIDADE, FISCALIDADE, GESTÃO FINANCEIRA E PLANEAMENTO (ANGOLA)",
          "Sistema integrado em Excel: UMA ÚNICA FONTE DE DADOS (02_DIÁRIO_LANÇAMENTOS) → múltiplas visões contabilísticas, fiscais, financeiras, orçamentais e de gestão.")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 100
    put(ws, "A4", "Estado do sistema:", "label", bold=True)
    put(ws, "B4", "=SYS_Estado", "grey", bold=True)
    status_cf(ws, "B4")
    put(ws, "A5", "Aviso legal:", "label", bold=True)
    put(ws, "C5", "Estruturado para conformidade com o quadro legal identificado e sujeito à validação contabilística, fiscal e técnica antes da utilização oficial. "
                  "Esta folha não é software de facturação certificado pela AGT e não emite documentos fiscais.", "note", wrap=True)
    section(ws, 7, 1, "LEGENDA DE CORES", 3)
    leg = [("Azul (texto) / fundo azul-claro", C_INPUT_FILL, C_INPUT_FONT, "Campo de entrada (input). Único local onde se escreve."),
           ("Preto", "FFFFFF", "000000", "Fórmula (não editar)."), ("Cinzento", C_CALC_FILL, "000000", "Processamento / informação protegida."),
           ("Verde", C_OK[0], C_OK[1], "Informação validada / conforme."), ("Amarelo", C_WARN[0], C_WARN[1], "Atenção / pendente / por validar."),
           ("Vermelho", C_ERR[0], C_ERR[1], "Erro / crítico / bloqueante.")]
    for i, (lab, bg, fg, desc) in enumerate(leg):
        c = ws.cell(row=8 + i, column=1, value=lab)
        c.fill = fill(bg)
        c.font = font(True, fg)
        put(ws, (8 + i, 3), desc, "label")
    section(ws, 15, 1, "COMO UTILIZAR (fluxo de trabalho)", 3)
    passos = ["1. Configure a entidade em 00_CONFIGURAÇÃO (ano, mês de reporte, regime, limites).",
              "2. Reveja/adapte o plano de contas (01_PLANO_CONTAS) — mantenha as colunas; só contas 'Movimento' recebem lançamentos.",
              "3. Registe terceiros em 01B_TERCEIROS (o Diário só pede o NIF).",
              "4. Registe cada operação UMA vez em 02A_OPERAÇÕES (venda, compra, recebimento, pagamento, salários, impostos…): os lançamentos, IVA, custos, INSS e IRT são gerados automaticamente. Situações especiais (abertura, notas de crédito, moeda estrangeira, acréscimos/diferimentos): 02_DIÁRIO_LANÇAMENTOS, zona manual.",
              "5. Corrija todas as linhas com '🔴 ERRO' (coluna Erros_Detectados).",
              "6. Consulte 99_CONTROLO_SISTEMA: o sistema deve estar 'SISTEMA CONFORME' antes do reporte.",
              "7. Feche o mês em 33_FECHO_MENSAL (Estado = S + data de fecho). Lançamentos posteriores nesse mês ficam assinalados.",
              "8. Relatórios: 22_DRE, 21_BALANÇO, 23_DFC, 29_DASHBOARD_EXECUTIVO, 38_RELATÓRIO_GESTÃO.",
              "9. Protecção: folhas protegidas SEM palavra-passe (Rever → Desproteger). Defina palavra-passe na implementação real.",
              "10. Para >2 000 linhas: aumente o intervalo (gerador src/build.py, variável MATRIZ_LINHAS) ou migre a fonte para Power Query / Power Pivot (ver docs/)."]
    for i, p in enumerate(passos):
        put(ws, (16 + i, 1), p, "label")
    section(ws, 27, 1, "ARQUITECTURA — 3 CAMADAS", 3)
    arq = [("CAMADA 1 — INPUT", "Configuração, plano, terceiros, Diário, orçamento, pressupostos, parâmetros fiscais."),
           ("CAMADA 2 — PROCESSAMENTO", "Colunas cinzentas do Diário, saldos do plano, IVA, Imposto Industrial, depreciações."),
           ("CAMADA 3 — OUTPUT", "Razão, balancete, demonstrações, KPI, dashboard, relatórios, alertas, controlo.")]
    for i, (a, b) in enumerate(arq):
        put(ws, (28 + i, 1), a, "label", bold=True)
        put(ws, (28 + i, 3), b, "label")
    section(ws, 32, 1, "ÍNDICE DE FOLHAS", 3)
    header(ws, 33, 1, ["Folha", "Camada", "Conteúdo"], None)
    for i, (s, l, d) in enumerate(INDEX_SHEETS):
        c = ws.cell(row=34 + i, column=1, value=s)
        c.hyperlink = f"#'{s}'!A1"
        c.font = font(False, "0563C1")
        put(ws, (34 + i, 2), l, "label")
        put(ws, (34 + i, 3), d, "label")
    protect(ws)


def build_controlo_interno(wb):
    ws = wb.create_sheet(S31)
    title(ws, "31 — CONTROLO INTERNO E TRILHO DE AUDITORIA", "Quem → fez o quê → quando → em que documento → qual o impacto contabilístico.", "CAMADA 3 — OUTPUT (responsáveis = input)")
    header(ws, 5, 1, ["Princípio", "Controlo", "Indicador automático", "Valor", "Estado", "Responsável", "Frequência"], [20, 50, 42, 10, 26, 18, 12])
    ctr = [
        ("Segregação de funções", "Quem lança não valida o mesmo lançamento", "Linhas com Validado_Por = Utilizador", '=COUNTIF(J_Segreg,"Violação")', 0),
        ("Autorização", "Todos os lançamentos validados por responsável", "Linhas por aprovar", '=COUNTIF(J_Segreg,"Pendente")', 1),
        ("Documentação", "Todo o lançamento tem documento de suporte", "Linhas sem suporte", '=COUNTIFS(J_ID,"<>",J_Suporte,"")', 0),
        ("Reconciliação", "Bancos, caixa, clientes e fornecedores reconciliados", "Reconciliações com diferença ou pendentes", '=(ABS(CX_Dif)>CFG_Tol)+(BK_Pend>0)+(ABS(BK_Dif)>CFG_Tol)+(ABS(CLI_Dif)>CFG_Tol)+(ABS(FRN_Dif)>CFG_Tol)', 1),
        ("Rastreabilidade", "Utilizador e data de inserção registados", "Linhas sem utilizador ou data de inserção", '=COUNTIFS(J_ID,"<>",J_User,"")+COUNTIFS(J_ID,"<>",J_DtIns,"")', 0),
        ("Revisão", "Revisão mensal do balancete e das demonstrações", "Meses até ao reporte não encerrados", '=CFG_MesRep-COUNTIF(FECHO_Estado,"S")', 1),
        ("Aprovação", "Orçamento aprovado e desvios justificados", "Rubricas com desvio 🔴 sem causa", '=SUMPRODUCT((LEFT(\'30_BUDGET_VS_ACTUAL\'!G6:G18,2)="🔴")*(\'30_BUDGET_VS_ACTUAL\'!H6:H18=""))', 1),
        ("Encerramento de períodos", "Sem lançamentos em períodos encerrados", "Alterações após o fecho", "=SUM(J_ErrPer)", 0),
        ("Controlo de alterações", "Alterações registadas (Data_Alteração)", "Linhas alteradas", '=COUNT(J_DtAlt)', 1),
    ]
    for i, (p, c, ind, f, crit) in enumerate(ctr):
        r = 6 + i
        put(ws, (r, 1), p, "label", bold=True)
        put(ws, (r, 2), c, "label")
        put(ws, (r, 3), ind, "label")
        put(ws, (r, 4), f, "calc", "0")
        put(ws, (r, 5), f'=IF(D{r}=0,"🟢 OK",IF({crit}=0,"🔴 Falha de controlo","🟡 Pendente"))', "calc")
        put(ws, (r, 6), "Director Financeiro", "input")
        put(ws, (r, 7), "Mensal", "input")
    status_cf(ws, "E6:E14")
    put(ws, "A16", "Estado do controlo interno", "label", bold=True)
    put(ws, "E16", '=IF(COUNTIF(E6:E14,"🔴*")>0,"🔴 Falhas de controlo",IF(COUNTIF(E6:E14,"🟡*")>0,"🟡 Controlo com pendências","🟢 Controlo interno OK"))', "grey", bold=True)
    status_cf(ws, "E16")
    name(wb, "CI_Estado", S31, "$E$16")
    section(ws, 18, 1, "TRILHO DE AUDITORIA (espelho do Diário, linha a linha)", 9)
    header(ws, 19, 1, ["Quem (utilizador)", "Fez o quê", "Quando", "ID", "Documento", "Conta", "Impacto (D − C)", "Validado por", "Estado"], None)
    JC = SB.JCOLS
    q_ = q(SB.S_J)
    for i in range(SB.JR1 - SB.JR0 + 1):
        r = 20 + i
        jr = SB.JR0 + i
        a = f"{q_}!{JC['ID']}{jr}"
        put(ws, (r, 1), f'=IF({a}="","",{q_}!{JC["User"]}{jr}&"")', "calc")
        put(ws, (r, 2), f'=IF({a}="","",IF(N({q_}!{JC["DtAlt"]}{jr})>0,"Alterou","Inseriu")&" — "&{q_}!{JC["Nat"]}{jr})', "calc")
        put(ws, (r, 3), f'=IF({a}="","",IF(N({q_}!{JC["DtAlt"]}{jr})>0,{q_}!{JC["DtAlt"]}{jr},{q_}!{JC["DtIns"]}{jr}))', "calc", DATE)
        put(ws, (r, 4), f'=IF({a}="","",{a})', "calc", "0")
        put(ws, (r, 5), f'=IF({a}="","",{q_}!{JC["TipoDoc"]}{jr}&" "&{q_}!{JC["Serie"]}{jr}&"/"&{q_}!{JC["NumDoc"]}{jr})', "calc")
        put(ws, (r, 6), f'=IF({a}="","",{q_}!{JC["Conta"]}{jr}&"")', "calc")
        put(ws, (r, 7), f'=IF({a}="","",{q_}!{JC["DC"]}{jr})', "calc", NUM)
        put(ws, (r, 8), f'=IF({a}="","",{q_}!{JC["Validador"]}{jr}&"")', "calc")
        put(ws, (r, 9), f'=IF({a}="","",{q_}!{JC["EstVal"]}{jr})', "calc")
    status_cf(ws, f"I20:I{20 + SB.JR1 - SB.JR0}")
    ws.freeze_panes = "A20"
    protect(ws)


def build_auditoria(wb):
    ws = wb.create_sheet(S32)
    title(ws, "32 — AUDITORIA: TESTES AUTOMÁTICOS", "Testes substantivos e de conformidade sobre a totalidade dos lançamentos (não amostragem).", "CAMADA 3 — OUTPUT")
    put(ws, "A4", "Parâmetros do teste de valores anormais", "label", bold=True)
    put(ws, "A5", "Média do montante por linha (excl. abertura)", "label")
    put(ws, "C5", '=IFERROR(AVERAGEIFS(J_Montante,J_IsAbe,0,J_Ano,CFG_Ano),0)', "grey", NUM)
    name(wb, "AUD_Media", S32, "$C$5")
    put(ws, "A6", "Desvio-padrão", "label")
    put(ws, "C6", '=IF(COUNTIFS(J_IsAbe,0,J_Ano,CFG_Ano)<2,0,SQRT(SUM(J_Mont2)/(COUNTIFS(J_IsAbe,0,J_Ano,CFG_Ano)-1)))', "grey", NUM)
    put(ws, "A7", "Limite (média + k × desvio-padrão)", "label")
    put(ws, "C7", "=C5+CFG_kAnom*C6", "grey", NUM)
    name(wb, "AUD_Limite", S32, "$C$7")
    ws.column_dimensions["A"].width = 44
    header(ws, 9, 1, ["#", "Teste", "Ocorrências", "Valor (Kz)", "Estado", "Procedimento recomendado"], [4, 44, 11, 15, 30, 60])
    ws.column_dimensions["A"].width = 6
    tests = [
        ("Débito ≠ crédito (linhas de lançamentos desequilibrados)", '=COUNTIF(J_Erros,"*NÃO EQUILIBRADO*")', '=ABS(SUM(J_Deb)-SUM(J_Cred))', "R", "Corrigir o lançamento no Diário até Σ D = Σ C."),
        ("Documentos duplicados", "=SUM(J_Dup)", '=SUMIFS(J_Montante,J_Dup,1)', "R", "Confirmar se se trata de dupla contabilização; estornar."),
        ("Lançamentos sem documento de suporte", '=COUNTIF(J_Erros,"*Sem documento de suporte*")', '=SUMIFS(J_Montante,J_Suporte,"",J_ID,"<>")', "R", "Obter e arquivar o suporte; caso contrário estornar."),
        ("Lançamentos fora do exercício", '=COUNTIF(J_Erros,"*Fora do exercício*")', "=0", "R", "Mover para o ficheiro do exercício correcto."),
        ("Valores anormais (> média + k·σ)", "=SUM(J_Anomalo)", '=SUMIFS(J_Montante,J_Anomalo,1)', "Y", "Inspeccionar documentos e autorizações dos lançamentos assinalados."),
        ("Contas incompatíveis (inexistentes ou de agregação)", '=COUNTIF(J_Erros,"*Conta inexistente*")+COUNTIF(J_Erros,"*não movimentável*")', "=0", "R", "Corrigir a conta para uma conta de movimento válida."),
        ("Alterações posteriores ao encerramento", "=SUM(J_ErrPer)", '=SUMIFS(J_Montante,J_ErrPer,1)', "R", "Reabrir o período com autorização formal ou estornar no período aberto."),
        ("Reconciliações pendentes (caixa, bancos, clientes, fornecedores, inventário, activos)",
         '=(ABS(CX_Dif)>CFG_Tol)+BK_Pend+(ABS(BK_Dif)>CFG_Tol)+(ABS(CLI_Dif)>CFG_Tol)+(ABS(FRN_Dif)>CFG_Tol)+(ABS(INV_Dif)>CFG_Tol)+COUNTIF(AF_Estado,"🔴*")', "=ABS(BK_Dif)+ABS(CX_Dif)", "Y", "Concluir reconciliações e contabilizar itens pendentes."),
        ("Saldos negativos anormais (caixa, bancos, clientes credores, fornecedores devedores, stock)",
         "=(CX_Saldo<0)+COUNTIF('06_BANCOS'!E6:E20,\"<0\")+COUNTIF(T_SaldoCli,\"<0\")+COUNTIF(T_SaldoForn,\"<0\")+COUNTIF(INV_Qtd,\"<0\")", "=0", "Y", "Analisar natureza: adiantamentos, erros de imputação ou descobertos."),
        ("Transacções com partes relacionadas", "=SUM(J_ParteRel)", "=SUMIFS(J_Montante,J_ParteRel,1)", "Y", "Verificar condições de mercado e necessidade de divulgação."),
        ("Operações fiscais inconsistentes", '=SUM(J_ErrTaxa)+COUNTIF(J_Erros,"*IVA sem código*")+COUNTIF(FAT_Val,"🔴*")', "=0", "R", "Corrigir código fiscal/base/taxa; regularizar documentos."),
        ("Violações de segregação de funções", '=COUNTIF(J_Segreg,"Violação")', "=0", "R", "Revalidar por utilizador distinto."),
    ]
    for i, (t, f, v, sev, proc) in enumerate(tests):
        r = 10 + i
        put(ws, (r, 1), i + 1, "label")
        put(ws, (r, 2), t, "label", wrap=True)
        put(ws, (r, 3), f, "calc", "0")
        put(ws, (r, 4), v, "calc", NUM)
        put(ws, (r, 5), f'=IF(C{r}=0,"🟢 Sem excepções",' + ('"🔴 "&C{r}&" excepção(ões)")' if sev == "R" else '"🟡 "&C{r}&" a analisar")').format(r=r), "calc")
        put(ws, (r, 6), proc, "label", wrap=True)
    r1 = 9 + len(tests)
    status_cf(ws, f"E10:E{r1}")
    put(ws, (r1 + 2, 2), "Resultado global da auditoria", "label", bold=True)
    put(ws, (r1 + 2, 5), f'=IF(COUNTIF(E10:E{r1},"🔴*")>0,"🔴 Excepções críticas",IF(COUNTIF(E10:E{r1},"🟡*")>0,"🟡 Excepções a analisar","🟢 Sem excepções"))', "grey", bold=True)
    status_cf(ws, f"E{r1 + 2}")
    name(wb, "AUD_Estado", S32, f"$E${r1 + 2}")
    protect(ws)


def build_fecho(wb):
    ws = wb.create_sheet(S33)
    title(ws, "33 — FECHO CONTABILÍSTICO MENSAL", "Controlos automáticos por mês + confirmação manual (✔). Período só pode ser encerrado sem lançamentos inválidos.",
          "CAMADA 1 — INPUT (✔, estado, data de fecho)  |  CAMADA 3 — controlos")
    header(ws, 5, 1, ["Tarefa"] + MESES, [40] + [14] * 12)
    ws.freeze_panes = "B6"
    auto = {
        "Caixa reconciliado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(INDEX(CX_Est,{m}),2)="🟢","✔","☐"))',
        "Bancos reconciliados": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF({m}<CFG_MesRep,IF(INDEX(FECHO_Estado,{m})="S","✔","☐"),IF(ABS(BK_Dif)<=CFG_Tol,"✔","☐")))',
        "Clientes reconciliados": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF({m}<CFG_MesRep,IF(INDEX(FECHO_Estado,{m})="S","✔","☐"),IF(ABS(CLI_Dif)<=CFG_Tol,"✔","☐")))',
        "Fornecedores reconciliados": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF({m}<CFG_MesRep,IF(INDEX(FECHO_Estado,{m})="S","✔","☐"),IF(ABS(FRN_Dif)<=CFG_Tol,"✔","☐")))',
        "Inventário reconciliado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF({m}<CFG_MesRep,IF(INDEX(FECHO_Estado,{m})="S","✔","☐"),IF(ABS(INV_Dif)<=CFG_Tol,"✔","☐")))',
        "Activos fixos actualizados": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(AF_Estado,2)="🟢","✔","☐"))',
        "Depreciações processadas": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(INDEX(AF_DepEst,{m}),2)="🟢","✔","☐"))',
        "IVA apurado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(INDEX(IVA_EstRow,{m}),2)="🔴","☐","✔"))',
        "Balancete validado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(AND(ABS(SUMIFS(J_Deb,J_Ano,CFG_Ano,J_Mes,{m})-SUMIFS(J_Cred,J_Ano,CFG_Ano,J_Mes,{m}))<0.005,COUNTIFS(J_Ano,CFG_Ano,J_Mes,{m},J_ErrFlag,1)=0),"✔","☐"))',
        "DRE validada": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(DR_Estado,2)="🟢","✔","☐"))',
        "Balanço validado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(BS_Check,2)="🟢","✔","☐"))',
        "Fluxo de caixa validado": lambda m, c: f'=IF({m}>CFG_MesRep,"",IF(LEFT(INDEX(FC_EstRow,{m}),2)="🟢","✔","☐"))',
    }
    manual = ["Acréscimos processados", "Diferimentos processados", "Impostos verificados"]
    order = ["Caixa reconciliado", "Bancos reconciliados", "Clientes reconciliados", "Fornecedores reconciliados", "Inventário reconciliado",
             "Activos fixos actualizados", "Depreciações processadas", "Acréscimos processados", "Diferimentos processados", "IVA apurado",
             "Impostos verificados", "Balancete validado", "DRE validada", "Balanço validado", "Fluxo de caixa validado"]
    r = 6
    for t in order:
        put(ws, (r, 1), ("☐ " if t in manual else "⚙ ") + t, "label")
        for m in range(1, 13):
            if t in auto:
                put(ws, (r, 1 + m), auto[t](m, CL(1 + m)), "calc", align="center")
            else:
                put(ws, (r, 1 + m), "✔" if m <= 2 and not D.PRODUCAO else None, "input", align="center")
        r += 1
    last_task = r - 1
    put(ws, (r, 1), "Lançamentos inválidos no mês", "label", bold=True)
    for m in range(1, 13):
        put(ws, (r, 1 + m), f"=COUNTIFS(J_Ano,CFG_Ano,J_Mes,{m},J_ErrFlag,1)", "calc", "0", align="center")
    rerr = r
    r += 1
    put(ws, (r, 1), "Tarefas concluídas", "label")
    for m in range(1, 13):
        c = CL(1 + m)
        put(ws, (r, 1 + m), f'=COUNTIF({c}6:{c}{last_task},"✔")&"/{len(order)}"', "calc", align="center")
    r += 2
    put(ws, (r, 1), "PERÍODO ENCERRADO? (S/N) — input", "label", bold=True)
    for m in range(1, 13):
        put(ws, (r, 1 + m), "S" if m <= 2 and not D.PRODUCAO else "N", "input", align="center")
    rest = r
    name(wb, "FECHO_Estado", S33, f"$B${r}:$M${r}")
    dv_list(ws, f"B{r}:M{r}", '"S,N"')
    r += 1
    put(ws, (r, 1), "Data de fecho — input", "label", bold=True)
    for m in range(1, 13):
        put(ws, (r, 1 + m), date(D.ANO, m + 1, 10) if m <= 2 and not D.PRODUCAO else None, "input", DATE)
    name(wb, "FECHO_Data", S33, f"$B${r}:$M${r}")
    r += 1
    put(ws, (r, 1), "ESTADO DO PERÍODO", "label", bold=True)
    for m in range(1, 13):
        c = CL(1 + m)
        put(ws, (r, 1 + m), f'=IF({c}{rest}="S",IF({c}{rerr}>0,"🔴 BLOQUEADO — "&{c}{rerr}&" erro(s)",IF(COUNTIF({c}6:{c}{last_task},"☐")>0,"🟡 Encerrado com tarefas por concluir","🟢 ENCERRADO")),IF({m}<=CFG_MesRep,"🟡 Aberto","—"))', "calc", align="center")
    status_cf(ws, f"B{r}:M{r}")
    name(wb, "FECHO_Periodo", S33, f"$B${r}:$M${r}")
    value_cf(ws, f"B6:M{last_task}", 'B6="☐"', C_WARN)
    value_cf(ws, f"B6:M{last_task}", 'B6="✔"', C_OK)
    put(ws, (r + 2, 1), "Regra: um período com 'S' e lançamentos inválidos fica BLOQUEADO. Qualquer lançamento inserido/alterado depois da data de fecho num mês encerrado é assinalado no Diário ('Alteração em período encerrado').", "note")
    protect(ws)


def build_encerramento(wb):
    ws = wb.create_sheet(S37)
    title(ws, "37 — ENCERRAMENTO DO EXERCÍCIO E ABERTURA DO SEGUINTE", "Proposta automática de lançamentos. Copiar (valores) para o Diário: apuramento com natureza 'Apuramento de resultados'; abertura no ficheiro do novo ano com natureza 'Abertura'. Nunca apagar históricos.",
          "CAMADA 3 — OUTPUT")
    put(ws, "A4", "Válido para o exercício completo quando o mês de reporte = 12. Mês de reporte actual:", "label")
    put(ws, "H4", "=CFG_MesRep", "grey", "0")
    section(ws, 6, 1, "A. APURAMENTO DO RESULTADO — ENCERRAMENTO DAS CONTAS DE RESULTADOS (→ 88)", 8)
    put(ws, "A7", "=MAX(PC_SeqEnc)", "grey", "0")
    header(ws, 8, 1, ["#", "Linha plano", "Conta", "Designação", "Saldo (D−C)", "Débito", "Crédito", "Natureza"], [5, 6, 10, 44, 15, 15, 15, 24])
    n = 120
    for i in range(n):
        r = 9 + i
        put(ws, (r, 1), f'=IF({i + 1}<=$A$7,{i + 1},"")', "grey", "0")
        put(ws, (r, 2), f'=IF(A{r}="","",MATCH(A{r},PC_SeqEnc,0))', "grey", "0")
        put(ws, (r, 3), f'=IF(B{r}="","",INDEX(PC_Cod,B{r}))', "calc")
        put(ws, (r, 4), f'=IF(B{r}="","",INDEX(PC_Desig,B{r}))', "calc")
        put(ws, (r, 5), f'=IF(B{r}="","",INDEX(PC_Saldo,B{r}))', "calc", NUM)
        put(ws, (r, 6), f'=IF(B{r}="","",MAX(0,-E{r}))', "calc", NUM)
        put(ws, (r, 7), f'=IF(B{r}="","",MAX(0,E{r}))', "calc", NUM)
        put(ws, (r, 8), f'=IF(B{r}="","","Apuramento de resultados")', "calc")
    r = 9 + n
    put(ws, (r, 3), "88", "label", bold=True)
    put(ws, (r, 4), "Resultado líquido do exercício (contrapartida)", "label", bold=True)
    put(ws, (r, 6), f"=MAX(0,SUM(G9:G{r - 1})-SUM(F9:F{r - 1}))", "grey", NUM, bold=True)
    put(ws, (r, 7), f"=MAX(0,SUM(F9:F{r - 1})-SUM(G9:G{r - 1}))", "grey", NUM, bold=True)
    r88 = r
    r += 1
    put(ws, (r, 4), "TOTAIS / CONTROLO", "label", bold=True)
    put(ws, (r, 6), f"=SUM(F9:F{r88})", "grey", NUM, bold=True)
    put(ws, (r, 7), f"=SUM(G9:G{r88})", "grey", NUM, bold=True)
    put(ws, (r, 8), f'=IF(ABS(F{r}-G{r})<0.005,"🟢 Apuramento equilibrado","🔴 Desequilibrado")', "grey")
    put(ws, (r + 1, 4), "Resultado apurado (crédito 88 = lucro) vs DR", "label")
    put(ws, (r + 1, 5), f"=G{r88}-F{r88}", "grey", NUM)
    put(ws, (r + 1, 8), f'=IF(ABS(E{r + 1}-DR_RL)<0.005,"🟢 = Resultado líquido da DR","🟡 Verificar: já existem lançamentos de apuramento")', "grey")
    status_cf(ws, f"H{r}:H{r + 1}")
    name(wb, "ENC_Estado", S37, f"$H${r}")
    name(wb, "ENC_RL", S37, f"$E${r + 1}")
    r += 4
    section(ws, r, 1, "B. ABERTURA DO EXERCÍCIO SEGUINTE — SALDOS INICIAIS (após apuramento)", 8)
    r += 1
    put(ws, (r, 1), "=MAX(PC_SeqAbe)", "grey", "0")
    cnt = f"$A${r}"
    header(ws, r + 1, 1, ["#", "Linha plano", "Conta", "Designação", "Saldo (D−C)", "Débito", "Crédito", "Natureza"], None)
    b0 = r + 2
    for i in range(250):
        rr = b0 + i
        put(ws, (rr, 1), f'=IF({i + 1}<={cnt},{i + 1},"")', "grey", "0")
        put(ws, (rr, 2), f'=IF(A{rr}="","",MATCH(A{rr},PC_SeqAbe,0))', "grey", "0")
        put(ws, (rr, 3), f'=IF(B{rr}="","",INDEX(PC_Cod,B{rr}))', "calc")
        put(ws, (rr, 4), f'=IF(B{rr}="","",INDEX(PC_Desig,B{rr}))', "calc")
        put(ws, (rr, 5), f'=IF(B{rr}="","",INDEX(PC_Saldo,B{rr}))', "calc", NUM)
        put(ws, (rr, 6), f'=IF(B{rr}="","",MAX(0,E{rr}))', "calc", NUM)
        put(ws, (rr, 7), f'=IF(B{rr}="","",MAX(0,-E{rr}))', "calc", NUM)
        put(ws, (rr, 8), f'=IF(B{rr}="","","Abertura")', "calc")
    rr = b0 + 250
    put(ws, (rr, 3), "81.1", "label", bold=True)
    put(ws, (rr, 4), "Resultado líquido do exercício → Resultados transitados (aplicação a deliberar em AG)", "label", bold=True)
    put(ws, (rr, 5), f"=-ENC_RL", "grey", NUM)
    put(ws, (rr, 6), f"=MAX(0,E{rr})", "grey", NUM)
    put(ws, (rr, 7), f"=MAX(0,-E{rr})", "grey", NUM)
    put(ws, (rr, 8), "Abertura", "label")
    put(ws, (rr + 1, 4), "TOTAIS / CONTROLO DE DIFERENÇAS", "label", bold=True)
    put(ws, (rr + 1, 6), f"=SUM(F{b0}:F{rr})", "grey", NUM, bold=True)
    put(ws, (rr + 1, 7), f"=SUM(G{b0}:G{rr})", "grey", NUM, bold=True)
    put(ws, (rr + 1, 8), f'=IF(ABS(F{rr + 1}-G{rr + 1})<0.005,"🟢 Abertura equilibrada","🔴 Diferença na abertura")', "grey")
    status_cf(ws, f"H{rr + 1}")
    name(wb, "ENC_AbeEstado", S37, f"$H${rr + 1}")
    protect(ws)


def build_relatorio(wb):
    ws = wb.create_sheet(S38)
    title(ws, "38 — RELATÓRIO DE GESTÃO E RELATÓRIO À ADMINISTRAÇÃO", "Gerado automaticamente. Recomendações derivam de regras parametrizadas sobre os dados — não de opiniões genéricas.", "CAMADA 3 — OUTPUT")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 70
    put(ws, "A4", "Tipo de relatório", "label", bold=True)
    put(ws, "B4", "Trimestral", "input")
    dv_list(ws, "B4", '"Mensal,Trimestral,Semestral,Anual"')
    put(ws, "A5", "Nº do período (mês / trimestre / semestre)", "label")
    put(ws, "B5", 1, "input", "0")
    put(ws, "C4", "Mês inicial", "label")
    put(ws, "D4", '=IF(B4="Mensal",B5,IF(B4="Trimestral",(B5-1)*3+1,IF(B4="Semestral",(B5-1)*6+1,1)))', "grey", "0")
    put(ws, "C5", "Mês final", "label")
    put(ws, "D5", '=MIN(CFG_MesRep,IF(B4="Mensal",B5,IF(B4="Trimestral",B5*3,IF(B4="Semestral",B5*6,12))))', "grey", "0")
    name(wb, "RP_Ini", S38, "$D$4")
    name(wb, "RP_Fim", S38, "$D$5")
    put(ws, "A6", '=CFG_Nome&" — Relatório "&LOWER(B4)&" — período: meses "&RP_Ini&" a "&RP_Fim&" de "&CFG_Ano', "label", bold=True)
    per = 'J_Ano,CFG_Ano,J_Mes,">="&RP_Ini,J_Mes,"<="&RP_Fim,J_IncDR,1'
    perc = 'J_Ano,CFG_Ano,J_Mes,">="&RP_Ini,J_Mes,"<="&RP_Fim'

    def dr(codes):
        return "=" + "+".join(f'-SUMIFS(J_DC,J_RubDR,"{c}",{per})' for c in codes)
    orc = lambda codes: "=" + "+".join(f'SUMPRODUCT((ORC_Codes="{c}")*(COLUMN(ORC_M)-COLUMN(INDEX(ORC_M,1,1))+1>=RP_Ini)*(COLUMN(ORC_M)-COLUMN(INDEX(ORC_M,1,1))+1<=RP_Fim)*ORC_M)' for c in codes)
    allc = [c for c, *_ in D.RUBRICAS_DR]
    header(ws, 8, 1, ["Secção / indicador", "Período", "Orçamento / referência", "Leitura"], None)
    blocks = [
        ("1. RESUMO EXECUTIVO", None),
        ("Receita (vendas + serviços)", (dr(["DR01", "DR02"]), orc(["DR01", "DR02"]), NUM)),
        ("Resultado líquido", (dr(allc), orc(allc), NUM)),
        ("Estado do sistema", ("=SYS_Estado", None, None)),
        ("2. EVOLUÇÃO DAS RECEITAS", None),
        ("Proveitos operacionais", (dr(["DR01", "DR02", "DR03"]), orc(["DR01", "DR02", "DR03"]), NUM)),
        ("3. EVOLUÇÃO DOS GASTOS", None),
        ("Custos operacionais", (dr(["DR04", "DR05", "DR06", "DR07", "DR08"]), orc(["DR04", "DR05", "DR06", "DR07", "DR08"]), NUM)),
        ("  dos quais: pessoal", (dr(["DR06"]), orc(["DR06"]), NUM)),
        ("  dos quais: FSE", (dr(["DR05"]), orc(["DR05"]), NUM)),
        ("4. RESULTADO", None),
        ("EBITDA", (dr(["DR01", "DR02", "DR03", "DR04", "DR05", "DR06", "DR08"]), None, NUM)),
        ("Resultado operacional", (dr(["DR01", "DR02", "DR03", "DR04", "DR05", "DR06", "DR07", "DR08"]), orc(["DR01", "DR02", "DR03", "DR04", "DR05", "DR06", "DR07", "DR08"]), NUM)),
        ("Resultado antes de impostos", (dr(allc[:-1]), orc(allc[:-1]), NUM)),
        ("5. CASH FLOW", None),
        ("Fluxo operacional", (f'=SUMIFS(J_DC,J_Fluxo,"Operacional",{perc})', None, NUM)),
        ("Fluxo de investimento", (f'=SUMIFS(J_DC,J_Fluxo,"Investimento",{perc})', None, NUM)),
        ("Fluxo de financiamento", (f'=SUMIFS(J_DC,J_Fluxo,"Financiamento",{perc})', None, NUM)),
        ("6. POSIÇÃO DE CAIXA", None),
        ("Meios monetários no fim do período", ("=INDEX(FC_SF,RP_Fim)", None, NUM)),
        ("Saldo mínimo projectado (13 semanas)", ("=TES_Min", "=CFG_SaldoMin", NUM)),
        ("7. CLIENTES (à data de reporte)", None),
        ("Saldo de clientes", ("=CLI_Total", None, NUM)), ("Saldo vencido", ("=CLI_Venc", None, NUM)), ("PMR (dias)", ("=CLI_PM", "=CFG_PMRalvo", "0")),
        ("8. FORNECEDORES", None),
        ("Saldo de fornecedores", ("=FRN_Total", None, NUM)), ("Obrigações vencidas", ("=FRN_Venc", None, NUM)), ("PMP (dias)", ("=FRN_PM", "=CFG_PMPalvo", "0")),
        ("9. DÍVIDA", None),
        ("Dívida financeira", (f"={DIVIDA}", None, NUM)), ("Dívida / EBITDA", ("=K_DEB", "=CFG_DivEBITDA", MULT)),
        ("10–12. ORÇAMENTO, EXECUÇÃO E DESVIOS", None),
        ("Rubricas com desvio desfavorável (acum.)", ("=BVA_Red", None, "0")), ("Execução dos custos (acum.)", ("=BVA_ExecCustos", None, PCT)),
        ("13. INDICADORES", None),
        ("Margem líquida (acum.)", ("=K_ML", "=CFG_MargMin", PCT)), ("Liquidez geral", ("=K_LG", None, MULT)), ("Autonomia financeira", ("=K_AF", None, PCT)),
        ("ROE (anualizado)", ("=K_ROE", None, PCT)), ("Indicadores críticos", ("=K_Red", None, "0")),
        ("14. RISCOS", None),
        ("Risco com maior exposição", ("=RSK_Top", None, None)), ("Riscos críticos", ("=RSK_Red", None, "0")),
        ("15. INVESTIMENTOS", None),
        ("Investimento em activos fixos no período", (f'=SUMIFS(J_Deb,J_Classe,"1",J_IsAbe,0,{perc})', None, NUM)),
        ("Projecto mais intensivo em capital", ("=PRJ_MaisCapital", None, None)),
        ("SITUAÇÃO FISCAL", None),
        ("Total a entregar ao Estado (saldos credores)", ("=FIS_Total", None, NUM)), ("IVA — estado", ("=IVA_Estado", None, None)),
        ("Facturação — estado", ("=FAT_Estado", None, None)), ("Obrigações fiscais vencidas", ("=CAL_Vencidas", None, "0")),
        ("SUSTENTABILIDADE", None), ("Diagnóstico", ("=SUS_Estado", None, None)), ("Ponto de equilíbrio", ("=BE_Estado", None, None)),
    ]
    r = 9
    for lab, v in blocks:
        if v is None:
            section(ws, r, 1, lab, 4)
            r += 1
            continue
        f, ref, fmt = v
        put(ws, (r, 1), lab, "label")
        put(ws, (r, 2), f, "calc", fmt)
        if ref:
            put(ws, (r, 3), ref, "calc", fmt)
            if fmt == NUM:
                put(ws, (r, 4), f'=IF(C{r}=0,"",IF(B{r}>=C{r},"🟢 Acima/igual à referência ("&TEXT(IF(C{r}=0,0,B{r}/C{r}-1),"0.0%")&")","🟡 Abaixo da referência ("&TEXT(IF(C{r}=0,0,B{r}/C{r}-1),"0.0%")&")"))', "calc")
        r += 1
    status_cf(ws, f"B9:D{r}")
    # 16. recomendações
    r += 1
    section(ws, r, 1, "16. RECOMENDAÇÕES DE GESTÃO (regras parametrizadas)", 4)
    r += 1
    header(ws, r, 1, ["Regra", "Accionada?", "Valor", "Recomendação"], None)
    rules = [
        ("Liquidez projectada abaixo do mínimo", "TES_Min<CFG_SaldoMin", "TES_Min", '"Reforçar liquidez: necessidade de "&TEXT(TES_Necess,"#,##0")&" Kz na "&TES_Semana&"; negociar linha de crédito ou antecipar cobranças."'),
        ("Clientes com saldo vencido > 90 dias", "CLI_V90>0", "CLI_V90", '"Accionar cobrança/contencioso dos saldos > 90 dias e rever limites de crédito."'),
        ("Saldo de clientes vencido", "CLI_Venc>0", "CLI_Venc", '"Cobrar "&TEXT(CLI_Venc,"#,##0")&" Kz vencidos; priorizar maiores saldos (07_CLIENTES)."'),
        ("Margem líquida abaixo do mínimo", "K_ML<CFG_MargMin", "K_ML", '"Rever preços e estrutura de custos: margem líquida abaixo do mínimo definido."'),
        ("Rubricas de custo acima do orçamento", "BVA_Exced>0", "BVA_Exced", '"Justificar desvios e definir acções correctivas em 30_BUDGET_VS_ACTUAL."'),
        ("IVA a pagar no mês", "IVA_PagarMes>0", "IVA_PagarMes", '"Reservar "&TEXT(IVA_PagarMes,"#,##0")&" Kz para o IVA até "&TEXT(EOMONTH(CFG_DataRef,1),"dd/mm/yyyy")&" (prazo parametrizado)."'),
        ("Obrigações fiscais vencidas", "CAL_Vencidas>0", "CAL_Vencidas", '"Regularizar de imediato as obrigações vencidas para limitar multas e juros."'),
        ("Documentos fiscais com erro/pendentes", 'LEFT(FAT_Estado,2)<>"🟢"', 'COUNTIF(FAT_Val,"🔴*")+COUNTIF(FAT_Val,"🟡*")', '"Corrigir/reenviar documentos à AGT (13_FACTURAÇÃO_FISCAL)."'),
        ("Endividamento acima do limite", "K_DEB>CFG_DivEBITDA", "K_DEB", '"Reduzir dívida ou aumentar EBITDA: Dívida/EBITDA acima do covenant parametrizado."'),
        ("Excedente de tesouraria", "TES_Excedente>0", "TES_Excedente", '"Avaliar aplicação de "&TEXT(TES_Excedente,"#,##0")&" Kz (prazo compatível com a previsão de 13 semanas)."'),
        ("Concentração de clientes", "CLI_Conc1>CFG_ConcCli", "CLI_Conc1", '"Diversificar a carteira: maior cliente representa "&TEXT(CLI_Conc1,"0%")&" do saldo."'),
        ("Taxas fiscais POR VALIDAR em uso", 'COUNTIF(TX_Alerta,"🟡 Em uso*")>0', 'COUNTIF(TX_Alerta,"🟡 Em uso*")', '"Validar a parametrização fiscal com técnico de contas / jurista antes do uso oficial."'),
        ("Desequilíbrio financeiro (CCL < NFM)", "SUS_TL<0", "SUS_TL", '"Financiar as NFM com capitais permanentes; renegociar prazos com fornecedores."'),
        ("Margem de segurança reduzida", 'AND(ISNUMBER(BE_MS),BE_MS<0.2)', "BE_MS", '"Aumentar margem de contribuição ou reduzir custos fixos."'),
        ("Sistema com pendências", 'LEFT(SYS_Estado,2)<>"🟢"', "SYS_Estado", '"Resolver pendências de 99_CONTROLO_SISTEMA antes do reporte oficial."'),
    ]
    r0 = r + 1
    for i, (lab, cond, val, rec) in enumerate(rules):
        rr = r0 + i
        put(ws, (rr, 1), lab, "label")
        put(ws, (rr, 2), f'=IF({cond},"✔ Accionada","—")', "calc")
        put(ws, (rr, 3), f"={val}", "calc", NUM)
        put(ws, (rr, 4), f'=IF({cond},{rec},"")', "calc", wrap=True)
    value_cf(ws, f"B{r0}:B{r0 + len(rules) - 1}", f'LEFT(B{r0},1)="✔"', C_WARN)
    name(wb, "REC_Count", S38, f"$B${r0 + len(rules) + 1}")
    put(ws, (r0 + len(rules) + 1, 1), "Nº de recomendações accionadas", "label", bold=True)
    put(ws, (r0 + len(rules) + 1, 2), f'=COUNTIF(B{r0}:B{r0 + len(rules) - 1},"✔*")', "grey", "0")
    protect(ws)


def build_alertas(wb):
    ws = wb.create_sheet(S39)
    title(ws, "39 — SISTEMA DE ALERTAS", "Condições e limites parametrizados em 00_CONFIGURAÇÃO / 28_KPI_FINANCEIROS.", "CAMADA 3 — OUTPUT")
    header(ws, 5, 1, ["#", "Alerta", "Valor", "Limite", "Estado", "Onde agir"], [4, 42, 16, 16, 30, 26])
    q06 = "'06_BANCOS'"
    al = [
        ("Caixa negativo", "=CX_Saldo", "0", "=IF(CX_Saldo<0,\"🔴 Crítico\",\"🟢 Normal\")", "05_CAIXA", NUM),
        ("Saldo bancário negativo", f"=COUNTIF({q06}!E6:E20,\"<0\")", "0", "=IF(C7>0,\"🔴 Crítico\",\"🟢 Normal\")", "06_BANCOS", "0"),
        ("Clientes com saldo vencido", "=CLI_Venc", "0", "=IF(CLI_V90>0,\"🔴 Crítico\",IF(CLI_Venc>0,\"🟡 Atenção\",\"🟢 Normal\"))", "07_CLIENTES", NUM),
        ("Fornecedores com saldo vencido", "=FRN_Venc", "0", "=IF(FRN_V90>0,\"🔴 Crítico\",IF(FRN_Venc>0,\"🟡 Atenção\",\"🟢 Normal\"))", "08_FORNECEDORES", NUM),
        ("Excesso de dívida (Dívida/EBITDA)", "=K_DEB", "=CFG_DivEBITDA", "=IF(C10>D10,\"🔴 Crítico\",\"🟢 Normal\")", "28_KPI_FINANCEIROS", MULT),
        ("Liquidez projectada abaixo do limite", "=TES_Min", "=CFG_SaldoMin", "=IF(C11<0,\"🔴 Crítico\",IF(C11<D11,\"🟡 Atenção\",\"🟢 Normal\"))", "18_TESOURARIA", NUM),
        ("Orçamento excedido (rubricas)", "=BVA_Exced", "0", "=IF(C12>0,\"🔴 Crítico\",IF(BVA_Red>0,\"🟡 Atenção\",\"🟢 Normal\"))", "30_BUDGET_VS_ACTUAL", "0"),
        ("Margem líquida abaixo do mínimo", "=K_ML", "=CFG_MargMin", "=IF(C13<0,\"🔴 Crítico\",IF(C13<D13,\"🟡 Atenção\",\"🟢 Normal\"))", "22_DRE", PCT),
        ("Queda anormal da receita (mês vs média anterior)", "=IF(CFG_MesRep<2,0,IFERROR(INDEX(DRM_Rec,CFG_MesRep)/AVERAGE(INDEX(DRM_Rec,1):INDEX(DRM_Rec,CFG_MesRep-1))-1,0))", "=CFG_QuedaRec", "=IF(C14<D14,\"🔴 Crítico\",\"🟢 Normal\")", "22_DRE", PCT),
        ("Aumento anormal de custos (mês vs média anterior)", "=IF(CFG_MesRep<2,0,IFERROR(INDEX(DRM_Cust,CFG_MesRep)/AVERAGE(INDEX(DRM_Cust,1):INDEX(DRM_Cust,CFG_MesRep-1))-1,0))", "=CFG_AumCust", "=IF(C15>D15,\"🟡 Atenção\",\"🟢 Normal\")", "22_DRE", PCT),
        ("Imposto Industrial por pagar", "=II_APagar", "0", "=IF(C16>0,\"🟡 Atenção\",\"🟢 Normal\")", "15_IMPOSTO_INDUSTRIAL", NUM),
        ("IVA por pagar (mês)", "=IVA_PagarMes", "0", "=IF(C17>0,\"🟡 Atenção\",\"🟢 Normal\")", "12_IVA", NUM),
        ("Documentos fiscais inconsistentes", '=COUNTIF(FAT_Val,"🔴*")', "0", "=IF(C18>0,\"🔴 Crítico\",IF(LEFT(FAT_Estado,2)=\"🟡\",\"🟡 Atenção\",\"🟢 Normal\"))", "13_FACTURAÇÃO_FISCAL", "0"),
        ("Lançamentos desequilibrados", '=COUNTIF(J_Erros,"*NÃO EQUILIBRADO*")', "0", "=IF(C19>0,\"🔴 Crítico\",\"🟢 Normal\")", "02_DIÁRIO_LANÇAMENTOS", "0"),
        ("Contas sem movimento esperado (reconciliáveis, com saldo)", '=COUNTIFS(PC_Rec,"S",PC_Tipo,"Movimento",PC_UltMov,"<"&(CFG_DataRef-CFG_DiasSemMov),PC_Saldo,"<>0")', "=CFG_DiasSemMov", "=IF(C20>0,\"🟡 Atenção\",\"🟢 Normal\")", "01_PLANO_CONTAS", "0"),
        ("Concentração excessiva de clientes", "=CLI_Conc1", "=CFG_ConcCli", "=IF(C21>D21,\"🟡 Atenção\",\"🟢 Normal\")", "07_CLIENTES", PCT),
        ("Concentração excessiva de fornecedores", "=FRN_Conc1", "=CFG_ConcForn", "=IF(C22>D22,\"🟡 Atenção\",\"🟢 Normal\")", "08_FORNECEDORES", PCT),
        ("Risco cambial elevado (passivo em ME / passivo)", '=IF(BS_PT=0,0,-SUMIFS(J_DC,J_Classe,"3",J_Moeda,"<>AOA",J_Moeda,"?*",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)/BS_PT)', "=CFG_ExpCamb", "=IF(C23>D23,\"🔴 Crítico\",\"🟢 Normal\")", "26_RISCO_FINANCEIRO", PCT),
        ("Obrigações fiscais com prazo ultrapassado", "=CAL_Vencidas", "0", "=IF(C24>0,\"🔴 Crítico\",IF(CAL_AVencer>0,\"🟡 Atenção\",\"🟢 Normal\"))", "34_CALENDÁRIO_FISCAL_AGT", "0"),
    ]
    for i, (lab, v, lim, st, where, fmt) in enumerate(al):
        r = 6 + i
        put(ws, (r, 1), i + 1, "label")
        put(ws, (r, 2), lab, "label")
        put(ws, (r, 3), v, "calc", fmt)
        put(ws, (r, 4), lim if str(lim).startswith("=") else float(lim), "calc", fmt)
        put(ws, (r, 5), st, "calc")
        c = ws.cell(row=r, column=6, value=where)
        c.hyperlink = f"#'{where}'!A1"
        c.font = font(False, "0563C1")
    r1 = 5 + len(al)
    status_cf(ws, f"E6:E{r1}")
    put(ws, (r1 + 2, 2), "Alertas activos (🔴 + 🟡)", "label", bold=True)
    put(ws, (r1 + 2, 3), f'=COUNTIF(E6:E{r1},"🔴*")+COUNTIF(E6:E{r1},"🟡*")', "grey", "0")
    put(ws, (r1 + 3, 2), "Alertas críticos (🔴)", "label")
    put(ws, (r1 + 3, 3), f'=COUNTIF(E6:E{r1},"🔴*")', "grey", "0")
    name(wb, "ALR_Count", S39, f"$C${r1 + 2}")
    name(wb, "ALR_Red", S39, f"$C${r1 + 3}")
    protect(ws)


def build_controlo(wb):
    ws = wb.create_sheet(S99)
    title(ws, "99 — MOTOR DE CONSISTÊNCIA DO SISTEMA", "Verificações de integridade e provas matemáticas de reconciliação global.", "CAMADA 3 — OUTPUT")
    header(ws, 5, 1, ["#", "Verificação", "Estado", "Crítica p/ bloqueio"], [4, 48, 52, 12])
    chk = [
        ("Diário equilibrado?", '=IF(AND(ABS(SUM(J_Deb)-SUM(J_Cred))<0.005,COUNTIF(J_Erros,"*NÃO EQUILIBRADO*")=0),"🟢 Diário equilibrado","🔴 ERRO — Diário desequilibrado")', "S"),
        ("Lançamentos sem erros de validação?", '=IF(SUM(J_ErrFlag)=0,"🟢 Sem erros","🔴 "&SUM(J_ErrFlag)&" linha(s) com erro")', "S"),
        ("Razão reconciliado com o Diário?", "=H{p1}", "S"),
        ("Balancete equilibrado?", '=IF(AND(LEFT(BT_Estado,2)="🟢",LEFT(BT_DiarioEst,2)="🟢"),"🟢 Balancete equilibrado e = Diário",BT_Estado&" | "&BT_DiarioEst)', "S"),
        ("Balanço equilibrado?", "=BS_Check", "S"),
        ("Plano de contas íntegro?", '=IF(COUNTIF(PC_Val,"🔴*")=0,"🟢 Plano íntegro","🔴 "&COUNTIF(PC_Val,"🔴*")&" conta(s) com erro")', "S"),
        ("Demonstração de resultados conciliada?", "=DR_Estado", "N"),
        ("Caixa reconciliado?", "=CX_Estado", "N"),
        ("Bancos reconciliados?", "=BK_Estado", "N"),
        ("IVA conciliado?", "=IVA_Estado", "N"),
        ("Fiscalidade conciliada (Imposto Industrial e retenções)?", '=IF(AND(LEFT(II_Estado,2)="🟢",LEFT(RET_Estado,2)="🟢"),"🟢 Fiscalidade conciliada",II_Estado&" | "&RET_Estado)', "N"),
        ("Clientes conciliados?", "=CLI_Estado", "N"),
        ("Fornecedores conciliados?", "=FRN_Estado", "N"),
        ("Inventário conciliado?", "=INV_Estado", "N"),
        ("Activos conciliados?", "=AF_Estado", "N"),
        ("Orçamento conciliado (mapa por rubrica = orçamentos parciais)?", "=H{p9}", "N"),
        ("Fluxo de caixa conciliado?", '=IF(AND(LEFT(FC_Estado,2)="🟢",LEFT(DFC_Estado,2)="🟢"),"🟢 Fluxo de caixa e DFC conciliados",FC_Estado&" | "&DFC_Estado)', "N"),
        ("Facturação conforme?", "=FAT_Estado", "N"),
        ("SAF-T (preparação)?", "=SAFT_Estado", "N"),
        ("Controlo interno?", "=CI_Estado", "N"),
        ("Auditoria?", "=AUD_Estado", "N"),
        ("Períodos encerrados sem bloqueio?", '=IF(COUNTIF(FECHO_Periodo,"🔴*")>0,"🔴 Período encerrado com erros","🟢 OK")', "S"),
    ]
    pr = 6 + len(chk) + 4
    for i, (lab, f, crit) in enumerate(chk):
        r = 6 + i
        put(ws, (r, 1), i + 1, "label")
        put(ws, (r, 2), lab, "label")
        put(ws, (r, 3), f.replace("{p1}", str(pr + 2)).replace("{p9}", str(pr + 11)), "calc")
        put(ws, (r, 4), crit, "label", align="center")
    r1 = 5 + len(chk)
    status_cf(ws, f"C6:C{r1}")
    put(ws, (r1 + 2, 2), "ESTADO GLOBAL DO SISTEMA", "label", bold=True)
    put(ws, (r1 + 2, 3), f'=IF(COUNTIFS(C6:C{r1},"🔴*",D6:D{r1},"S")>0,"🔴 SISTEMA BLOQUEADO",IF(COUNTIF(C6:C{r1},"🔴*")+COUNTIF(C6:C{r1},"🟡*")>0,"🟡 SISTEMA COM PENDÊNCIAS","🟢 SISTEMA CONFORME"))', "grey", bold=True)
    ws.cell(row=r1 + 2, column=3).font = font(True, "000000", 13)
    status_cf(ws, f"C{r1 + 2}")
    name(wb, "SYS_Estado", S99, f"$C${r1 + 2}")
    # provas de reconciliação global
    section(ws, pr, 1, "TESTE DE RECONCILIAÇÃO GLOBAL — PROVAS MATEMÁTICAS", 8)
    header(ws, pr + 1, 1, ["#", "Prova", "Valor A", "Valor B", "Diferença", "", "", "Estado"], [4, 48, 18, 18, 14, 2, 2, 30])
    yt = 'J_Ano,CFG_Ano,J_Mes,"<="&CFG_MesRep'
    proofs = [
        ("Diário = Razão (Σ débitos do Diário vs Σ débitos das contas de movimento)", f"=SUMIFS(J_Deb,{yt})", '=SUMIFS(PC_Deb,PC_Tipo,"Movimento")'),
        ("Razão = Balancete (Σ saldos devedores − credores = 0)", '=SUMIFS(PC_Dev,PC_Tipo,"Movimento")', '=SUMIFS(PC_Cr,PC_Tipo,"Movimento")'),
        ("Balancete = Demonstrações (Activo vs CP + Passivo)", "=BS_AT", "=BS_CP+BS_PT"),
        ("Resultado no Balanço = Resultado na DR", "=BS_BCP_RL", "=DR_RL"),
        ("Caixa contabilístico = movimento de caixa (05_CAIXA)", '=SUMIFS(J_DC,J_G2,"45",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "=CX_Saldo"),
        ("Bancos contabilísticos = reconciliação bancária (extracto ajustado)", "=BK_Saldo", "='06_BANCOS'!M22"),
        ("IVA contabilístico (saldo 34.5) = mapa fiscal de IVA (apuramento acumulado)", "=IVA_Saldo345", "=IVA_ApurAcum"),
        ("Resultado contabilístico → fiscal → imposto (estimado vs contabilizado)", "=II_Estimado", '=SUMIFS(J_DC,J_G2,"87",' + yt + ')'),
        ("Fluxos contabilísticos → DFC (saldo final DFC vs contas de meios monetários)", "=DFC_Fim", '=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)'),
        ("Orçamento: Σ mapa por rubrica = Σ orçamentos parciais (RL orçado)", "=SUM(ORC_M)", "=SUM(ORC_RL)"),
    ]
    for i, (lab, a, b) in enumerate(proofs):
        r = pr + 2 + i
        put(ws, (r, 1), i + 1, "label")
        put(ws, (r, 2), lab, "label", wrap=True)
        put(ws, (r, 3), a, "calc", NUM)
        put(ws, (r, 4), b, "calc", NUM)
        put(ws, (r, 5), f"=C{r}-D{r}", "calc", NUM)
        put(ws, (r, 8), f'=IF(ABS(E{r})<=CFG_Tol,"🟢 Provado","🔴 Diferença")', "calc")
    status_cf(ws, f"H{pr + 2}:H{pr + 1 + len(proofs)}")
    protect(ws)
