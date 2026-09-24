"""Camada INPUT + motor: configuração, tabelas, plano de contas, terceiros, diário de lançamentos."""
from datetime import date
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table, TableStyleInfo
from core import *
import data as D

S_CFG = "00_CONFIGURAÇÃO"
S_PC = "01_PLANO_CONTAS"
S_TAB = "01A_TABELAS"
S_TER = "01B_TERCEIROS"
S_J = "02_DIÁRIO_LANÇAMENTOS"
S_33 = "33_FECHO_MENSAL"

JR0 = 6
JR1 = JR0 + int(__import__("os").environ.get("MATRIZ_LINHAS", "2000")) - 1  # linhas do diário
PR0, PR1 = 6, 405           # linhas do plano
TR0, TR1 = 6, 305           # linhas de terceiros

J_IN = [
    ("ID", "ID_Lançamento", 8, "0"), ("Data", "Data", 11, DATE), ("TipoDoc", "Tipo_Documento", 8, "@"),
    ("Serie", "Série", 7, "@"), ("NumDoc", "Nº_Documento", 9, "0"), ("DataDoc", "Data_Documento", 11, DATE),
    ("Venc", "Data_Vencimento", 11, DATE), ("NIF", "NIF_Terceiro", 12, "@"), ("Conta", "Conta", 9, "@"),
    ("Desc", "Descrição", 34, None), ("Deb", "Débito", 14, NUM), ("Cred", "Crédito", 14, NUM),
    ("CC", "Centro_Custo", 7, "@"), ("CR", "Centro_Responsabilidade", 8, "@"), ("Proj", "Projecto", 7, "@"),
    ("Fonte", "Fonte_Financiamento", 10, "@"), ("Nat", "Natureza_Operação", 24, "@"), ("CodF", "Código_Fiscal", 10, "@"),
    ("Base", "Base_Tributável", 13, NUM), ("Ref", "Ref_Origem_ID", 8, "0"), ("Moeda", "Moeda", 6, "@"),
    ("Cambio", "Taxa_Câmbio", 9, "#,##0.0000"), ("ValME", "Valor_Moeda_Origem", 12, NUM), ("FormaPag", "Forma_Pagamento", 12, "@"),
    ("Artigo", "Artigo", 7, "@"), ("Qtd", "Quantidade", 9, "#,##0.00"), ("Suporte", "Documento_Suporte", 20, "@"),
    ("User", "Utilizador", 9, "@"), ("Validador", "Validado_Por", 9, "@"), ("DtIns", "Data_Inserção", 11, DATE),
    ("DtAlt", "Data_Alteração", 11, DATE), ("EstDoc", "Estado_Documento", 10, "@"), ("Hash", "Hash_AGT", 7, "@"),
    ("Cert", "Nº_Certificado_Software", 16, "@"), ("EstAGT", "Estado_Comunicação_AGT", 12, "@"), ("DtCom", "Data_Comunicação", 11, DATE),
    ("ErroCom", "Erro_Comunicação", 18, "@"), ("Obs", "Observação", 40, None),
]

# Colunas calculadas (camada de PROCESSAMENTO) — [X] = coluna X na mesma linha
J_CALC = [
    ("Ano", "Exercício", 7, "0", '=IF([ID]="","",IF(ISNUMBER([Data]),YEAR([Data]),""))'),
    ("Mes", "Mês", 5, "0", '=IF([ID]="","",IF(ISNUMBER([Data]),MONTH([Data]),""))'),
    ("Trim", "Trimestre", 5, "0", '=IF([Mes]="","",ROUNDUP([Mes]/3,0))'),
    ("Sem", "Semestre", 5, "0", '=IF([Mes]="","",IF([Mes]<=6,1,2))'),
    ("Classe", "Classe", 5, "@", '=IF([ID]="","",LEFT([Conta],1))'),
    ("G2", "Conta_2D", 5, "@", '=IF([ID]="","",LEFT([Conta],2))'),
    ("NomeConta", "Designação_Conta", 28, None, '=IF([ID]="","",IFERROR(INDEX(PC_Desig,MATCH([Conta],PC_Cod,0)),"#CONTA INEXISTENTE"))'),
    ("ContaOK", "Conta_Existe", 5, "0", '=IF([ID]="","",IF(ISNUMBER(MATCH([Conta],PC_Cod,0)),1,0))'),
    ("ContaMov", "Conta_Movimento", 5, "0", '=IF([ID]="","",IF([ContaOK]=1,IF(INDEX(PC_Tipo,MATCH([Conta],PC_Cod,0))="Movimento",1,0),0))'),
    ("ContaCorr", "Conta_Corrente", 5, "@", '=IF([ID]="","",IF([ContaOK]=1,INDEX(PC_CCorr,MATCH([Conta],PC_Cod,0))&"","N"))'),
    ("Terceiro", "Nome_Terceiro", 22, None, '=IF(OR([ID]="",[NIF]=""),"",IFERROR(INDEX(T_Nome,MATCH([NIF],T_NIF,0)),"#NIF NÃO REGISTADO"))'),
    ("TipoTerc", "Tipo_Terceiro", 10, None, '=IF(OR([ID]="",[NIF]=""),"",IFERROR(INDEX(T_Tipo,MATCH([NIF],T_NIF,0)),""))'),
    ("NIFfmt", "NIF_Formato_OK", 5, "0", '=IF(OR([ID]="",[NIF]=""),"",IF(IFERROR(INDEX(T_Pais,MATCH([NIF],T_NIF,0)),"Angola")<>"Angola",1,IF(OR([NIF]=CFG_NIFCF,AND(LEN([NIF])=10,ISNUMBER(-[NIF])),AND(LEN([NIF])=14,ISNUMBER(-LEFT([NIF],9)),ISNUMBER(-RIGHT([NIF],3)))),1,0)))'),
    ("DC", "Valor_D_menos_C", 13, NUM, '=N([Deb])-N([Cred])'),
    ("Montante", "Montante_Linha", 13, NUM, '=N([Deb])+N([Cred])'),
    ("SaldoID", "Saldo_Lançamento", 11, NUM, '=IF([ID]="","",IF([ID]=[ID-1],[SaldoID-1],SUMIF(J_ID,[ID],J_DC)))'),
    ("ChaveDoc", "Chave_Documento", 18, None, '=IF(OR([ID]="",[NumDoc]=""),"",[TipoDoc]&"|"&[Serie]&"|"&[NumDoc]&"|"&[NIF])'),
    ("Dup", "Documento_Duplicado", 5, "0", '=IF([ChaveDoc]="",0,IF(INDEX(J_ID,MATCH([ChaveDoc],J_ChaveDoc,0))<>[ID],1,0))'),
    ("IsAbe", "É_Abertura", 5, "0", '=IF([ID]="","",IF([Nat]="Abertura",1,0))'),
    ("IncDR", "Inclui_DR", 5, "0", '=IF([ID]="","",IF(IFERROR(INDEX(NAT_IncDR,MATCH([Nat],NAT_Nome,0)),"S")="S",1,0))'),
    ("RubDR", "Rubrica_DR", 7, None, '=IF([ID]="","",IFERROR(INDEX(PC_RubDR,MATCH([Conta],PC_Cod,0))&"",""))'),
    ("Caixa", "Meio_Monetário", 5, "0", '=IF([ID]="","",IF(ISNUMBER(MATCH([G2],L_CaixaPref,0)),1,0))'),
    ("Fluxo", "Fluxo_Caixa", 12, None, '=IF([ID]="","",IF([Caixa]<>1,"",IF([IsAbe]=1,"Abertura",IFERROR(INDEX(NAT_Fluxo,MATCH([Nat],NAT_Nome,0)),"NÃO CLASSIFICADO"))))'),
    ("LinhaDFC", "Linha_DFC", 26, None, '=IF([ID]="","",IF(OR([Caixa]<>1,[IsAbe]=1),"",IFERROR(INDEX(NAT_DFC,MATCH([Nat],NAT_Nome,0)),"NÃO CLASSIFICADO")))'),
    ("TaxaF", "Taxa_Fiscal", 7, PCT, '=IF(OR([ID]="",[CodF]=""),"",IFERROR(INDEX(TX_Taxa,MATCH([CodF],TX_Cod,0)),"#CÓDIGO"))'),
    ("TaxaEst", "Estado_Taxa", 14, None, '=IF(OR([ID]="",[CodF]=""),"",IFERROR(INDEX(TX_Estado,MATCH([CodF],TX_Cod,0)),""))'),
    ("TipoIVA", "Tipo_IVA", 16, None, '=IF(OR([ID]="",LEFT([CodF],3)<>"IVA"),"",IF(LEFT([Conta],6)="34.5.3","Liquidado",IF(LEFT([Conta],6)="34.5.2","Dedutível",IF(LEFT([Conta],6)="34.5.4","Regularização",IF([Classe]="6","Base — venda s/ IVA","Base — aquisição s/ IVA")))))'),
    ("IVACalc", "IVA_Calculado", 12, NUM, '=IF(OR([ID]="",[TipoIVA]=""),"",IF(ISNUMBER([TaxaF]),ROUND(N([Base])*[TaxaF],2),""))'),
    ("ErrTaxa", "Erro_Taxa", 5, "0", '=IF([ID]="","",IF(OR([TipoIVA]="Liquidado",[TipoIVA]="Dedutível",[TipoIVA]="Regularização",[CodF]="RET_SERV"),IF(ISNUMBER([TaxaF]),IF(ABS([Montante]-N([Base])*[TaxaF])>CFG_Tol,1,0),1),IF([TaxaF]="#CÓDIGO",1,0)))'),
    ("MesFech", "Mês_Encerrado", 5, "0", '=IF([ID]="","",IF(AND(ISNUMBER([Mes]),[Ano]=CFG_Ano),IF(INDEX(FECHO_Estado,[Mes])="S",1,0),0))'),
    ("ErrPer", "Alteração_Período_Encerrado", 5, "0", '=IF([ID]="","",IF([MesFech]=1,IF(OR(N([DtIns])>INDEX(FECHO_Data,[Mes]),N([DtAlt])>INDEX(FECHO_Data,[Mes])),1,0),0))'),
    ("Segreg", "Segregação", 10, None, '=IF([ID]="","",IF([Validador]="","Pendente",IF([Validador]=[User],"Violação","OK")))'),
    ("ParteRel", "Parte_Relacionada", 5, "0", '=IF(OR([ID]="",[NIF]=""),0,IF(IFERROR(INDEX(T_PR,MATCH([NIF],T_NIF,0)),"N")="S",1,0))'),
    ("Mont2", "Desvio²_Montante", 10, NUM0, '=IF(OR([ID]="",[IsAbe]=1),0,([Montante]-AUD_Media)^2)'),
    ("Anomalo", "Valor_Anómalo", 5, "0", '=IF(OR([ID]="",[IsAbe]=1),0,IF([Montante]>AUD_Limite,1,0))'),
    ("Erros", "Erros_Detectados", 40, None,
     '=IF([ID]="","",IF(NOT(ISNUMBER([Data])),"Data inválida; ","")'
     '&IF(AND(ISNUMBER([Data]),[Ano]<>CFG_Ano),"Fora do exercício; ","")'
     '&IF([ContaOK]=0,"Conta inexistente; ","")'
     '&IF(AND([ContaOK]=1,[ContaMov]=0),"Conta não movimentável; ","")'
     '&IF(ABS([SaldoID])>0.005,"ERRO — LANÇAMENTO NÃO EQUILIBRADO; ","")'
     '&IF(OR(N([Deb])<0,N([Cred])<0),"Valor negativo; ","")'
     '&IF(OR(AND(N([Deb])>0,N([Cred])>0),[Montante]=0),"Débito/Crédito inválido na linha; ","")'
     '&IF([Dup]=1,"Documento duplicado; ","")'
     '&IF(AND([ContaCorr]="S",OR([NIF]="",LEFT([Terceiro],1)="#",[NIFfmt]=0)),"NIF inválido/não registado; ","")'
     '&IF([Suporte]="","Sem documento de suporte; ","")'
     '&IF([ErrPer]=1,"Alteração em período encerrado; ","")'
     '&IF([ErrTaxa]=1,"Taxa fiscal incompatível; ","")'
     '&IF(AND(OR([Conta]="34.5.2",[Conta]="34.5.3",LEFT([Conta],6)="34.5.4"),[CodF]="",[Nat]<>"Apuramento de IVA"),"IVA sem código fiscal; ","")'
     '&IF(AND(N([Ref])<>0,COUNTIF(J_ID,[Ref])=0),"Ref_Origem inexistente; ","")'
     '&IF([Nat]="","Natureza em falta; ","")'
     '&IF(AND([Nat]<>"",ISNA(MATCH([Nat],NAT_Nome,0))),"Natureza não parametrizada; ","")'
     '&IF(AND([Moeda]<>"",[Moeda]<>CFG_Moeda,OR(N([Cambio])=0,ABS(N([ValME])*N([Cambio])-[Montante])>CFG_Tol)),"Câmbio inconsistente; ","")'
     '&IF(AND([Caixa]=1,[Fluxo]="NÃO CLASSIFICADO"),"Fluxo de caixa não classificado; ",""))'),
    ("ErrFlag", "Tem_Erro", 5, "0", '=IF([ID]="",0,IF([Erros]="",0,1))'),
    ("EstVal", "Estado_Validação", 14, None, '=IF([ID]="","",IF([ErrFlag]=1,"🔴 ERRO",IF([Segreg]="OK","🟢 Validado",IF([Segreg]="Violação","🟡 Segregação","🟡 Por aprovar"))))'),
    ("EstFisc", "Estado_Fiscal", 16, None, '=IF(OR([ID]="",[CodF]=""),"",IF([ErrTaxa]=1,"🔴 Inconsistente",IF(ISNUMBER(SEARCH("VALIDAR",[TaxaEst])),"🟡 Taxa por validar","🟢 Conforme parametrização")))'),
    ("Primeira", "Primeira_Linha_ID", 5, "0", '=IF([ID]="",0,IF(MATCH([ID],J_ID,0)=ROW()-5,1,0))'),
    ("EmitDoc", "Documento_Emitido", 5, "@", '=IF([ID]="","",IFERROR(INDEX(TD_Emit,MATCH([TipoDoc],TD_Cod,0)),"N"))'),
    ("SeqFact", "Seq_Facturação", 6, "0", '=IF(AND([Primeira]=1,[EmitDoc]="S"),1,0)+N([SeqFact-1])'),
    ("SeqCli", "Seq_Clientes", 6, "0", '=IF(AND([ID]<>"",[G2]="31",N([Deb])>0,N([Ref])=0,N([Data])<=CFG_DataRef),1,0)+N([SeqCli-1])'),
    ("SeqForn", "Seq_Fornecedores", 6, "0", '=IF(AND([ID]<>"",[G2]="32",N([Cred])>0,N([Ref])=0,N([Data])<=CFG_DataRef),1,0)+N([SeqForn-1])'),
    ("SeqCaixa", "Seq_Caixa", 6, "0", '=IF(AND([ID]<>"",[G2]="45",[Ano]=CFG_Ano,N([Data])<=CFG_DataRef),1,0)+N([SeqCaixa-1])'),
    ("SeqDG", "Seq_Diário_Geral", 6, "0",
     '=IF(AND([ID]<>"",[Ano]=CFG_Ano,N([Mes])>=DG_MesIni,N([Mes])<=DG_MesFim,OR(DG_Conta="",LEFT([Conta],LEN(DG_Conta))=DG_Conta&""),OR(DG_NIF="",[NIF]=DG_NIF&""),OR(DG_CC="",[CC]=DG_CC&""),OR(DG_Proj="",[Proj]=DG_Proj&""),OR(DG_Nat="",[Nat]=DG_Nat&""),OR(DG_User="",[User]=DG_User&""),OR(DG_TipoDoc="",[TipoDoc]=DG_TipoDoc&""),OR(DG_Doc="",[NumDoc]&""=DG_Doc&"")),1,0)+N([SeqDG-1])'),
    ("SeqRZ", "Seq_Razão", 6, "0",
     '=IF(AND([ID]<>"",[Ano]=CFG_Ano,[IsAbe]=0,N([Mes])>=RZ_MesIni,N([Mes])<=RZ_MesFim,OR([Conta]=RZ_Conta&"",LEFT([Conta],LEN(RZ_Conta)+1)=RZ_Conta&".",AND(LEN(RZ_Conta)=1,[Classe]=RZ_Conta&"")),OR(RZ_NIF="",[NIF]=RZ_NIF&""),OR(RZ_Proj="",[Proj]=RZ_Proj&""),OR(RZ_CC="",[CC]=RZ_CC&"")),1,0)+N([SeqRZ-1])'),
    ("SeqArt", "Seq_Kardex", 6, "0", '=IF(AND([ID]<>"",[Classe]="2",[Artigo]=INV_ArtSel&"",N([Qtd])<>0),1,0)+N([SeqArt-1])'),
]

JCOLS = {}


def _journal_cols():
    i = 1
    for k, *_ in J_IN:
        JCOLS[k] = CL(i)
        i += 1
    for k, *_ in J_CALC:
        JCOLS[k] = CL(i)
        i += 1


_journal_cols()


def build_config(wb):
    ws = wb.create_sheet(S_CFG)
    title(ws, "00 — CONFIGURAÇÃO DA ENTIDADE", "Parametrização central. Todas as folhas lêem estes valores; nada é repetido manualmente.", "CAMADA 1 — INPUT")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 60
    rows = [
        ("IDENTIFICAÇÃO", None, None, None, None),
        ("Nome da entidade", "EMPRESA MODELO, LDA (fictícia)", "CFG_Nome", None, "Entidade fictícia para demonstração."),
        ("Denominação social", "Empresa Modelo — Comércio e Serviços, Lda", "CFG_Denom", None, ""),
        ("NIF", "5000000000", "CFG_NIF", "@", "10 dígitos (pessoa colectiva). Validação de existência requer consulta à AGT."),
        ("Forma jurídica", "Sociedade por quotas", "CFG_Forma", None, ""),
        ("Sector de actividade (CAE)", "Comércio a retalho e consultoria", "CFG_Sector", None, ""),
        ("Sede", "Rua Exemplo, n.º 1", "CFG_Sede", None, ""),
        ("Município", "Luanda", "CFG_Municipio", None, ""),
        ("Província", "Luanda", "CFG_Provincia", None, ""),
        ("Contactos", "geral@empresamodelo.ao | +244 900 000 000", "CFG_Contactos", None, "Fictício."),
        ("Ano económico (exercício)", D.ANO, "CFG_Ano", "0", "Exercício a que respeitam os lançamentos. Dados de outros anos ficam no histórico e não entram nos mapas."),
        ("Mês de reporte (1–12)", 3, "CFG_MesRep", "0", "Todos os mapas acumulados (Balanço, DR, KPI) são calculados até este mês."),
        ("ENQUADRAMENTO", None, None, None, None),
        ("Regime fiscal (Imposto Industrial)", "Regime Geral", "CFG_RegII", None, "Regime Geral / Regime Simplificado — confirmar enquadramento na AGT."),
        ("Regime de IVA", "Regime Geral", "CFG_RegIVA", None, "Regime Geral / Simplificado / Exclusão — confirmar."),
        ("Regime contabilístico", "PGC — Decreto n.º 82/01", "CFG_Plano", None, "Algumas entidades usam IFRS ou planos sectoriais (BNA, ARSEG): ver 01_PLANO_CONTAS."),
        ("Tipo de entidade", "Sociedade comercial", "CFG_TipoEnt", None, ""),
        ("Dimensão (micro/pequena/média/grande)", "Pequena empresa", "CFG_Dim", None, "Classificação legal a confirmar (critérios da legislação das MPME)."),
        ("Grande Contribuinte (S/N)", "N", "CFG_GC", None, "Determina a data de obrigatoriedade da facturação electrónica (DP 71/25)."),
        ("Moeda funcional", "AOA", "CFG_Moeda", None, "Kwanza."),
        ("NIF genérico de consumidor final", "999999999", "CFG_NIFCF", "@", "PARÂMETRO DO MODELO — confirmar o identificador exigido pela AGT/software validado."),
        ("Prefixos das contas de meios monetários (DFC)", "43", "CFG_Caixa1", "@", "Caixa e equivalentes para a DFC."),
        ("", "45", "CFG_Caixa2", "@", ""),
        ("", "", "CFG_Caixa3", "@", "Acrescente, se aplicável (ex.: 41 se forem equivalentes de caixa)."),
        ("PARÂMETROS FINANCEIROS", None, None, None, None),
        ("Salário mínimo nacional (Kz)", None, "CFG_SalMin", NUM, "POR VALIDAR — introduzir o valor do diploma em vigor."),
        ("Taxa de juro de referência (BNA)", None, "CFG_TxBNA", PCT, "POR VALIDAR — introduzir a taxa BNA vigente."),
        ("Taxa de juro média da dívida", 0.18, "CFG_TxJuro", PCT, "Parâmetro do modelo (contrato de mútuo de teste)."),
        ("Taxa de câmbio USD/AOA (fecho)", 920, "CFG_USD", "#,##0.00", "Parâmetro — actualizar com a taxa de referência do BNA à data de reporte."),
        ("Taxa de câmbio EUR/AOA (fecho)", 1000, "CFG_EUR", "#,##0.00", "Parâmetro — actualizar."),
        ("Prazo médio de recebimento — alvo (dias)", 45, "CFG_PMRalvo", "0", "Política interna."),
        ("Prazo médio de pagamento — alvo (dias)", 45, "CFG_PMPalvo", "0", "Política interna."),
        ("Limite de tesouraria (linha de crédito disponível)", 2_000_000, "CFG_LimTes", NUM, "Política interna / contrato bancário."),
        ("Reserva mínima de liquidez (Kz)", 1_500_000, "CFG_SaldoMin", NUM, "Política interna — saldo mínimo de caixa + bancos."),
        ("Fundo fixo de caixa (Kz)", 300_000, "CFG_FundoFixo", NUM, "Política interna."),
        ("Tolerância de arredondamento (Kz)", 1, "CFG_Tol", NUM, "Usada nas validações de IVA, câmbio e reconciliações."),
        ("Depreciação: contar mês de aquisição (1=Sim, 0=Não)", 1, "CFG_IniDep", "0", "Política contabilística — confirmar regra fiscal aplicável."),
        ("Auditoria: nº de desvios-padrão para valor anómalo", 3, "CFG_kAnom", "0", "Teste estatístico de valores anormais."),
        ("Vendas + serviços do exercício anterior (Kz)", 12_000_000, "CFG_VendasN1", NUM, "Da DR do exercício anterior (fictício)."),
        ("EBITDA do exercício anterior (Kz)", 3_000_000, "CFG_EBITDAN1", NUM, "Da DR do exercício anterior (fictício)."),
        ("Probabilidade de cobrança de saldos vencidos (tesouraria)", 0.5, "CFG_ProbCob", PCT, "Pressuposto da previsão de 13 semanas."),
        ("Salários líquidos mensais previstos (Kz)", 348_000, "CFG_SalPrev", NUM, "Pressuposto da tesouraria (a pagar na última semana do mês)."),
        ("Dia de pagamento de salários", 28, "CFG_DiaSal", "0", "Pressuposto da tesouraria."),
        ("Serviço da dívida mensal previsto (Kz)", 45_000, "CFG_DividaPrev", NUM, "Juros mensais do empréstimo de teste."),
        ("LIMITES DE ALERTA (parametrizáveis — política da entidade)", None, None, None, None),
        ("Concentração máxima num cliente (% saldo)", 0.40, "CFG_ConcCli", PCT, "Referência técnica inicial: ajustar à política de risco."),
        ("Concentração máxima num fornecedor (% saldo)", 0.40, "CFG_ConcForn", PCT, "Idem."),
        ("Dívida financeira / EBITDA máximo", 3.0, "CFG_DivEBITDA", MULT, "Covenant típico — ajustar ao contrato de financiamento."),
        ("Margem líquida mínima", 0.05, "CFG_MargMin", PCT, "Política interna."),
        ("Queda anormal da receita (mês vs média)", -0.25, "CFG_QuedaRec", PCT, "Alerta quando a receita do mês cai abaixo deste desvio face à média do ano."),
        ("Aumento anormal de custos (mês vs média)", 0.25, "CFG_AumCust", PCT, ""),
        ("Tolerância de desvio orçamental", 0.10, "CFG_OrcTol", PCT, "Desvio desfavorável acima do qual o orçamento se considera excedido."),
        ("Exposição cambial máxima (% passivo em ME)", 0.30, "CFG_ExpCamb", PCT, ""),
        ("Dias sem movimento para conta 'sem movimento esperado'", 45, "CFG_DiasSemMov", "0", "Aplica-se a contas marcadas com Reconciliação = S."),
    ]
    r = 5
    for lab, val, nm, fmt, note in rows:
        if val is None and nm is None:
            section(ws, r, 1, lab, 3)
            r += 1
            continue
        put(ws, f"A{r}", lab, "label")
        if nm:
            put(ws, f"B{r}", val, "input", fmt)
            name(wb, nm, S_CFG, f"$B${r}")
        put(ws, f"C{r}", note, "note", wrap=True)
        r += 1
    # derivados (calculados)
    section(ws, r, 1, "DERIVADOS (calculados — não editar)", 3)
    r += 1
    for lab, f, nm, fmt in [
        ("Data de início do exercício", "=DATE(CFG_Ano,1,1)", "CFG_DataIni", DATE),
        ("Data de referência (fim do mês de reporte)", "=EOMONTH(DATE(CFG_Ano,CFG_MesRep,1),0)", "CFG_DataRef", DATE),
        ("Dias decorridos no exercício", "=CFG_DataRef-CFG_DataIni+1", "CFG_Dias", "0"),
        ("Factor de anualização", "=12/CFG_MesRep", "CFG_Anual", "0.00"),
    ]:
        put(ws, f"A{r}", lab, "label")
        put(ws, f"B{r}", f, "grey", fmt)
        name(wb, nm, S_CFG, f"$B${r}")
        r += 1
    # lista de prefixos de caixa
    cells = [wb.defined_names[n].attr_text.split("!")[1] for n in ("CFG_Caixa1", "CFG_Caixa3")]
    name(wb, "L_CaixaPref", S_CFG, cells[0] + ":" + cells[1])
    mr = wb.defined_names["CFG_MesRep"].attr_text.split("!")[1].replace("$", "")
    dv_list(ws, mr, '"1,2,3,4,5,6,7,8,9,10,11,12"')
    protect(ws)
    return ws


def build_tables(wb):
    ws = wb.create_sheet(S_TAB)
    title(ws, "01A — TABELAS DE PARAMETRIZAÇÃO", "Naturezas de operação, rubricas das demonstrações, tipos de documento, linhas da DFC.", "CAMADA 1 — INPUT (parametrização estrutural)")
    # Naturezas
    header(ws, 5, 1, ["Natureza da operação", "Fluxo de caixa", "Linha DFC", "Inclui na DR (S/N)"], [30, 14, 42, 10])
    for i, (n, f, l, inc) in enumerate(D.NATUREZAS):
        r = 6 + i
        for j, v in enumerate((n, f, l, inc)):
            put(ws, (r, 1 + j), v, "input")
    n1 = 6 + 59
    name(wb, "NAT_Nome", S_TAB, f"$A$6:$A${n1}")
    name(wb, "NAT_Fluxo", S_TAB, f"$B$6:$B${n1}")
    name(wb, "NAT_DFC", S_TAB, f"$C$6:$C${n1}")
    name(wb, "NAT_IncDR", S_TAB, f"$D$6:$D${n1}")
    for r in range(6 + len(D.NATUREZAS), n1 + 1):
        for j in range(4):
            put(ws, (r, 1 + j), None, "input")
    name(wb, "L_Natureza", S_TAB, f"$A$6:$A${5 + len(D.NATUREZAS)}")
    # Tipos de documento
    header(ws, 5, 6, ["Tipo doc.", "Descrição", "Emitido pela entidade", "Tipo SAF-T"], [8, 36, 10, 20])
    for i, row in enumerate(D.TIPOS_DOC):
        for j, v in enumerate(row):
            put(ws, (6 + i, 6 + j), v, "input")
    t1 = 5 + len(D.TIPOS_DOC)
    name(wb, "TD_Cod", S_TAB, f"$F$6:$F${t1}")
    name(wb, "TD_Emit", S_TAB, f"$H$6:$H${t1}")
    name(wb, "TD_SAFT", S_TAB, f"$I$6:$I${t1}")
    # Rubricas de balanço
    header(ws, 5, 11, ["Rubrica balanço", "Designação", "Secção"], [14, 40, 20])
    for i, row in enumerate(D.RUBRICAS_BAL):
        for j, v in enumerate(row):
            put(ws, (6 + i, 11 + j), v, "input")
    b1 = 5 + len(D.RUBRICAS_BAL)
    name(wb, "RB_Cod", S_TAB, f"$K$6:$K${b1}")
    name(wb, "RB_Nome", S_TAB, f"$L$6:$L${b1}")
    # Rubricas DR
    header(ws, 5, 15, ["Rubrica DR", "Designação", "Natureza", "Sinal (S=proveito/C=custo)", "% variável (break-even)"], [9, 44, 14, 10, 10])
    var = {"DR04": 1.0, "DR05": 0.3, "DR06": 0.0, "DR07": 0.0, "DR08": 0.5, "DR10": 0.0, "DR12": 0.0, "DR13": 0.0}
    for i, row in enumerate(D.RUBRICAS_DR):
        for j, v in enumerate(row):
            put(ws, (6 + i, 15 + j), v, "input")
        put(ws, (6 + i, 19), var.get(row[0], 0.0) if row[3] == "C" else None, "input", PCT)
    d1 = 5 + len(D.RUBRICAS_DR)
    name(wb, "RD_Cod", S_TAB, f"$O$6:$O${d1}")
    name(wb, "RD_Nome", S_TAB, f"$P$6:$P${d1}")
    name(wb, "RD_Var", S_TAB, f"$S$6:$S${d1}")
    # Linhas DFC
    header(ws, 5, 21, ["Actividade", "Linha DFC"], [14, 42])
    for i, (a, l) in enumerate(D.DFC_LINHAS):
        put(ws, (6 + i, 21), a, "input")
        put(ws, (6 + i, 22), l, "input")
    # Listas simples
    header(ws, 5, 24, ["Tipos de conta", "S/N", "Moedas", "Formas de pagamento", "Centros de custo", "Estado comunicação AGT", "Estado documento", "Tipos de terceiro"], [12, 5, 7, 16, 10, 16, 12, 12])
    lists = [
        ["Movimento", "Agregação"], ["S", "N"], ["AOA", "USD", "EUR", "ZAR", "CNY"],
        ["Numerário", "Transferência", "Multicaixa/TPA", "Cheque", "Débito directo", "Compensação"],
        ["ADM", "COM", "OPS", "FIN", "RH"], ["Comunicado", "Pendente", "Erro", "Não aplicável"],
        ["Normal", "Anulado", "Rectificado"], ["Cliente", "Fornecedor", "Cliente/Fornecedor", "Pessoal", "Estado", "Banco", "Outro"],
    ]
    names_ = ["L_TipoConta", "L_SN", "L_Moeda", "L_FormaPag", "L_CC", "L_EstAGT", "L_EstDoc", "L_TipoTerc"]
    for j, (lst, nm) in enumerate(zip(lists, names_)):
        col = 24 + j
        for i, v in enumerate(lst):
            put(ws, (6 + i, col), v, "input")
        name(wb, nm, S_TAB, f"${CL(col)}$6:${CL(col)}${5 + len(lst)}")
    protect(ws)
    return ws


def build_plano(wb):
    ws = wb.create_sheet(S_PC)
    title(ws, "01 — PLANO DE CONTAS (PGC — Decreto n.º 82/01, com subcontas analíticas sugeridas)",
          "Classe → Conta → Subconta → Sub-subconta → Analítica. Apenas contas 'Movimento' recebem lançamentos. Para outro plano, substitua as linhas mantendo as colunas.",
          "CAMADA 1 — INPUT (A:R)  |  CAMADA 2 — PROCESSAMENTO (S:AF, cinzento)")
    heads = ["Código", "Designação", "Classe", "Nível", "Tipo (Movimento/Agregação)", "Natureza (saldo normal)", "Balanço/Resultado",
             "Rubrica balanço (saldo devedor)", "Rubrica balanço (saldo credor)", "Rubrica DR", "Rubrica fiscal", "C. custo permitido",
             "Projecto permitido", "Conta corrente", "Reconciliação", "Indicador fiscal", "Encerra no fim do exercício", "Plano de origem",
             "Débito acumulado", "Crédito acumulado", "Saldo (D−C)", "Saldo devedor", "Saldo credor", "Saldo abertura (D−C)",
             "Abertura devedor", "Abertura credor", "Validação", "Seq. encerramento", "Seq. abertura N+1", "Último movimento", "Nº linhas", "aux Déb (mov.)", "aux Créd (mov.)", "aux Abertura (mov.)"]
    widths = [9, 44, 6, 6, 12, 10, 10, 15, 15, 8, 8, 6, 6, 6, 6, 7, 8, 12, 14, 14, 14, 14, 14, 14, 14, 14, 26, 6, 6, 11, 6, 12, 12, 12]
    header(ws, 5, 1, heads, widths)
    ws.freeze_panes = "C6"
    crit = ',J_Ano,CFG_Ano,J_Mes,"<="&CFG_MesRep'

    def acc(measure, extra):
        return (f'IF($D{{r}}=1,SUMIFS({measure},J_Classe,$A{{r}}{extra}),'
                f'SUMIFS({measure},J_Conta,$A{{r}}{extra})+SUMIFS({measure},J_Conta,$A{{r}}&".*"{extra}))')

    for i in range(PR1 - PR0 + 1):
        r = PR0 + i
        if i < len(D.PLANO):
            code, des, t = D.PLANO[i]
            nat, br, rd, rc, dr, ccorr, rec, fis, enc = D.classify(code)
            if t == "A":
                rd = rc = dr = ""
            vals = [code, des, None, None, "Movimento" if t == "M" else "Agregação", nat, br, rd, rc, dr, fis, "S", "S" if code[0] in "1267" else "N",
                    ccorr, rec, "S" if fis else "N", enc, "PGC 82/01" if code.count(".") < 1 else "Analítica (sugerida)"]
        else:
            vals = [None] * 18
        for j, v in enumerate(vals):
            if j in (2, 3):
                continue
            put(ws, (r, 1 + j), v, "input", "@" if j == 0 else None)
        put(ws, (r, 3), f'=IF(A{r}="","",LEFT(A{r},1))', "grey")
        put(ws, (r, 4), f'=IF(A{r}="","",IF(LEN(A{r})=1,1,2+LEN(A{r})-LEN(SUBSTITUTE(A{r},".",""))))', "grey")
        # colunas auxiliares AF:AH = valores apenas das contas de movimento (evita referências circulares)
        put(ws, (r, 32), f'=IF(OR($A{r}="",$E{r}<>"Movimento"),0,SUMIFS(J_Deb,J_Conta,$A{r}{crit}))', "grey", NUM)
        put(ws, (r, 33), f'=IF(OR($A{r}="",$E{r}<>"Movimento"),0,SUMIFS(J_Cred,J_Conta,$A{r}{crit}))', "grey", NUM)
        put(ws, (r, 34), f'=IF(OR($A{r}="",$E{r}<>"Movimento"),0,SUMIFS(J_DC,J_Conta,$A{r},J_Ano,CFG_Ano,J_IsAbe,1))', "grey", NUM)

        def agg(h):
            return (f'=IF($A{r}="","",IF($E{r}="Movimento",{h}{r},IF($D{r}=1,SUMIFS(${h}${PR0}:${h}${PR1},$C${PR0}:$C${PR1},$A{r},$E${PR0}:$E${PR1},"Movimento"),'
                    f'SUMIFS(${h}${PR0}:${h}${PR1},$A${PR0}:$A${PR1},$A{r}&".*",$E${PR0}:$E${PR1},"Movimento"))))')
        put(ws, (r, 19), agg("AF"), "grey", NUM)
        put(ws, (r, 20), agg("AG"), "grey", NUM)
        put(ws, (r, 21), f'=IF($A{r}="","",S{r}-T{r})', "grey", NUM)
        put(ws, (r, 22), f'=IF($A{r}="","",IF($E{r}="Movimento",MAX(0,U{r}),0))', "grey", NUM)
        put(ws, (r, 23), f'=IF($A{r}="","",IF($E{r}="Movimento",MAX(0,-U{r}),0))', "grey", NUM)
        put(ws, (r, 24), agg("AH"), "grey", NUM)
        put(ws, (r, 25), f'=IF($A{r}="","",IF($E{r}="Movimento",MAX(0,X{r}),0))', "grey", NUM)
        put(ws, (r, 26), f'=IF($A{r}="","",IF($E{r}="Movimento",MAX(0,-X{r}),0))', "grey", NUM)
        put(ws, (r, 27), f'=IF(A{r}="","",IF(COUNTIF(PC_Cod,A{r})>1,"🔴 Código duplicado",IF(AND(E{r}="Movimento",OR(H{r}="",I{r}="")),"🔴 Sem rubrica de balanço",'
                         f'IF(AND(E{r}="Movimento",OR(C{r}="6",C{r}="7",LEFT(A{r},2)="87"),J{r}=""),"🔴 Sem rubrica DR",'
                         f'IF(AND(E{r}="Agregação",COUNTIF(J_Conta,A{r})>0),"🔴 Movimentos em conta de agregação","🟢 OK")))))', "grey")
        put(ws, (r, 28), f'=IF(A{r}="",N(AB{r - 1}),IF(AND(E{r}="Movimento",OR(C{r}="6",C{r}="7",AND(C{r}="8",VALUE(LEFT(A{r},2))>=82,VALUE(LEFT(A{r},2))<=87)),ABS(N(U{r}))>0.005),1,0)+N(AB{r - 1}))', "grey", "0")
        put(ws, (r, 29), f'=IF(AND(E{r}="Movimento",OR(C{r}="1",C{r}="2",C{r}="3",C{r}="4",C{r}="5",LEFT(A{r},2)="81",LEFT(A{r},2)="89"),ABS(N(U{r}))>0.005),1,0)+N(AC{r - 1})', "grey", "0")
        put(ws, (r, 30), f'=IF(OR(A{r}="",E{r}<>"Movimento"),"",IF(COUNTIF(J_Conta,A{r})=0,"",_xlfn.MAXIFS(J_Data,J_Conta,A{r})))', "grey", DATE)
        put(ws, (r, 31), f'=IF(A{r}="","",COUNTIF(J_Conta,A{r}))', "grey", "0")
    rng = lambda c: f"${c}${PR0}:${c}${PR1}"
    for nm, c in [("PC_Cod", "A"), ("PC_Desig", "B"), ("PC_Classe", "C"), ("PC_Nivel", "D"), ("PC_Tipo", "E"), ("PC_RubD", "H"), ("PC_RubC", "I"),
                  ("PC_RubDR", "J"), ("PC_CCorr", "N"), ("PC_Rec", "O"), ("PC_Deb", "S"), ("PC_Cred", "T"), ("PC_Saldo", "U"), ("PC_Dev", "V"),
                  ("PC_Cr", "W"), ("PC_Abe", "X"), ("PC_AbeDev", "Y"), ("PC_AbeCr", "Z"), ("PC_Val", "AA"), ("PC_SeqEnc", "AB"), ("PC_SeqAbe", "AC"),
                  ("PC_UltMov", "AD")]:
        name(wb, nm, S_PC, rng(c))
    dv_list(ws, f"E{PR0}:E{PR1}", "=L_TipoConta")
    for c in "LMNOQ":
        dv_list(ws, f"{c}{PR0}:{c}{PR1}", "=L_SN")
    dv_list(ws, f"H{PR0}:I{PR1}", "=RB_Cod")
    dv_list(ws, f"J{PR0}:J{PR1}", "=RD_Cod")
    status_cf(ws, f"AA{PR0}:AA{PR1}")
    ws.auto_filter.ref = f"A5:AE{PR1}"
    protect(ws)
    return ws


def build_terceiros(wb):
    ws = wb.create_sheet(S_TER)
    title(ws, "01B — FICHEIRO DE TERCEIROS (clientes, fornecedores, outros)",
          "O Diário só recebe o NIF; nome, tipo e país são lidos daqui (princípio da não duplicação).", "CAMADA 1 — INPUT (A:G)  |  CAMADA 2 — PROCESSAMENTO (H:P)")
    heads = ["NIF", "Nome / denominação", "Tipo de terceiro", "País", "Prazo (dias)", "Limite de crédito (Kz)", "Parte relacionada (S/N)",
             "NIF formato OK", "Saldo cliente (31)", "Saldo fornecedor (32)", "Facturado a cliente (ano)", "Facturado por fornecedor (ano)",
             "Vencido > 90 dias (cliente)", "Excede limite de crédito", "Quota no saldo de clientes", "Quota no saldo de fornecedores"]
    header(ws, 5, 1, heads, [13, 32, 14, 12, 8, 14, 9, 8, 14, 14, 14, 14, 14, 10, 10, 10])
    ws.freeze_panes = "C6"
    ref = '"<="&CFG_DataRef'
    for i in range(TR1 - TR0 + 1):
        r = TR0 + i
        vals = list(D.TERCEIROS[i]) if i < len(D.TERCEIROS) else [None] * 7
        for j, v in enumerate(vals):
            put(ws, (r, 1 + j), v, "input", "@" if j == 0 else (NUM if j == 5 else None))
        put(ws, (r, 8), f'=IF(A{r}="","",IF(D{r}<>"Angola",1,IF(OR(A{r}=CFG_NIFCF,AND(LEN(A{r})=10,ISNUMBER(-A{r})),AND(LEN(A{r})=14,ISNUMBER(-LEFT(A{r},9)),ISNUMBER(-RIGHT(A{r},3)))),1,0)))', "grey", "0")
        put(ws, (r, 9), f'=IF(A{r}="","",SUMIFS(J_DC,J_NIF,A{r},J_G2,"31",J_Ano,CFG_Ano,J_Data,{ref}))', "grey", NUM)
        put(ws, (r, 10), f'=IF(A{r}="","",-SUMIFS(J_DC,J_NIF,A{r},J_G2,"32",J_Ano,CFG_Ano,J_Data,{ref}))', "grey", NUM)
        put(ws, (r, 11), f'=IF(A{r}="","",SUMIFS(J_Deb,J_NIF,A{r},J_G2,"31",J_Ano,CFG_Ano,J_Data,{ref},J_Ref,""))', "grey", NUM)
        put(ws, (r, 12), f'=IF(A{r}="","",SUMIFS(J_Cred,J_NIF,A{r},J_G2,"32",J_Ano,CFG_Ano,J_Data,{ref},J_Ref,""))', "grey", NUM)
        put(ws, (r, 13), f'=IF(A{r}="","",SUMIFS(CLI_V91,CLI_NIF,A{r}))', "grey", NUM)
        put(ws, (r, 14), f'=IF(OR(A{r}="",N(F{r})=0),"",IF(I{r}>F{r},"🔴 Excede","🟢 OK"))', "grey")
        put(ws, (r, 15), f'=IF(OR(A{r}="",SUMIF(T_SaldoCli,">0")=0),"",MAX(0,I{r})/SUMIF(T_SaldoCli,">0"))', "grey", PCT)
        put(ws, (r, 16), f'=IF(OR(A{r}="",SUMIF(T_SaldoForn,">0")=0),"",MAX(0,J{r})/SUMIF(T_SaldoForn,">0"))', "grey", PCT)
    rng = lambda c: f"${c}${TR0}:${c}${TR1}"
    for nm, c in [("T_NIF", "A"), ("T_Nome", "B"), ("T_Tipo", "C"), ("T_Pais", "D"), ("T_Prazo", "E"), ("T_Limite", "F"), ("T_PR", "G"),
                  ("T_NIFok", "H"), ("T_SaldoCli", "I"), ("T_SaldoForn", "J"), ("T_FactCli", "K"), ("T_FactForn", "L"), ("T_V91", "M"),
                  ("T_ExcLim", "N"), ("T_QuotaCli", "O"), ("T_QuotaForn", "P")]:
        name(wb, nm, S_TER, rng(c))
    dv_list(ws, f"C{TR0}:C{TR1}", "=L_TipoTerc")
    dv_list(ws, f"G{TR0}:G{TR1}", "=L_SN")
    status_cf(ws, f"N{TR0}:N{TR1}")
    protect(ws)
    return ws


def build_journal(wb):
    ws = wb.create_sheet(S_J)
    title(ws, "02 — DIÁRIO DE LANÇAMENTOS (FONTE ÚNICA DE DADOS)",
          "Partidas dobradas: cada ID_Lançamento tem ≥2 linhas e Σ Débito = Σ Crédito. Colunas azuis = input; cinzentas = processamento automático (não editar).",
          "CAMADA 1 — INPUT (A:AL)  |  CAMADA 2 — PROCESSAMENTO (AM em diante)")
    labels = [h for _, h, _, _ in J_IN] + [h for _, h, _, _, _ in J_CALC]
    widths = [w for _, _, w, _ in J_IN] + [w for _, _, w, _, _ in J_CALC]
    header(ws, 5, 1, labels, widths)
    for j in range(len(J_IN), len(labels)):
        ws.cell(row=5, column=1 + j).fill = fill("595959")
    ws.freeze_panes = "B6"
    # resumo de controlo no topo
    put(ws, "D3", "Linhas:", "label")
    put(ws, "E3", "=COUNT(J_ID)", "grey", "0")
    put(ws, "F3", "Σ Débito:", "label")
    put(ws, "G3", "=SUM(J_Deb)", "grey", NUM)
    put(ws, "H3", "Σ Crédito:", "label")
    put(ws, "I3", "=SUM(J_Cred)", "grey", NUM)
    put(ws, "J3", '=IF(ABS(G3-I3)>0.005,"🔴 ERRO — DIÁRIO NÃO EQUILIBRADO","🟢 Diário equilibrado")&" | Linhas com erro: "&SUM(J_ErrFlag)', "grey")
    status_cf(ws, "J3")
    data_rows = D.JOURNAL
    for i in range(JR1 - JR0 + 1):
        r = JR0 + i
        rec = data_rows[i] if i < len(data_rows) else None
        for j, (k, _, _, fmt) in enumerate(J_IN):
            v = rec[k] if rec else None
            c = ws.cell(row=r, column=1 + j, value=v)
            c.font = font(False, C_INPUT_FONT)
            c.protection = Protection(locked=False)
            if fmt:
                c.number_format = fmt
        for j, (k, _, _, fmt, tpl) in enumerate(J_CALC):
            c = ws.cell(row=r, column=len(J_IN) + 1 + j, value=render(tpl, JCOLS, r))
            if fmt:
                c.number_format = fmt
    # fundo de input para a área (apenas cabeçalho de colunas é suficiente visualmente)
    for j in range(len(J_IN)):
        ws.cell(row=4, column=1 + j).fill = fill(C_INPUT_FILL)
    for j in range(len(J_CALC)):
        ws.cell(row=4, column=len(J_IN) + 1 + j).fill = fill(C_CALC_FILL)
    for k, col in JCOLS.items():
        name(wb, "J_" + k, S_J, f"${col}${JR0}:${col}${JR1}")
    # validações de dados
    dv_list(ws, f"{JCOLS['TipoDoc']}{JR0}:{JCOLS['TipoDoc']}{JR1}", "=TD_Cod")
    dv_list(ws, f"{JCOLS['Conta']}{JR0}:{JCOLS['Conta']}{JR1}", "=PC_Cod")
    dv_list(ws, f"{JCOLS['Nat']}{JR0}:{JCOLS['Nat']}{JR1}", "=L_Natureza")
    dv_list(ws, f"{JCOLS['CodF']}{JR0}:{JCOLS['CodF']}{JR1}", "=TX_Cod")
    dv_list(ws, f"{JCOLS['Moeda']}{JR0}:{JCOLS['Moeda']}{JR1}", "=L_Moeda")
    dv_list(ws, f"{JCOLS['FormaPag']}{JR0}:{JCOLS['FormaPag']}{JR1}", "=L_FormaPag")
    dv_list(ws, f"{JCOLS['CC']}{JR0}:{JCOLS['CC']}{JR1}", "=L_CC")
    dv_list(ws, f"{JCOLS['EstAGT']}{JR0}:{JCOLS['EstAGT']}{JR1}", "=L_EstAGT")
    dv_list(ws, f"{JCOLS['EstDoc']}{JR0}:{JCOLS['EstDoc']}{JR1}", "=L_EstDoc")
    dv_list(ws, f"{JCOLS['NIF']}{JR0}:{JCOLS['NIF']}{JR1}", "=T_NIF")
    from openpyxl.worksheet.datavalidation import DataValidation
    dvd = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)", showErrorMessage=True,
                         errorTitle="Data inválida", error="Introduza uma data válida.")
    ws.add_data_validation(dvd)
    dvd.add(f"B{JR0}:B{JR1}")
    dvn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True, errorTitle="Valor negativo",
                         error="Débito e Crédito não podem ser negativos. Use o lado oposto.")
    ws.add_data_validation(dvn)
    dvn.add(f"K{JR0}:L{JR1}")
    status_cf(ws, f"{JCOLS['EstVal']}{JR0}:{JCOLS['EstFisc']}{JR1}")
    value_cf(ws, f"{JCOLS['Erros']}{JR0}:{JCOLS['Erros']}{JR1}", f'LEN({JCOLS["Erros"]}{JR0})>0', C_ERR)
    last = CL(len(J_IN) + len(J_CALC))
    ws.auto_filter.ref = f"A5:{last}{JR1}"
    protect(ws)
    return ws
