"""Gera MATRIZ_PRO_MASTER_CONTABILIDADE_ANGOLA.xlsx.

Uso:  python src/build.py [saida.xlsx]
      MATRIZ_LINHAS=5000 python src/build.py   (capacidade do Diário; por omissão 2000 linhas)
"""
import os
import sys
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(__file__))
import sheets_base as SB  # noqa: E402
import sheets_ledgers as SL  # noqa: E402
import sheets_fiscal as SF  # noqa: E402
import sheets_reports as SR  # noqa: E402
import sheets_analytics as SA  # noqa: E402
import sheets_mgmt as SM  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "dist", "MATRIZ_PRO_MASTER_CONTABILIDADE_ANGOLA.xlsx")

ORDER = [
    SM.build_leiame, SB.build_config, SB.build_plano, SB.build_tables, SB.build_terceiros, SB.build_journal,
    SL.build_diario_geral, SL.build_razao, SL.build_caixa, SL.build_bancos, SL.build_clientes, SL.build_fornecedores,
    SL.build_inventarios, SL.build_activos,
    SF.build_fiscalidade, SF.build_iva, SF.build_facturacao, SF.build_saft, SF.build_ii,
    SR.build_orcamento, SA.build_planeamento, SR.build_tesouraria, SR.build_fluxo, SR.build_balancete, SR.build_balanco,
    SR.build_dre, SR.build_dfc,
    SA.build_projectos, SA.build_investimentos, SA.build_risco, SA.build_sustentabilidade, SA.build_kpi,
    None,  # dashboard (precisa das folhas de origem dos gráficos já criadas)
    SR.build_bva, SM.build_controlo_interno, SM.build_auditoria, SM.build_fecho, SF.build_calendario, SF.build_base_legal,
    SA.build_breakeven, SM.build_encerramento, SM.build_relatorio, SM.build_alertas, SM.build_controlo,
]


LAYER = {"INPUT": "2F5597", "PROC": "7F7F7F", "OUTPUT": "548235"}


def finish(wb):
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.properties import PageSetupProperties
    import sheets_mgmt
    layer = {s: l for s, l, _ in sheets_mgmt.INDEX_SHEETS}
    for ws in wb.worksheets:
        l = layer.get(ws.title, "")
        ws.sheet_properties.tabColor = LAYER["INPUT"] if l.startswith("INPUT") else (LAYER["PROC"] if l.startswith("PROC") else LAYER["OUTPUT"])
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = 9
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.gridLines = False
        ws.oddFooter.center.text = "&A — página &P de &N"
        ws.oddFooter.right.text = "Matriz PRO MASTER — Angola"
    # Tabelas Excel (ListObjects) sobre as bases de dados de input — prontas para Power Query
    spec = [("02_DIÁRIO_LANÇAMENTOS", "tbl_Diario", 5, SB.JR1), ("01_PLANO_CONTAS", "tbl_PlanoContas", 5, SB.PR1),
            ("01B_TERCEIROS", "tbl_Terceiros", 5, SB.TR1), ("10_ACTIVOS_FIXOS", "tbl_Activos", 5, 105)]
    from openpyxl.utils import get_column_letter
    for sh, nm, h, last in spec:
        ws = wb[sh]
        ws.auto_filter.ref = None
        ref = f"A{h}:{get_column_letter(ws.max_column)}{last}"
        t = Table(displayName=nm, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
        ws.add_table(t)


def main(out=OUT):
    wb = Workbook()
    wb.remove(wb.active)
    for fn in ORDER:
        if fn is None:
            continue
        fn(wb)
    SA.build_dashboard(wb)
    # posicionar o dashboard depois de 28_KPI_FINANCEIROS
    ws = wb["29_DASHBOARD_EXECUTIVO"]
    wb.move_sheet(ws, offset=wb.sheetnames.index("30_BUDGET_VS_ACTUAL") - wb.sheetnames.index("29_DASHBOARD_EXECUTIVO"))
    finish(wb)
    wb.active = wb.sheetnames.index("LEIA-ME")
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print("OK", os.path.abspath(out), len(wb.sheetnames), "folhas,", len(wb.defined_names), "nomes definidos")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else OUT)
