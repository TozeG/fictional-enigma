"""Entrada por OPERAÇÕES (modo aplicação): 02A_OPERAÇÕES (input) → 02B_MOTOR (partidas dobradas) → zona automática do Diário.

O utilizador descreve a operação em linguagem de gestão (vendi, comprei, recebi, paguei salários…);
o motor gera os lançamentos, o IVA, o custo das mercadorias (custo médio ponderado móvel), o INSS e o IRT
(tabela oficial do exercício), sem que seja necessário conhecer partidas dobradas.
"""
from openpyxl.utils import get_column_letter as CL
from core import *
import data as D

S_OP = "02A_OPERAÇÕES"
S_MOT = "02B_MOTOR"
OR0 = 7                     # primeira linha de operações
OPS = int(__import__("os").environ.get("MATRIZ_OPERACOES", "300"))
SLOTS = 6
OP_ID0 = 100000             # ID do lançamento gerado = 100000 + nº da operação

# (tipo, natureza, tipo doc se "A crédito", tipo doc se pago/recebido, [(origem da conta, lado, montante)], ajuda)
# Origem: código literal | MEIO_CLI | MEIO_FOR | MEIO_FORI | TES | CONTA:<código por omissão>
TIPOS = [
    ("Venda de mercadorias", "Venda", "FT", "FR",
     [("MEIO_CLI", "D", "TOTAL"), ("CONTA:61.1", "C", "BASE"), ("34.5.3", "C", "IVA"), ("71.1", "D", "CUSTO"), ("26.1", "C", "CUSTO")],
     "Valor = preço sem IVA. Artigo + quantidade → custo automático (custo médio). Pago/recebido: 'A crédito' ou caixa/banco."),
    ("Prestação de serviços", "Venda", "FT", "FR",
     [("MEIO_CLI", "D", "TOTAL"), ("CONTA:62.1", "C", "BASE"), ("34.5.3", "C", "IVA")],
     "Valor = preço sem IVA."),
    ("Compra de mercadorias", "Compra", "FC", "FRF",
     [("26.1", "D", "BASE"), ("34.5.2", "D", "IVA"), ("MEIO_FOR", "C", "TOTAL")],
     "Valor = custo sem IVA. Indique artigo e quantidade."),
    ("Despesa / fornecimento de serviços", "Compra", "FC", "FRF",
     [("CONTA:75.2.9", "D", "BASE"), ("34.5.2", "D", "IVA"), ("MEIO_FOR", "C", "TOTAL")],
     "Conta específica = conta de custo (ex.: 75.2.1 material, 75.2.4 electricidade)."),
    ("Aquisição de imobilizado", "Aquisição de imobilizado", "FC", "FRF",
     [("CONTA:11.5.1", "D", "BASE"), ("34.5.2", "D", "IVA"), ("MEIO_FORI", "C", "TOTAL")],
     "Conta específica = conta do activo. Registe também o bem em 10_ACTIVOS_FIXOS."),
    ("Recebimento de cliente", "Recebimento de cliente", "RC", "RC",
     [("TES", "D", "BASE"), ("CONTA:31.1.1", "C", "BASE")],
     "Valor recebido. Ref. = nº da operação da venda liquidada. Pago/recebido = caixa/banco."),
    ("Recebimento de cliente com retenção na fonte", "Recebimento de cliente", "RC", "RC",
     [("TES", "D", "LIQRET"), ("34.1.2", "D", "RET"), ("CONTA:31.1.1", "C", "BASE")],
     "Valor = total da factura liquidada (com IVA). Ref. = nº da operação da prestação de serviços. Retenção = taxa RET_SERV (11_FISCALIDADE_AGT) × valor do serviço sem IVA."),
    ("Pagamento a fornecedor", "Pagamento a fornecedor", "RCF", "RCF",
     [("CONTA:32.1.1", "D", "BASE"), ("TES", "C", "BASE")],
     "Valor pago. Ref. = nº da operação da compra liquidada."),
    ("Salários (processar e pagar)", "Pagamento ao pessoal", "FS", "FS",
     [("72.2", "D", "BASE"), ("72.5", "D", "INSS_E"), ("34.3.1", "C", "IRT"), ("34.9.1", "C", "INSS_TOT"), ("TES", "C", "LIQ")],
     "Valor = remuneração bruta. INSS e IRT calculados pela tabela oficial do exercício."),
    ("Pagamento de impostos", "Pagamento de impostos", "DLI", "DLI",
     [("CONTA:34.5.6", "D", "BASE"), ("TES", "C", "BASE")],
     "Conta específica = imposto pago (34.5.6 IVA, 34.3.1 IRT, 34.9.1 INSS, 34.1.1 II)."),
    ("Transferência entre contas", "Transferência interna", "DI", "DI",
     [("CONTA:43.1.1", "D", "BASE"), ("TES", "C", "BASE")],
     "Conta específica = conta de destino; Pago/recebido = conta de origem."),
    ("Empréstimo recebido", "Empréstimo obtido", "EXT", "EXT",
     [("TES", "D", "BASE"), ("CONTA:33.1.1", "C", "BASE")], "Conta específica = 33.1.1 (M/L prazo) ou 33.1.2 (curto prazo)."),
    ("Reembolso de empréstimo", "Reembolso de empréstimo", "EXT", "EXT",
     [("CONTA:33.1.1", "D", "BASE"), ("TES", "C", "BASE")], ""),
    ("Juros e encargos bancários", "Juros pagos", "EXT", "EXT",
     [("CONTA:76.1", "D", "BASE"), ("TES", "C", "BASE")], "Conta específica: 76.1 juros; 76.8 serviços bancários."),
    ("Entrada de capital", "Aumento de capital", "DI", "DI",
     [("TES", "D", "BASE"), ("CONTA:51.1", "C", "BASE")], ""),
    ("Depreciação do mês", "Depreciação", "DI", "DI",
     [("73.1", "D", "DEP"), ("CONTA:18.1.5", "C", "DEP")],
     "Valor vazio = depreciação do mês calculada em 10_ACTIVOS_FIXOS."),
]
VALKEYS = ["BASE", "IVA", "TOTAL", "CUSTO", "INSS_E", "INSS_TOT", "IRT", "LIQ", "DEP", "RET", "LIQRET"]

IN_COLS = [  # (chave, cabeçalho, largura, formato)
    ("Data", "Data", 11, DATE), ("Tipo", "Tipo de operação", 30, None), ("TipoDoc", "Tipo doc. (opcional)", 8, "@"),
    ("Serie", "Série", 6, "@"), ("NumDoc", "Nº documento", 9, "0"), ("NIF", "NIF cliente / fornecedor", 12, "@"),
    ("Desc", "Descrição", 34, None), ("Valor", "Valor (Kz) — sem IVA / bruto", 15, NUM), ("CodF", "Código IVA", 10, "@"),
    ("Meio", "Pago / recebido por", 14, "@"), ("ContaEsp", "Conta específica (opcional)", 10, "@"), ("Ref", "Ref. operação liquidada (nº)", 9, "0"),
    ("Venc", "Vencimento", 11, DATE), ("Artigo", "Artigo", 8, "@"), ("Qtd", "Quantidade", 9, "#,##0.00"), ("CC", "C. custo", 7, "@"),
    ("Proj", "Projecto", 8, "@"), ("Suporte", "Documento de suporte", 18, "@"), ("User", "Utilizador", 10, "@"),
    ("Validador", "Validado por", 10, "@"), ("Hash", "Hash AGT", 7, "@"), ("Cert", "Nº certificado software", 14, "@"),
    ("EstAGT", "Estado AGT", 11, "@"), ("Obs", "Observação", 28, None),
]
OC = {k: CL(2 + i) for i, (k, *_) in enumerate(IN_COLS)}   # coluna A = nº operação
CALC = ["Taxa", "IVA", "Total", "Custo", "INSS_T", "INSS_E", "INSS_TOT", "IRT", "LIQ", "DQ", "DV", "CMP", "ID", "DocEf", "Nat", "Linhas", "Estado", "DEP", "BaseRef", "RET", "LIQRET"]
for i, k in enumerate(CALC):
    OC[k] = CL(2 + len(IN_COLS) + i)
LAST_OP = OR0 + OPS - 1


def add_templates(wb):
    """Modelos de lançamento em 01A_TABELAS (parametrizáveis)."""
    ws = wb["01A_TABELAS"]
    c0 = 34
    header(ws, 5, c0, ["Chave", "Tipo de operação", "Linha", "Origem da conta", "Lado", "Montante"], [30, 30, 5, 14, 5, 9])
    r = 6
    for nome, *_rest in TIPOS:
        for s, (src, lado, val) in enumerate(_rest[3], start=1):
            put(ws, (r, c0), f'=IF({CL(c0 + 1)}{r}="","",{CL(c0 + 1)}{r}&"|"&{CL(c0 + 2)}{r})', "grey")
            put(ws, (r, c0 + 1), nome, "input")
            put(ws, (r, c0 + 2), s, "input", "0")
            put(ws, (r, c0 + 3), D.CONTA_DEFAULT.get(src, src), "input")
            put(ws, (r, c0 + 4), lado, "input")
            put(ws, (r, c0 + 5), val, "input")
            r += 1
    r1 = r + 20
    for rr in range(r, r1):
        put(ws, (rr, c0), f'=IF({CL(c0 + 1)}{rr}="","",{CL(c0 + 1)}{rr}&"|"&{CL(c0 + 2)}{rr})', "grey")
        for j in range(1, 6):
            put(ws, (rr, c0 + j), None, "input")
    for nm, j in [("MOD_Key", 0), ("MOD_Src", 3), ("MOD_Lado", 4), ("MOD_Val", 5)]:
        name(wb, nm, "01A_TABELAS", f"${CL(c0 + j)}$6:${CL(c0 + j)}${r1 - 1}")
    # tipos de operação
    t0 = c0 + 7
    header(ws, 5, t0, ["Tipo de operação", "Natureza", "Doc. a crédito", "Doc. pago/recebido", "Ajuda"], [30, 22, 8, 8, 60])
    for i, (nome, nat, dc, dp, _slots, ajuda) in enumerate(TIPOS):
        for j, v in enumerate((nome, nat, dc, dp, ajuda)):
            put(ws, (6 + i, t0 + j), v, "input")
    n1 = 5 + len(TIPOS)
    for nm, j in [("TOP_Nome", 0), ("TOP_Nat", 1), ("TOP_DocC", 2), ("TOP_DocP", 3), ("TOP_Ajuda", 4)]:
        name(wb, nm, "01A_TABELAS", f"${CL(t0 + j)}$6:${CL(t0 + j)}${n1}")
    name(wb, "L_TipoOp", "01A_TABELAS", f"${CL(t0)}$6:${CL(t0)}${n1}")
    # montantes e meios de pagamento
    v0 = t0 + 6
    header(ws, 5, v0, ["Montante (chave)", "Pago / recebido por"], [12, 16])
    for i, k in enumerate(VALKEYS):
        put(ws, (6 + i, v0), k, "grey")
    name(wb, "L_ValKeys", "01A_TABELAS", f"${CL(v0)}$6:${CL(v0)}${5 + len(VALKEYS)}")
    meios = ["A crédito"] + [c for c, _, t in D.PLANO if t == "M" and c[:2] in ("43", "45")]
    for i in range(25):
        put(ws, (6 + i, v0 + 1), meios[i] if i < len(meios) else None, "input", "@")
    name(wb, "L_Meio", "01A_TABELAS", f"${CL(v0 + 1)}$6:${CL(v0 + 1)}${5 + len(meios)}")


def build_operacoes(wb):
    ws = wb.create_sheet(S_OP)
    title(ws, "02A — OPERAÇÕES (entrada simplificada — modo aplicação)",
          "Uma linha por operação. A matriz gera sozinha os lançamentos, IVA, custo das mercadorias, INSS e IRT. Lançamentos especiais (abertura, notas de crédito, "
          "moeda estrangeira, acréscimos/diferimentos) continuam no 02_DIÁRIO (zona manual).",
          "CAMADA 1 — INPUT (azul)  |  cálculo automático (cinzento)")
    heads = ["Nº"] + [h for _, h, _, _ in IN_COLS] + ["Taxa IVA", "IVA", "Total", "Custo mercadorias", "INSS trabalhador", "INSS empresa",
                                                      "INSS total", "IRT", "Líquido a pagar", "Δ qtd stock", "Δ valor stock", "Custo médio antes",
                                                      "ID lançamento", "Tipo doc. efectivo", "Natureza", "Linhas no Diário", "Estado", "Depreciação do mês",
                                                      "Base do serviço liquidado", "Retenção na fonte", "Líquido recebido"]
    widths = [6] + [w for _, _, w, _ in IN_COLS] + [7, 12, 13, 13, 12, 12, 12, 12, 13, 9, 12, 11, 10, 8, 20, 7, 30, 12, 13, 12, 13]
    header(ws, OR0 - 1, 1, heads, widths)
    for j in range(len(IN_COLS) + 1, len(heads)):
        ws.cell(row=OR0 - 1, column=1 + j).fill = fill("595959")
    ws.freeze_panes = f"D{OR0}"
    put(ws, "A4", "Ajuda do tipo seleccionado na linha activa: ver 01A_TABELAS (coluna Ajuda). Ref. operação = nº (coluna A) da venda/compra que o recebimento/pagamento liquida.", "note")
    ops = D.OPERACOES if hasattr(D, "OPERACOES") else []
    sal = '"Salários (processar e pagar)"'
    for i in range(OPS):
        r = OR0 + i
        put(ws, (r, 1), i + 1, "grey", "0")
        rec = ops[i] if i < len(ops) else {}
        for k, _, _, fmt in IN_COLS:
            put(ws, f"{OC[k]}{r}", rec.get(k), "input", fmt)
        g = lambda k: f"{OC[k]}{r}"
        tipo = g("Tipo")
        f = {
            "Taxa": f'=IF({g("CodF")}="",0,IFERROR(N(INDEX(TX_Taxa,MATCH({g("CodF")},TX_Cod,0))),0))',
            "IVA": f"=ROUND(N({g('Valor')})*{g('Taxa')},2)",
            "Total": f"=N({g('Valor')})+{g('IVA')}",
            "Custo": f'=IF(AND({tipo}="Venda de mercadorias",{g("Artigo")}<>""),ROUND(N({g("Qtd")})*N({g("CMP")}),2),0)',
            "INSS_T": f"=IF({tipo}={sal},ROUND(N({g('Valor')})*TX_INSS_T,2),0)",
            "INSS_E": f"=IF({tipo}={sal},ROUND(N({g('Valor')})*TX_INSS_E,2),0)",
            "INSS_TOT": f"={g('INSS_T')}+{g('INSS_E')}",
            "IRT": (f"=IF({tipo}<>{sal},0,IF(N({g('Valor')})-{g('INSS_T')}<=IRT_Isencao,0,IFERROR(ROUND(INDEX(IRT_PF,MATCH(N({g('Valor')})-{g('INSS_T')},IRT_Inf,1))"
                    f"+(N({g('Valor')})-{g('INSS_T')}-INDEX(IRT_Inf,MATCH(N({g('Valor')})-{g('INSS_T')},IRT_Inf,1)))*INDEX(IRT_Tx,MATCH(N({g('Valor')})-{g('INSS_T')},IRT_Inf,1)),2),0)))"),
            "LIQ": f"=N({g('Valor')})-{g('INSS_T')}-{g('IRT')}",
            "DQ": f'=IF({g("Artigo")}="",0,IF({tipo}="Compra de mercadorias",N({g("Qtd")}),IF({tipo}="Venda de mercadorias",-N({g("Qtd")}),0)))',
            "DV": f'=IF({g("Artigo")}="",0,IF({tipo}="Compra de mercadorias",N({g("Valor")}),IF({tipo}="Venda de mercadorias",-{g("Custo")},0)))',
            "CMP": (f'=IF({g("Artigo")}="","",IFERROR((SUMIFS(MJ_DC,MJ_Artigo,{g("Artigo")},MJ_Classe,"2",MJ_Data,"<="&N({g("Data")}))'
                    f'+SUMIFS(${OC["DV"]}${OR0 - 1}:{OC["DV"]}{r - 1},${OC["Artigo"]}${OR0 - 1}:{OC["Artigo"]}{r - 1},{g("Artigo")}))'
                    f'/(SUMIFS(MJ_Qtd,MJ_Artigo,{g("Artigo")},MJ_Data,"<="&N({g("Data")}))'
                    f'+SUMIFS(${OC["DQ"]}${OR0 - 1}:{OC["DQ"]}{r - 1},${OC["Artigo"]}${OR0 - 1}:{OC["Artigo"]}{r - 1},{g("Artigo")})),0))'),
            "ID": f'=IF(OR({g("Data")}="",{tipo}=""),"",{OP_ID0}+A{r})',
            "DocEf": f'=IF({g("ID")}="","",IF({g("TipoDoc")}<>"",{g("TipoDoc")},IFERROR(IF({g("Meio")}="A crédito",INDEX(TOP_DocC,MATCH({tipo},TOP_Nome,0)),INDEX(TOP_DocP,MATCH({tipo},TOP_Nome,0))),"DI")))',
            "Nat": f'=IF({g("ID")}="","",IFERROR(INDEX(TOP_Nat,MATCH({tipo},TOP_Nome,0)),""))',
            "Linhas": f'=IF({g("ID")}="","",COUNTIF(J_ID,{g("ID")}))',
            "Estado": (f'=IF({g("ID")}="","",IF(ISNA(MATCH({tipo},TOP_Nome,0)),"🔴 Tipo de operação desconhecido",IF({g("Linhas")}<2,"🔴 Não gerou lançamento — verifique valor e meio",'
                       f'IF(COUNTIFS(J_ID,{g("ID")},J_ErrFlag,1)>0,"🔴 "&COUNTIFS(J_ID,{g("ID")},J_ErrFlag,1)&" linha(s) com erro — ver Erros_Detectados no Diário",'
                       f'IF(AND({tipo}={sal},COUNT(IRT_Inf)=0),"🟡 Tabela IRT em falta",IF({g("Validador")}="","🟡 Registada — por aprovar","🟢 Registada"))))))'),
            "BaseRef": f'=IF(N({g("Ref")})=0,0,IFERROR(N(INDEX(${OC["Valor"]}${OR0}:${OC["Valor"]}${LAST_OP},{g("Ref")})),0))',
            "RET": f'=IF({tipo}<>"Recebimento de cliente com retenção na fonte",0,ROUND({g("BaseRef")}*TX_RET,2))',
            "LIQRET": f"=N({g('Valor')})-{g('RET')}",
            "DEP": f'=IF({tipo}<>"Depreciação do mês",0,IF(N({g("Valor")})<>0,N({g("Valor")}),IFERROR(INDEX(AF_DepMesReg,MONTH({g("Data")})),0)))',
        }
        for k in CALC:
            put(ws, f"{OC[k]}{r}", f[k], "grey", {"Taxa": PCT, "CMP": NUM, "ID": "0", "Linhas": "0", "DQ": "#,##0.00"}.get(k, None if k in ("DocEf", "Nat", "Estado") else NUM))
    rng = lambda k: f"{OC[k]}{OR0}:{OC[k]}{LAST_OP}"
    dv_list(ws, rng("Tipo"), "=L_TipoOp")
    dv_list(ws, rng("CodF"), "=TX_Cod")
    dv_list(ws, rng("Meio"), "=L_Meio")
    dv_list(ws, rng("ContaEsp"), "=PC_Cod")
    dv_list(ws, rng("NIF"), "=T_NIF")
    dv_list(ws, rng("TipoDoc"), "=TD_Cod")
    dv_list(ws, rng("CC"), "=L_CC")
    dv_list(ws, rng("EstAGT"), "=L_EstAGT")
    status_cf(ws, rng("Estado"))
    put(ws, "D3", "Operações:", "label")
    put(ws, "E3", f'=COUNT({rng("ID")})', "grey", "0")
    put(ws, "F3", "Com erro:", "label")
    put(ws, "G3", f'=COUNTIF({rng("Estado")},"🔴*")', "grey", "0")
    name(wb, "OP_Erros", S_OP, f"$G$3")
    protect(ws)


def build_motor(wb):
    ws = wb.create_sheet(S_MOT)
    title(ws, "02B — MOTOR DE LANÇAMENTOS (gerado — não editar)",
          f"Cada operação de 02A gera até {SLOTS} linhas de partidas dobradas segundo os modelos de 01A_TABELAS. Estas linhas alimentam a zona automática do 02_DIÁRIO.",
          "CAMADA 2 — PROCESSAMENTO")
    import sheets_base as SB
    aux = ["Op", "Linha", "Tipo", "Chave", "Origem", "Lado", "Montante (chave)", "Montante", "Meio", "Conta", "Activa"]
    heads = aux + [h for _, h, _, _ in SB.J_IN]
    header(ws, 5, 1, heads, [5, 5, 26, 30, 12, 5, 9, 13, 9, 9, 5] + [10] * len(SB.J_IN))
    ws.freeze_panes = "C6"
    A = {k: CL(1 + i) for i, k in enumerate(["Op", "Slot", "Tipo", "Key", "Src", "Lado", "VK", "Mont", "Meio", "Conta", "Act"])}
    JC = {k: CL(len(aux) + 1 + i) for i, (k, *_) in enumerate(SB.J_IN)}
    q_ = q(S_OP)
    for i in range(OPS):
        orow = OR0 + i
        o = lambda k: f"{q_}!{OC[k]}{orow}"
        for s in range(1, SLOTS + 1):
            r = 6 + i * SLOTS + (s - 1)
            ws[f"A{r}"] = i + 1
            ws[f"B{r}"] = s
            f = {
                "Tipo": f'={o("Tipo")}&""',
                "Key": f'=C{r}&"|"&B{r}',
                "Src": f'=IF(C{r}="","",IFERROR(INDEX(MOD_Src,MATCH(D{r},MOD_Key,0))&"",""))',
                "Lado": f'=IF(E{r}="","",IFERROR(INDEX(MOD_Lado,MATCH(D{r},MOD_Key,0))&"",""))',
                "VK": f'=IF(E{r}="","",IFERROR(INDEX(MOD_Val,MATCH(D{r},MOD_Key,0))&"",""))',
                "Mont": (f'=IF(E{r}="",0,IFERROR(CHOOSE(MATCH(G{r},L_ValKeys,0),N({o("Valor")}),{o("IVA")},{o("Total")},{o("Custo")},{o("INSS_E")},'
                         f'{o("INSS_TOT")},{o("IRT")},{o("LIQ")},{o("DEP")},{o("RET")},{o("LIQRET")}),0))'),
                "Meio": f'={o("Meio")}&""',
                "Conta": (f'=IF(E{r}="","",IF(E{r}="MEIO_CLI",IF(I{r}="A crédito","31.1.1",I{r}),IF(E{r}="MEIO_FOR",IF(I{r}="A crédito","32.1.1",I{r}),'
                          f'IF(E{r}="MEIO_FORI",IF(I{r}="A crédito","32.2",I{r}),IF(E{r}="TES",I{r},IF(LEFT(E{r},6)="CONTA:",IF({o("ContaEsp")}<>"",{o("ContaEsp")}&"",MID(E{r},7,20)),E{r}))))))'),
                "Act": f'=IF(AND({o("ID")}<>"",E{r}<>"",ABS(H{r})>0.005),1,0)',
            }
            for k, v in f.items():
                ws[f"{A[k]}{r}"] = v
            idc = f"${JC['ID']}{r}"
            iva_line = f'OR(J{r}="34.5.2",J{r}="34.5.3")'
            base_isenta = f'AND(G{r}="BASE",{o("Taxa")}=0,{o("CodF")}<>"",LEFT(J{r},1)="6")'
            jf = {
                "ID": f'=IF(K{r}=1,{o("ID")},"")',
                "Data": f'=IF({idc}="","",{o("Data")})',
                "TipoDoc": f'=IF({idc}="","",{o("DocEf")})',
                "Serie": f'=IF({idc}="","",{o("Serie")}&"")',
                "NumDoc": f'=IF(OR({idc}="",{o("NumDoc")}=""),"",{o("NumDoc")})',
                "DataDoc": f'=IF({idc}="","",{o("Data")})',
                "Venc": f'=IF(OR({idc}="",{o("Venc")}=""),"",{o("Venc")})',
                "NIF": f'=IF({idc}="","",{o("NIF")}&"")',
                "Conta": f'=IF({idc}="","",J{r})',
                "Desc": f'=IF({idc}="","",{o("Desc")}&"")',
                "Deb": f'=IF(AND({idc}<>"",F{r}="D"),H{r},"")',
                "Cred": f'=IF(AND({idc}<>"",F{r}="C"),H{r},"")',
                "CC": f'=IF({idc}="","",{o("CC")}&"")',
                "CR": '=""', "Fonte": '=""', "Moeda": '=""', "Cambio": '=""', "ValME": '=""', "DtAlt": '=""', "DtCom": '=""', "ErroCom": '=""',
                "Proj": f'=IF({idc}="","",{o("Proj")}&"")',
                "Nat": f'=IF({idc}="","",{o("Nat")})',
                "CodF": f'=IF({idc}="","",IF(G{r}="RET","RET_SERV",IF(OR({iva_line},{base_isenta}),{o("CodF")}&"","")))',
                "Base": f'=IF(OR({idc}="",{JC["CodF"]}{r}=""),"",IF(G{r}="RET",{o("BaseRef")},N({o("Valor")})))',
                "Ref": f'=IF(OR({idc}="",{o("Ref")}=""),"",IF(OR(LEFT(J{r},2)="31",LEFT(J{r},2)="32"),{OP_ID0}+{o("Ref")},""))',
                "FormaPag": f'=IF({idc}="","",IF(LEFT(J{r},2)="45","Numerário",IF(LEFT(J{r},2)="43","Transferência","")))',
                "Artigo": f'=IF(OR({idc}="",{o("Artigo")}=""),"",IF(LEFT(J{r},1)="2",{o("Artigo")}&"",""))',
                "Qtd": f'=IF({JC["Artigo"]}{r}="","",{o("DQ")})',
                "Suporte": f'=IF({idc}="","",{o("Suporte")}&"")',
                "User": f'=IF({idc}="","",{o("User")}&"")',
                "Validador": f'=IF({idc}="","",{o("Validador")}&"")',
                "DtIns": f'=IF({idc}="","",{o("Data")})',
                "EstDoc": f'=IF(OR({idc}="",B{r}<>1),"","Normal")',
                "Hash": f'=IF(OR({idc}="",B{r}<>1),"",{o("Hash")}&"")',
                "Cert": f'=IF(OR({idc}="",B{r}<>1),"",{o("Cert")}&"")',
                "EstAGT": f'=IF(OR({idc}="",B{r}<>1),"",{o("EstAGT")}&"")',
                "Obs": f'=IF({idc}="","","Gerado da operação nº "&A{r})',
            }
            for k, *_ in SB.J_IN:
                ws[f"{JC[k]}{r}"] = jf[k]
    protect(ws)
    return JC


def motor_ref(k, row_offset):
    """Referência para a célula do motor que alimenta a linha automática do Diário."""
    import sheets_base as SB
    col = CL(11 + 1 + [x[0] for x in SB.J_IN].index(k))
    return f"{q(S_MOT)}!{col}{6 + row_offset}"
