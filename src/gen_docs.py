"""Gera a documentação derivada do próprio modelo (sem duplicar manualmente):
docs/04_DICIONARIO_DADOS.md, docs/05_MAPA_DEPENDENCIAS.md, docs/06_MATRIZ_LEGISLACAO.md, docs/08_CASOS_TESTE.md

Uso: python src/gen_docs.py [workbook.xlsx]
"""
import os
import re
import sys
from collections import defaultdict, OrderedDict

sys.path.insert(0, os.path.dirname(__file__))
import data as D  # noqa: E402
import sheets_base as SB  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")

DESC = {
    "ID": ("Inteiro", "Identificador único do lançamento. Todas as linhas da mesma operação partilham o ID.", "Obrigatório; Σ Débito = Σ Crédito por ID."),
    "Data": ("Data", "Data contabilística do lançamento.", "Data válida, dentro do exercício (00_CONFIGURAÇÃO)."),
    "TipoDoc": ("Texto (lista)", "Tipo de documento (FT, FR, NC, ND, RC, FC, FRF, NCF, RCF, DU, EXT, FS, DI, DLI, ABE).", "Lista 01A_TABELAS."),
    "Serie": ("Texto", "Série do documento.", "—"),
    "NumDoc": ("Inteiro", "Número do documento na série.", "Duplicados detectados (tipo+série+nº+NIF); sequência validada em 13_FACTURAÇÃO_FISCAL."),
    "DataDoc": ("Data", "Data do documento de suporte.", "—"),
    "Venc": ("Data", "Data de vencimento (conta corrente).", "Usada no aging e na tesouraria de 13 semanas."),
    "NIF": ("Texto (lista)", "NIF do terceiro (clientes, fornecedores).", "Obrigatório em contas correntes (31/32); formato 10 dígitos (colectivas) ou 14 caracteres (BI); deve existir em 01B_TERCEIROS."),
    "Conta": ("Texto (lista)", "Conta do PGC (apenas contas de Movimento).", "Existe no plano; não pode ser conta de agregação."),
    "Desc": ("Texto", "Descrição da linha.", "—"),
    "Deb": ("Número", "Valor a débito (Kz).", "≥ 0; exclusivo com Crédito."),
    "Cred": ("Número", "Valor a crédito (Kz).", "≥ 0; exclusivo com Débito."),
    "CC": ("Texto (lista)", "Centro de custo.", "Lista 01A_TABELAS."),
    "CR": ("Texto", "Centro de responsabilidade.", "—"),
    "Proj": ("Texto", "Código do projecto (liga a 24_GESTÃO_DE_PROJECTOS).", "—"),
    "Fonte": ("Texto", "Fonte de financiamento.", "—"),
    "Nat": ("Texto (lista)", "Natureza da operação — determina a classificação na DFC e a inclusão na DR.", "Obrigatória; parametrizada em 01A_TABELAS."),
    "CodF": ("Texto (lista)", "Código fiscal (IVA_GER, IVA_ISE, RET_SERV, …).", "Obrigatório nas linhas de IVA; taxa × base = valor (tolerância CFG_Tol)."),
    "Base": ("Número", "Base tributável associada ao código fiscal.", "Usada no mapa de IVA e retenções."),
    "Ref": ("Inteiro", "ID do documento de origem liquidado/regularizado (recibos, pagamentos, NC).", "Tem de existir no Diário."),
    "Moeda": ("Texto (lista)", "Moeda da operação (vazio = AOA).", "Se ≠ AOA: câmbio e valor em ME obrigatórios."),
    "Cambio": ("Número", "Taxa de câmbio aplicada.", "Valor ME × câmbio = valor em Kz."),
    "ValME": ("Número", "Valor na moeda de origem.", "—"),
    "FormaPag": ("Texto (lista)", "Forma de pagamento/recebimento.", "—"),
    "Artigo": ("Texto", "Código do artigo (linhas da classe 2).", "Alimenta 09_INVENTÁRIOS e o kardex."),
    "Qtd": ("Número", "Quantidade (+ entrada, − saída).", "—"),
    "Suporte": ("Texto", "Referência do documento de suporte / arquivo.", "Obrigatório."),
    "User": ("Texto", "Utilizador que registou.", "Rastreabilidade."),
    "Validador": ("Texto", "Utilizador que validou.", "Diferente do Utilizador (segregação de funções)."),
    "DtIns": ("Data", "Data de inserção.", "Comparada com a data de fecho do mês."),
    "DtAlt": ("Data", "Data da última alteração.", "Alteração após fecho = erro."),
    "EstDoc": ("Texto (lista)", "Normal / Anulado / Rectificado.", "—"),
    "Hash": ("Texto", "Hash/assinatura do documento emitido (software validado).", "Obrigatório em documentos emitidos."),
    "Cert": ("Texto", "Nº do certificado/validação do software de facturação.", "Obrigatório em documentos emitidos."),
    "EstAGT": ("Texto (lista)", "Estado da comunicação electrónica à AGT.", "Erro/Pendente geram alertas."),
    "DtCom": ("Data", "Data de comunicação à AGT.", "—"),
    "ErroCom": ("Texto", "Mensagem de erro da comunicação.", "—"),
    "Obs": ("Texto", "Observação.", "—"),
}


def dicionario():
    L = ["# 04 — Dicionário de Dados", "", "Gerado por `src/gen_docs.py` a partir de `src/sheets_base.py` (fonte única).", "",
         "## 02_DIÁRIO_LANÇAMENTOS — campos de INPUT", "", "| Coluna | Campo | Tipo | Descrição | Validação |", "|---|---|---|---|---|"]
    for k, h, *_ in SB.J_IN:
        t, d, v = DESC.get(k, ("", "", ""))
        L.append(f"| {SB.JCOLS[k]} | {h} | {t} | {d} | {v} |")
    L += ["", "## 02_DIÁRIO_LANÇAMENTOS — campos CALCULADOS (camada de processamento)", "",
          "| Coluna | Campo | Nome definido | Fórmula (linha 6) |", "|---|---|---|---|"]
    from core import render
    for k, h, _, _, tpl in SB.J_CALC:
        f = render(tpl, SB.JCOLS, 6).replace("|", "¦")
        if len(f) > 180:
            f = f[:177] + "…"
        L.append(f"| {SB.JCOLS[k]} | {h} | `J_{k}` | `{f}` |")
    L += ["", "## 01_PLANO_CONTAS", "", "| Campo | Descrição |", "|---|---|"]
    for a, b in [("Código", "Código PGC (texto, com pontos)."), ("Tipo", "Movimento (recebe lançamentos) ou Agregação (soma as filhas)."),
                 ("Rubrica balanço (saldo devedor/credor)", "Rubrica do Balanço consoante o sinal do saldo (ex.: 43 devedor → Disponibilidades; credor → Empréstimos de curto prazo)."),
                 ("Rubrica DR", "Linha da Demonstração de Resultados (classes 6, 7 e 87)."), ("Conta corrente", "S = exige NIF (31, 32)."),
                 ("Reconciliação", "S = conta sujeita a reconciliação/alerta de ausência de movimento."), ("Saldos", "Débito, crédito, saldo, abertura — calculados.")]:
        L.append(f"| {a} | {b} |")
    L += ["", "## Nomes definidos (interface entre folhas)", "", "Os mapas comunicam por **nomes definidos** (ex.: `DR_RL`, `BS_AT`, `IVA_Estado`). A lista completa com a folha de origem está em `05_MAPA_DEPENDENCIAS.md`."]
    open(os.path.join(DOCS, "04_DICIONARIO_DADOS.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


def dependencias(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    name_sheet = {}
    for nm, dn in wb.defined_names.items():
        sh = dn.attr_text.split("!")[0].strip("'")
        name_sheet[nm] = sh
    tok = re.compile(r"[A-Za-z_][A-Za-z0-9_\.]*")
    deps = defaultdict(set)
    name_users = defaultdict(set)
    for ws in wb.worksheets:
        seen = set()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    for t in set(tok.findall(v)):
                        if t in name_sheet and t not in seen:
                            seen.add(t)
                    for m in re.findall(r"'([^']+)'!", v):
                        if m != ws.title:
                            deps[ws.title].add(m)
        for t in seen:
            src = name_sheet[t]
            name_users[t].add(ws.title)
            if src != ws.title:
                deps[ws.title].add(src)
    order = wb.sheetnames
    L = ["# 05 — Mapa de Dependências", "", "Extraído automaticamente das fórmulas do ficheiro gerado (`src/gen_docs.py`).", "",
         "## Fluxo de dados (visão de alto nível)", "", "```mermaid", "flowchart LR"]
    groups = OrderedDict([("INPUT", ["00_CONFIGURAÇÃO", "01_PLANO_CONTAS", "01A_TABELAS", "01B_TERCEIROS", "02_DIÁRIO_LANÇAMENTOS", "16_ORÇAMENTO_EMPRESARIAL", "11_FISCALIDADE_AGT"]),
                          ("LIVROS", ["03_DIÁRIO_GERAL", "04_RAZÃO", "05_CAIXA", "06_BANCOS", "07_CLIENTES", "08_FORNECEDORES", "09_INVENTÁRIOS", "10_ACTIVOS_FIXOS"]),
                          ("FISCAL", ["12_IVA", "13_FACTURAÇÃO_FISCAL", "14_SAFT", "15_IMPOSTO_INDUSTRIAL", "34_CALENDÁRIO_FISCAL_AGT"]),
                          ("DEMONSTRAÇÕES", ["20_BALANCETE", "21_BALANÇO", "22_DRE", "19_FLUXO_DE_CAIXA", "23_DFC"]),
                          ("GESTÃO", ["17_PLANEAMENTO_FINANCEIRO", "18_TESOURARIA", "24_GESTÃO_DE_PROJECTOS", "25_ANÁLISE_DE_INVESTIMENTOS", "26_RISCO_FINANCEIRO",
                                      "27_SUSTENTABILIDADE", "28_KPI_FINANCEIROS", "30_BUDGET_VS_ACTUAL", "36_BREAK_EVEN"]),
                          ("CONTROLO", ["31_CONTROLO_INTERNO", "32_AUDITORIA", "33_FECHO_MENSAL", "37_ENCERRAMENTO_EXERCÍCIO", "39_ALERTAS", "99_CONTROLO_SISTEMA"]),
                          ("REPORTE", ["29_DASHBOARD_EXECUTIVO", "38_RELATÓRIO_GESTÃO"])])
    for g, sheets in groups.items():
        L.append(f'  subgraph {g}')
        for s in sheets:
            L.append(f'    S_{re.sub(r"[^A-Za-z0-9]", "_", s)}["{s}"]')
        L.append("  end")
    L += ["  INPUT --> LIVROS --> DEMONSTRAÇÕES --> GESTÃO --> REPORTE", "  INPUT --> FISCAL --> DEMONSTRAÇÕES", "  LIVROS --> CONTROLO", "  FISCAL --> CONTROLO",
          "  DEMONSTRAÇÕES --> CONTROLO --> REPORTE", "```", "", "## Dependências por folha (folhas de que cada folha lê dados)", "", "| Folha | Lê de |", "|---|---|"]
    for s in order:
        L.append(f"| {s} | {', '.join(sorted(deps.get(s, []), key=order.index)) or '—'} |")
    L += ["", "## Nomes definidos: origem e utilizadores", "", "| Nome | Folha de origem | Usado em |", "|---|---|---|"]
    for nm in sorted(name_sheet):
        users = sorted((name_users.get(nm, set()) - {name_sheet[nm]}), key=order.index)
        L.append(f"| `{nm}` | {name_sheet[nm]} | {', '.join(users) if users else '(interno)'} |")
    open(os.path.join(DOCS, "05_MAPA_DEPENDENCIAS.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


def legislacao():
    L = ["# 06 — Matriz de Legislação", "",
         "> Pesquisa efectuada em 24/09/2026 com base em fontes secundárias (MINFIN, consultoras, imprensa especializada). **Nenhuma regra deve ser usada oficialmente sem confirmação no Diário da República.**",
         "", "## Diplomas", "", "| Diploma | Número | Data | Artigo | Matéria | Regra | Aplicabilidade | Vigor | Revogação | Fonte | Estado |", "|---|---|---|---|---|---|---|---|---|---|---|"]
    for b in D.BASE_LEGAL:
        L.append("| " + " | ".join(str(x).replace("|", "/") for i, x in enumerate(b) if i != 10) + " |")
    L += ["", "## Parametrização fiscal (11_FISCALIDADE_AGT)", "", "| Código | Imposto | Descrição | Taxa | Natureza | Diploma | Artigo | Vigência | Estado | Observação |", "|---|---|---|---|---|---|---|---|---|---|"]
    for t in D.TAXAS:
        taxa = f"{t[3]:.1%}" if isinstance(t[3], float) else "— (não parametrizada)"
        L.append(f"| {t[0]} | {t[1]} | {t[2]} | {taxa} | {t[4]} | {t[5]} | {t[6]} | {t[7] or '—'} | {t[10]} | {t[11]} |")
    L += ["", "## Pontos em aberto (validação obrigatória)", "",
          "1. **Tabela do IRT 2026 (Grupo A)** — confirmado: isenção até 150 000 Kz, 12 escalões, taxas 13%–25% (Lei n.º 14/25). Localização: art. 21.º, n.º 3 e **Anexo I** da Lei n.º 14/25. Em aberto: limites e parcelas fixas de cada escalão — copiar do Anexo I para o separador IRT_ESCALOES do modelo.",
          "2. **Número e data da Lei do OGE 2026** — referida como Lei n.º 14/25, de 30 de Dezembro, em fontes secundárias; confirmar.",
          "3. **Retenção na fonte de 6,5% sobre serviços (Imposto Industrial)** — usada nos testes; confirmar artigo, incidência e dispensas.",
          "4. **Contas de IVA no PGC (34.5.x)** — codificação analítica usada é a sugerida; confirmar com o instrutivo/decreto executivo que as criou.",
          "5. **Prazos declarativos (IVA, IRT, INSS, retenções, Imposto de Selo, Modelo 1)** — parametrizados como 'último dia do mês seguinte' e '31/05'; validar cada um.",
          "6. **Taxas de amortização fiscalmente aceites** — diploma a identificar; vida útil actual = política contabilística.",
          "7. **Exigibilidade do IVA em adiantamentos** e **autoliquidação em serviços de não residentes** — confirmar artigos do CIVA.",
          "8. **Estrutura XSD do SAF-T (AO)** vigente e requisitos de comunicação electrónica (DP 71/25 e regulamentação)."]
    open(os.path.join(DOCS, "06_MATRIZ_LEGISLACAO.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


REFLEXOS = {
    "TESTE 00": ("Abertura", "Saldos iniciais em 01_PLANO_CONTAS (abertura), 21_BALANÇO (coluna abertura), 05_CAIXA, 19/23 (saldo inicial), kardex A001 (1 000 un)."),
    "TESTE 01": ("Venda a dinheiro (FR A/1)", "05_CAIXA (+114 000), 22_DRE Jan (vendas 100 000; CMV −60 000), 12_IVA Jan (liquidado 14 000), 13_FACTURAÇÃO (FR A/1), 19/23 (recebimentos de clientes), 09_INVENTÁRIOS (−60 un)."),
    "TESTE 02": ("Venda a crédito (FT A/1)", "07_CLIENTES (doc ID 3), 22_DRE, 12_IVA, 13_FACTURAÇÃO, aging, KPI PMR."),
    "TESTE 03": ("Compra a crédito (FC B/889)", "08_FORNECEDORES (doc ID 4), 12_IVA dedutível 77 000, 09_INVENTÁRIOS (+500 un a 1 100), CMP móvel."),
    "TESTE 04": ("Compra a pronto (FRF G/1502)", "05_CAIXA (−57 000), 22_DRE FSE, 12_IVA dedutível 7 000, DFC pagamentos a fornecedores."),
    "TESTE 05": ("Pagamento a fornecedor", "08_FORNECEDORES (doc 4 em aberto 227 000), 06_BANCOS, 19/23."),
    "TESTE 06": ("Recebimento de cliente (parcial + com retenção)", "07_CLIENTES (doc 3: 213 000; doc 12: 0), 11_FISCALIDADE retenções 195 000, 15_IMPOSTO_INDUSTRIAL (dedução à colecta), 19/23."),
    "TESTE 07": ("Salários (processamento + pagamento)", "22_DRE pessoal 432 000, 11_FISCALIDADE (IRT/INSS), 34_CALENDÁRIO, DFC pagamentos ao pessoal."),
    "TESTE 08": ("Aquisição de activo (AF001)", "10_ACTIVOS_FIXOS, 21_BALANÇO imobilizado, 23_DFC investimento, 24_PROJECTOS (PRJ01), 12_IVA dedutível 168 000."),
    "TESTE 09": ("Depreciação (Fev + Mar)", "10_ACTIVOS (estado 'Processada' por mês), 22_DRE amortizações 50 000, 33_FECHO_MENSAL, EBITDA."),
    "TESTE 10": ("Empréstimo bancário", "21_BALANÇO passivo não corrente 3 000 000, 23_DFC financiamento, KPI Dívida/EBITDA, 26_RISCO."),
    "TESTE 11": ("Pagamento de juros", "22_DRE custos financeiros, 23_DFC financiamento (juros pagos), cobertura de juros."),
    "TESTE 12": ("Operação sujeita a IVA (serviços FT A/2)", "12_IVA liquidado 420 000, 22_DRE serviços 3 000 000, 24_PROJECTOS PRJ02."),
    "TESTE 13": ("Operação isenta (FT A/3)", "12_IVA base isenta 200 000 (IVA_ISE), checklist AGT (motivo de isenção), comunicação pendente → 🟡."),
    "TESTE 14": ("Nota de crédito (NC A/1)", "12_IVA regularização a favor 7 000, 22_DRE vendas −50 000, 07_CLIENTES (reduz doc 3), 13_FACTURAÇÃO."),
    "TESTE 15": ("Operação cambial (USD)", "Validação câmbio × valor ME, 08_FORNECEDORES estrangeiro, diferença de câmbio 20 000 em 22_DRE, 39_ALERTAS risco cambial."),
    "TESTE 16": ("Investimento financeiro (depósito a prazo)", "21_BALANÇO aplicações 1 000 000, 23_DFC investimento (não é equivalente de caixa)."),
    "TESTE 17": ("Financiamento (empréstimo + aumento de capital)", "21_BALANÇO capital 7 000 000, 23_DFC financiamento 4 955 000."),
    "TESTE 18": ("Despesa antecipada (seguro anual)", "37.4.1 diferimento 1 100 000 no activo, 22_DRE seguros 100 000 (1/12)."),
    "TESTE 19": ("Receita antecipada (adiantamento FR A/2)", "37.6.1 proveito diferido 600 000, IVA liquidado 84 000, 13_FACTURAÇÃO erro de comunicação → 🔴 (detecção)."),
    "TESTE 20": ("Operação de encerramento", "37_ENCERRAMENTO_EXERCÍCIO: apuramento proposto equilibrado; resultado apurado = RL da DR; abertura N+1 equilibrada."),
}


def casos():
    L = ["# 08 — Casos de Teste (teste de integridade)", "",
         "Os lançamentos de teste estão em `src/data.py` (lista `JOURNAL`) e são carregados no Diário da matriz entregue. Valores fictícios; entidade fictícia.",
         "Resultados automáticos em `09_RELATORIO_VALIDACAO.md`.", "",
         "| Teste | Operação | IDs | Reflexos verificados |", "|---|---|---|---|"]
    ids = defaultdict(set)
    for l in D.JOURNAL:
        for m in re.findall(r"TESTE \d\d", l["Obs"] or ""):
            ids[m].add(l["ID"])
    ids["TESTE 20"] = {"(proposta 37)"}
    for k in sorted(REFLEXOS):
        op, rf = REFLEXOS[k]
        L.append(f"| {k} | {op} | {', '.join(str(x) for x in sorted(ids[k], key=str))} | {rf} |")
    L += ["", "## Lançamentos adicionais do cenário", "", "| ID | Operação |", "|---|---|",
          "| 26 | Acréscimo de custos (electricidade Março) |", "| 27 | Estimativa de Imposto Industrial do 1.º trimestre |",
          "| 28, 30, 32 | Apuramento mensal do IVA (34.5.2/34.5.3/34.5.4 → 34.5.6) |", "| 29, 31 | Pagamentos ao Estado (IRT/INSS de Janeiro; IVA de Fevereiro) |", "",
          "## Testes negativos (injecção de erros)", "",
          "| Erro injectado | Resultado esperado |", "|---|---|",
          "| Lançamento desequilibrado | 'ERRO — LANÇAMENTO NÃO EQUILIBRADO' + SISTEMA BLOQUEADO |", "| Documento duplicado | 'Documento duplicado' + bloqueio |",
          "| Conta inexistente | 'Conta inexistente' + bloqueio |", "| Conta de agregação | 'Conta não movimentável' + bloqueio |",
          "| Alteração em mês encerrado | 'Alteração em período encerrado' + bloqueio |", "| Taxa de IVA incompatível | 'Taxa fiscal incompatível' + bloqueio |",
          "| Cliente sem NIF | 'NIF inválido/não registado' + bloqueio |", "| Sem documento de suporte | 'Sem documento de suporte' + bloqueio |",
          "| Fora do exercício | 'Fora do exercício' + bloqueio |", "| Câmbio inconsistente | 'Câmbio inconsistente' + bloqueio |",
          "", "## Valores de referência do cenário (31/03/2026)", "",
          "| Indicador | Valor (Kz) |", "|---|---|", "| Receita (vendas + serviços) | 3 750 000,00 |", "| Resultado operacional | 1 673 614,00 |",
          "| Resultado antes de impostos | 1 608 614,00 |", "| Imposto Industrial estimado (25%) | 402 153,50 |", "| Resultado líquido | 1 206 460,50 |",
          "| Activo total = CP + Passivo | 13 799 614,00 |", "| Capital próprio | 9 406 460,50 |", "| IVA apurado Jan–Mar | 329 000,00 (pago 245 000; a pagar 84 000) |",
          "| Meios monetários | 8 856 000,00 |", "| Clientes em aberto | 413 000,00 |", "| Fornecedores em aberto | 227 000,00 |", "| Existências (A001: 1 040 un) | 1 085 614,00 |"]
    open(os.path.join(DOCS, "08_CASOS_TESTE.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "dist", "MATRIZ_PRO_MASTER_CONTABILIDADE_ANGOLA.xlsx")
    os.makedirs(DOCS, exist_ok=True)
    dicionario()
    dependencias(xlsx)
    legislacao()
    casos()
    print("docs gerados")
