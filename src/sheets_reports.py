"""Demonstrações financeiras, orçamento, tesouraria e fluxos de caixa."""
from datetime import date
from openpyxl.utils import get_column_letter as CL
from core import *
from sheets_ledgers import MESES
import data as D

S16, S18, S19, S20, S21, S22, S23, S30 = ("16_ORÇAMENTO_EMPRESARIAL", "18_TESOURARIA", "19_FLUXO_DE_CAIXA", "20_BALANCETE", "21_BALANÇO",
                                          "22_DRE", "23_DFC", "30_BUDGET_VS_ACTUAL")
YTD = 'J_Ano,CFG_Ano,J_Mes,"<="&CFG_MesRep'


# ---------------------------------------------------------------------------
def build_balancete(wb):
    ws = wb.create_sheet(S20)
    title(ws, "20 — BALANCETE (mensal / trimestral / semestral / anual / acumulado)", "Linhas espelham o Plano de Contas. Saldo inicial inclui abertura e meses anteriores ao período.", "CAMADA 3 — OUTPUT")
    put(ws, "B4", "Tipo de período", "label")
    put(ws, "C4", "Acumulado", "input")
    dv_list(ws, "C4", '"Mensal,Trimestral,Semestral,Anual,Acumulado"')
    put(ws, "B5", "Nº do período (mês 1–12, trimestre 1–4, semestre 1–2)", "label")
    put(ws, "C5", 3, "input", "0")
    put(ws, "E4", "Mês inicial", "label")
    put(ws, "F4", '=IF(C4="Mensal",C5,IF(C4="Trimestral",(C5-1)*3+1,IF(C4="Semestral",(C5-1)*6+1,1)))', "grey", "0")
    put(ws, "E5", "Mês final", "label")
    put(ws, "F5", '=IF(C4="Mensal",C5,IF(C4="Trimestral",C5*3,IF(C4="Semestral",C5*6,IF(C4="Anual",12,CFG_MesRep))))', "grey", "0")
    name(wb, "BT_Ini", S20, "$F$4")
    name(wb, "BT_Fim", S20, "$F$5")
    heads = ["Código", "Designação", "Tipo", "Nível", "Saldo inicial", "Débito", "Crédito", "Saldo final", "Saldo devedor", "Saldo credor",
             "aux SI", "aux D", "aux C"]
    header(ws, 9, 1, heads, [9, 44, 11, 5, 15, 15, 15, 15, 15, 15, 12, 12, 12])
    ws.freeze_panes = "C10"
    r0, r1 = 10, 409
    per = 'J_Ano,CFG_Ano,J_Mes,">="&BT_Ini,J_Mes,"<="&BT_Fim,J_IsAbe,0'
    for i in range(r1 - r0 + 1):
        r = r0 + i
        put(ws, f"A{r}", f'=IF(INDEX(PC_Cod,{i + 1})="","",INDEX(PC_Cod,{i + 1}))', "calc")
        put(ws, f"B{r}", f'=IF(A{r}="","",INDEX(PC_Desig,{i + 1}))', "calc")
        put(ws, f"C{r}", f'=IF(A{r}="","",INDEX(PC_Tipo,{i + 1}))', "calc")
        put(ws, f"D{r}", f'=IF(A{r}="","",INDEX(PC_Nivel,{i + 1}))', "calc", "0")
        mov = f'OR(A{r}="",C{r}<>"Movimento")'
        put(ws, f"K{r}", f'=IF({mov},0,SUMIFS(J_DC,J_Conta,A{r},J_Ano,CFG_Ano,J_Mes,"<"&BT_Ini,J_IsAbe,0)+SUMIFS(J_DC,J_Conta,A{r},J_Ano,CFG_Ano,J_IsAbe,1))', "grey", NUM)
        put(ws, f"L{r}", f'=IF({mov},0,SUMIFS(J_Deb,J_Conta,A{r},{per}))', "grey", NUM)
        put(ws, f"M{r}", f'=IF({mov},0,SUMIFS(J_Cred,J_Conta,A{r},{per}))', "grey", NUM)
        for c, h in zip("EFG", "KLM"):
            put(ws, f"{c}{r}", f'=IF(A{r}="","",IF(C{r}="Movimento",{h}{r},IF(D{r}=1,SUMIFS(${h}${r0}:${h}${r1},$A${r0}:$A${r1},"?*",$C${r0}:$C${r1},"Movimento",$A${r0}:$A${r1},A{r}&"*"),'
                                  f'SUMIFS(${h}${r0}:${h}${r1},$A${r0}:$A${r1},A{r}&".*",$C${r0}:$C${r1},"Movimento"))))', "calc", NUM)
        put(ws, f"H{r}", f'=IF(A{r}="","",E{r}+F{r}-G{r})', "calc", NUM, bold=False)
        put(ws, f"I{r}", f'=IF(A{r}="","",MAX(0,H{r}))', "calc", NUM)
        put(ws, f"J{r}", f'=IF(A{r}="","",MAX(0,-H{r}))', "calc", NUM)
    value_cf(ws, f"A{r0}:J{r1}", f'$D{r0}<=2', ("DDEBF7", "1F3864"))
    put(ws, "B7", "TOTAIS (contas de movimento)", "label", bold=True)
    for c in "EFGHIJ":
        put(ws, f"{c}7", f'=SUMIFS({c}{r0}:{c}{r1},$C${r0}:$C${r1},"Movimento")', "grey", NUM, bold=True)
    put(ws, "B8", "Validação", "label", bold=True)
    put(ws, "F8", '=IF(ABS(F7-G7)<0.005,"🟢 Débitos = Créditos","🔴 ERRO — Débitos ≠ Créditos")', "grey")
    put(ws, "H8", '=IF(AND(ABS(E7)<0.005,ABS(H7)<0.005,ABS(I7-J7)<0.005),"🟢 Balancete equilibrado","🔴 Balancete desequilibrado")', "grey")
    put(ws, "I8", '=SUMIFS(J_Deb,' + per + ')', "grey", NUM)
    put(ws, "J8", '=IF(ABS(I8-F7)<0.005,"🟢 Diário = Balancete","🔴 Diário ≠ Balancete (conta inexistente?)")', "grey")
    status_cf(ws, "F8:J8")
    name(wb, "BT_Estado", S20, "$H$8")
    name(wb, "BT_DiarioEst", S20, "$J$8")
    name(wb, "BT_TotD", S20, "$F$7")
    name(wb, "BT_TotC", S20, "$G$7")
    name(wb, "BT_TotSF", S20, "$H$7")
    protect(ws)


def build_balanco(wb):
    ws = wb.create_sheet(S21)
    title(ws, "21 — BALANÇO", "Construído a partir dos saldos do Plano (rubrica por sinal do saldo). Data de reporte vs abertura do exercício.", "CAMADA 3 — OUTPUT")
    header(ws, 5, 1, ["Código", "Rubrica", "Data de reporte", "Abertura do exercício", "Variação"], [14, 52, 17, 17, 15])
    put(ws, "C4", "=CFG_DataRef", "grey", DATE)
    put(ws, "D4", "=CFG_DataIni", "grey", DATE)
    secs = [("ACTIVO", None), ("Activo não corrente", "A"), ("Activo corrente", "A"), ("CAPITAL PRÓPRIO E PASSIVO", None),
            ("Capital próprio", "P"), ("Passivo não corrente", "P"), ("Passivo corrente", "P")]
    r = 6
    tot = {}
    for sec, side in secs:
        if side is None:
            section(ws, r, 1, sec, 5)
            r += 1
            continue
        put(ws, (r, 2), sec, "label", bold=True)
        r += 1
        first = r
        for code, nome, s in D.RUBRICAS_BAL:
            if s != sec:
                continue
            put(ws, (r, 1), code, "note")
            put(ws, (r, 2), "   " + nome, "label")
            if side == "A":
                put(ws, (r, 3), f"=SUMIFS(PC_Dev,PC_RubD,A{r})-SUMIFS(PC_Cr,PC_RubC,A{r})", "calc", NUM)
                put(ws, (r, 4), f"=SUMIFS(PC_AbeDev,PC_RubD,A{r})-SUMIFS(PC_AbeCr,PC_RubC,A{r})", "calc", NUM)
            else:
                put(ws, (r, 3), f"=SUMIFS(PC_Cr,PC_RubC,A{r})-SUMIFS(PC_Dev,PC_RubD,A{r})", "calc", NUM)
                put(ws, (r, 4), f"=SUMIFS(PC_AbeCr,PC_RubC,A{r})-SUMIFS(PC_AbeDev,PC_RubD,A{r})", "calc", NUM)
            put(ws, (r, 5), f"=C{r}-D{r}", "calc", NUM)
            name(wb, "BS_" + code, S21, f"$C${r}")
            name(wb, "BSA_" + code, S21, f"$D${r}")
            r += 1
        put(ws, (r, 2), "Total " + sec.lower(), "label", bold=True)
        for c in "CDE":
            put(ws, (r, "CDE".index(c) + 3), f"=SUM({c}{first}:{c}{r - 1})", "grey", NUM, bold=True)
        tot[sec] = r
        r += 1
        if sec == "Activo corrente":
            put(ws, (r, 2), "TOTAL DO ACTIVO", "label", bold=True)
            for c in "CDE":
                put(ws, (r, "CDE".index(c) + 3), f"={c}{tot['Activo não corrente']}+{c}{tot['Activo corrente']}", "grey", NUM, bold=True)
            tot["AT"] = r
            r += 2
    put(ws, (r, 2), "TOTAL DO PASSIVO", "label", bold=True)
    for c in "CDE":
        put(ws, (r, "CDE".index(c) + 3), f"={c}{tot['Passivo não corrente']}+{c}{tot['Passivo corrente']}", "grey", NUM, bold=True)
    tot["PT"] = r
    r += 1
    put(ws, (r, 2), "TOTAL DO CAPITAL PRÓPRIO E PASSIVO", "label", bold=True)
    for c in "CDE":
        put(ws, (r, "CDE".index(c) + 3), f"={c}{tot['Capital próprio']}+{c}{tot['PT']}", "grey", NUM, bold=True)
    tot["CPP"] = r
    r += 2
    put(ws, (r, 2), "TESTE: ACTIVO = CAPITAL PRÓPRIO + PASSIVO", "label", bold=True)
    put(ws, (r, 3), f'=IF(ABS(C{tot["AT"]}-C{tot["CPP"]})<0.005,"🟢 Balanço equilibrado","🔴 ERRO — Balanço desequilibrado")', "grey", bold=True)
    put(ws, (r, 4), f'=IF(ABS(D{tot["AT"]}-D{tot["CPP"]})<0.005,"🟢 Abertura equilibrada","🔴 Abertura desequilibrada")', "grey")
    status_cf(ws, f"C{r}:D{r}")
    name(wb, "BS_Check", S21, f"$C${r}")
    put(ws, (r + 1, 2), "TESTE: Resultado líquido no Balanço = Resultado líquido na DR", "label")
    put(ws, (r + 1, 3), '=IF(ABS(BS_BCP_RL-DR_RL)<0.005,"🟢 Balanço ↔ DR coerentes","🟡 Existem lançamentos de apuramento/encerramento no período")', "grey")
    status_cf(ws, f"C{r + 1}")
    name(wb, "BS_RLCheck", S21, f"$C${r + 1}")
    put(ws, (r + 2, 2), "Contas de movimento sem rubrica atribuída (ver 01_PLANO_CONTAS)", "label")
    put(ws, (r + 2, 3), '=COUNTIF(PC_Val,"🔴 Sem rubrica*")', "grey", "0")
    for k, t in [("BS_ANC", "Activo não corrente"), ("BS_AC", "Activo corrente"), ("BS_CP", "Capital próprio"), ("BS_PNC", "Passivo não corrente"),
                 ("BS_PC", "Passivo corrente"), ("BS_AT", "AT"), ("BS_PT", "PT")]:
        name(wb, k, S21, f"$C${tot[t]}")
        name(wb, k.replace("BS_", "BSA_"), S21, f"$D${tot[t]}")
    protect(ws)


def build_dre(wb):
    ws = wb.create_sheet(S22)
    title(ws, "22 — DEMONSTRAÇÃO DE RESULTADOS (por natureza)", "Proveitos positivos, custos negativos. Lançamentos de apuramento/encerramento excluídos.", "CAMADA 3 — OUTPUT")
    cols = MESES + ["T1", "T2", "T3", "T4", "S1", "S2", "Acumulado", "Ano"]
    header(ws, 5, 1, ["Código", "Rubrica"] + cols, [7, 46] + [12] * 20)
    ws.freeze_panes = "C6"
    put(ws, "U4", '="até "&TEXT(CFG_DataRef,"dd/mm/yyyy")', "grey")

    def crit(j):
        if j < 12:
            return f",J_Mes,{j + 1}"
        if j < 16:
            return f",J_Trim,{j - 11}"
        if j < 18:
            return f",J_Sem,{j - 15}"
        if j == 18:
            return ',J_Mes,"<="&CFG_MesRep'
        return ""
    struct = [("DR01",), ("DR02",), ("SUB", "Receita (vendas + serviços)", ["DR01", "DR02"], "REC"), ("DR03",),
              ("SUB", "PROVEITOS OPERACIONAIS", ["DR01", "DR02", "DR03"], "PROV"), ("DR04",), ("DR05",), ("DR06",), ("DR07",), ("DR08",),
              ("SUB", "CUSTOS OPERACIONAIS", ["DR04", "DR05", "DR06", "DR07", "DR08"], "CUST"),
              ("SUB", "RESULTADO OPERACIONAL (EBIT)", ["DR01", "DR02", "DR03", "DR04", "DR05", "DR06", "DR07", "DR08"], "RO"),
              ("DR09",), ("DR10",), ("SUB", "RESULTADO FINANCEIRO", ["DR09", "DR10"], "RF"), ("DR11",), ("DR12",),
              ("SUB", "RESULTADOS NÃO OPERACIONAIS", ["DR11", "DR12"], "RNO"),
              ("SUB", "RESULTADO ANTES DE IMPOSTOS", ["DR01", "DR02", "DR03", "DR04", "DR05", "DR06", "DR07", "DR08", "DR09", "DR10", "DR11", "DR12"], "RAI"),
              ("DR13",), ("SUB", "RESULTADO LÍQUIDO DO EXERCÍCIO", [c for c, *_ in D.RUBRICAS_DR], "RL")]
    nomes = {c: n for c, n, *_ in D.RUBRICAS_DR}
    rowof = {}
    r = 6
    for it in struct:
        if it[0] != "SUB":
            code = it[0]
            rowof[code] = r
            put(ws, (r, 1), code, "note")
            put(ws, (r, 2), nomes[code], "label")
            for j in range(20):
                put(ws, (r, 3 + j), f'=-SUMIFS(J_DC,J_RubDR,$A{r},J_Ano,CFG_Ano,J_IncDR,1{crit(j)})', "calc", NUM)
        else:
            _, lab, codes, key = it
            rowof[key] = r
            put(ws, (r, 2), lab, "label", bold=True)
            for j in range(20):
                c = CL(3 + j)
                put(ws, (r, 3 + j), "=" + "+".join(f"{c}{rowof[k]}" for k in codes), "grey", NUM, bold=True)
        r += 1
    r += 1
    extra = [("EBITDA (RO + amortizações)", lambda c: f"={c}{rowof['RO']}-{c}{rowof['DR07']}", NUM, "EBITDA"),
             ("Margem bruta", lambda c: f"=IF({c}{rowof['REC']}=0,0,({c}{rowof['REC']}+{c}{rowof['DR04']})/{c}{rowof['REC']})", PCT, "MB"),
             ("Margem EBITDA", lambda c: f"=IF({c}{rowof['REC']}=0,0,{c}{rowof['EBITDA']}/{c}{rowof['REC']})", PCT, "MEBITDA"),
             ("Margem operacional", lambda c: f"=IF({c}{rowof['REC']}=0,0,{c}{rowof['RO']}/{c}{rowof['REC']})", PCT, "MO"),
             ("Margem líquida", lambda c: f"=IF({c}{rowof['REC']}=0,0,{c}{rowof['RL']}/{c}{rowof['REC']})", PCT, "ML")]
    for lab, f, fmt, key in extra:
        rowof[key] = r
        put(ws, (r, 2), lab, "label", bold=key == "EBITDA")
        for j in range(20):
            put(ws, (r, 3 + j), f(CL(3 + j)), "grey", fmt)
        r += 1
    # nomes (coluna acumulado = U)
    for key, nm in [("REC", "DR_Rec"), ("PROV", "DR_Prov"), ("DR04", "DR_CMV"), ("DR05", "DR_FSE"), ("DR06", "DR_Pessoal"), ("DR07", "DR_Amort"),
                    ("CUST", "DR_Cust"), ("RO", "DR_RO"), ("DR10", "DR_Juros"), ("RF", "DR_RF"), ("RAI", "DR_RAI"), ("DR13", "DR_Imp"),
                    ("RL", "DR_RL"), ("EBITDA", "DR_EBITDA"), ("MB", "DR_MB"), ("MEBITDA", "DR_MEBITDA"), ("MO", "DR_MO"), ("ML", "DR_ML")]:
        name(wb, nm, S22, f"$U${rowof[key]}")
        name(wb, nm.replace("DR_", "DRM_"), S22, f"$C${rowof[key]}:$N${rowof[key]}")
    name(wb, "DR_Codes", S22, f"$A$6:$A${r}")
    name(wb, "DR_Acum", S22, f"$U$6:$U${r}")
    name(wb, "DR_Months", S22, f"$C$6:$N${r}")
    r += 1
    put(ws, (r, 2), "Controlo: RL da DR vs contas de resultados (classes 6, 7 e 87)", "label")
    put(ws, (r, 21), '=IF(ABS(DR_RL-(-SUMIFS(J_DC,J_Classe,"6",'+YTD+',J_IncDR,1)-SUMIFS(J_DC,J_Classe,"7",'+YTD+',J_IncDR,1)-SUMIFS(J_DC,J_G2,"87",'+YTD+',J_IncDR,1)))<0.005,"🟢 DR conciliada","🔴 Contas de resultados sem rubrica DR")', "grey")
    status_cf(ws, f"U{r}")
    name(wb, "DR_Estado", S22, f"$U${r}")
    protect(ws)
    return rowof


def build_fluxo(wb):
    ws = wb.create_sheet(S19)
    title(ws, "19 — FLUXO DE CAIXA MENSAL (método directo)", "Classificação automática: natureza da operação da linha de meios monetários → actividade → linha da DFC.", "CAMADA 3 — OUTPUT")
    header(ws, 5, 1, ["Rubrica"] + MESES + ["Acumulado"], [46] + [12] * 13)
    ws.freeze_panes = "B6"
    r = 6
    put(ws, (r, 1), "Saldo inicial de meios monetários", "label", bold=True)
    for m in range(1, 13):
        put(ws, (r, 1 + m), "=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_IsAbe,1)" if m == 1 else f"={CL(m)}{{SF}}", "grey", NUM)
    put(ws, (r, 14), "=B6", "grey", NUM)
    rsi = r
    r += 1
    act_rows = {}
    for act in ["Operacional", "Investimento", "Financiamento"]:
        section(ws, r, 1, f"Actividades de {act.lower()}", 14)
        r += 1
        first = r
        for a, lin in D.DFC_LINHAS:
            if a != act:
                continue
            put(ws, (r, 1), "   " + lin, "label")
            for m in range(1, 13):
                put(ws, (r, 1 + m), f'=SUMIFS(J_DC,J_LinhaDFC,"{lin}",J_Ano,CFG_Ano,J_Mes,{m})', "calc", NUM)
            put(ws, (r, 14), f'=SUMPRODUCT((COLUMN(B{r}:M{r})-1<=CFG_MesRep)*B{r}:M{r})', "grey", NUM)
            r += 1
        put(ws, (r, 1), f"Fluxo de caixa das actividades de {act.lower()}", "label", bold=True)
        for m in range(1, 14):
            put(ws, (r, 1 + m), f'=SUMIFS(J_DC,J_Fluxo,"{act}",J_Ano,CFG_Ano,J_Mes,{m})' if m <= 12 else f'=SUMPRODUCT((COLUMN(B{r}:M{r})-1<=CFG_MesRep)*B{r}:M{r})', "grey", NUM, bold=True)
        put(ws, (r, 15), f'=IF(ABS(SUM(B{r}:M{r})-SUM(B{first}:M{r - 1}))<0.005,"🟢","🔴 Linhas ≠ total")', "grey")
        act_rows[act] = r
        r += 1
    put(ws, (r, 1), "Transferências internas entre meios monetários (deve ser 0)", "label")
    for m in range(1, 13):
        put(ws, (r, 1 + m), f'=SUMIFS(J_DC,J_Fluxo,"Excluído",J_Ano,CFG_Ano,J_Mes,{m})', "calc", NUM)
    rtr = r
    r += 1
    put(ws, (r, 1), "Movimentos de caixa não classificados", "label")
    for m in range(1, 13):
        put(ws, (r, 1 + m), f'=SUMIFS(J_DC,J_Fluxo,"NÃO CLASSIFICADO",J_Ano,CFG_Ano,J_Mes,{m})', "calc", NUM)
    rnc = r
    r += 1
    put(ws, (r, 1), "VARIAÇÃO LÍQUIDA DE CAIXA", "label", bold=True)
    for m in range(1, 13):
        c = CL(1 + m)
        put(ws, (r, 1 + m), f"={c}{act_rows['Operacional']}+{c}{act_rows['Investimento']}+{c}{act_rows['Financiamento']}+{c}{rtr}+{c}{rnc}", "grey", NUM, bold=True)
    put(ws, (r, 14), f'=SUMPRODUCT((COLUMN(B{r}:M{r})-1<=CFG_MesRep)*B{r}:M{r})', "grey", NUM, bold=True)
    rvar = r
    r += 1
    put(ws, (r, 1), "SALDO FINAL DE MEIOS MONETÁRIOS", "label", bold=True)
    for m in range(1, 13):
        c = CL(1 + m)
        put(ws, (r, 1 + m), f"={c}{rsi}+{c}{rvar}", "grey", NUM, bold=True)
    put(ws, (r, 14), "=INDEX(B{0}:M{0},CFG_MesRep)".format(r), "grey", NUM, bold=True)
    rsf = r
    for m in range(2, 13):
        ws.cell(row=rsi, column=1 + m).value = f"={CL(m)}{rsf}"
    r += 1
    put(ws, (r, 1), "Controlo: saldo contabilístico das contas de meios monetários", "label")
    for m in range(1, 13):
        put(ws, (r, 1 + m), f'=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_Mes,"<="&{m})', "calc", NUM)
    rchk = r
    r += 1
    put(ws, (r, 1), "Estado", "label", bold=True)
    for m in range(1, 13):
        c = CL(1 + m)
        put(ws, (r, 1 + m), f'=IF({m}>CFG_MesRep,"",IF(ABS({c}{rsf}-{c}{rchk})>0.005,"🔴 Diferença",IF(ABS({c}{rnc})>0.005,"🟡 Não classificado",IF({c}{rsf}<0,"🔴 Saldo negativo","🟢 OK"))))', "calc")
    status_cf(ws, f"B{r}:M{r}")
    put(ws, (r, 14), f'=IF(COUNTIF(B{r}:M{r},"🔴*")>0,"🔴 Fluxo de caixa não conciliado",IF(COUNTIF(B{r}:M{r},"🟡*")>0,"🟡 Movimentos por classificar","🟢 Fluxo de caixa conciliado"))', "grey", bold=True)
    status_cf(ws, f"N{r}")
    name(wb, "FC_Estado", S19, f"$N${r}")
    name(wb, "FC_EstRow", S19, f"$B${r}:$M${r}")
    name(wb, "FC_SF", S19, f"$B${rsf}:$M${rsf}")
    name(wb, "FC_Op", S19, f"$B${act_rows['Operacional']}:$M${act_rows['Operacional']}")
    name(wb, "FC_OpAcum", S19, f"$N${act_rows['Operacional']}")
    name(wb, "FC_InvAcum", S19, f"$N${act_rows['Investimento']}")
    name(wb, "FC_FinAcum", S19, f"$N${act_rows['Financiamento']}")
    name(wb, "FC_Var", S19, f"$B${rvar}:$M${rvar}")
    name(wb, "FC_Saldo", S19, f"$N${rsf}")
    protect(ws)


def build_dfc(wb):
    ws = wb.create_sheet(S23)
    title(ws, "23 — DEMONSTRAÇÃO DOS FLUXOS DE CAIXA (método directo)", "Recebimentos e pagamentos efectivos nas contas de meios monetários (prefixos em 00_CONFIGURAÇÃO), classificados pela natureza da operação.", "CAMADA 3 — OUTPUT")
    header(ws, 5, 1, ["Rubrica", "Acumulado até ao mês de reporte", "Exercício completo"], [58, 20, 20])
    r = 6
    subs = {}
    for act in ["Operacional", "Investimento", "Financiamento"]:
        section(ws, r, 1, f"FLUXOS DE CAIXA DAS ACTIVIDADES {'OPERACIONAIS' if act == 'Operacional' else 'DE ' + act.upper()}", 3)
        r += 1
        first = r
        for a, lin in D.DFC_LINHAS:
            if a != act:
                continue
            put(ws, (r, 1), "   " + lin, "label")
            put(ws, (r, 2), f'=SUMIFS(J_DC,J_LinhaDFC,"{lin}",{YTD})', "calc", NUM)
            put(ws, (r, 3), f'=SUMIFS(J_DC,J_LinhaDFC,"{lin}",J_Ano,CFG_Ano)', "calc", NUM)
            r += 1
        put(ws, (r, 1), f"Fluxo líquido — actividades {act.lower()}{'is' if act == 'Operacional' else ''}".replace("operacionalis", "operacionais"), "label", bold=True)
        put(ws, (r, 2), f"=SUM(B{first}:B{r - 1})", "grey", NUM, bold=True)
        put(ws, (r, 3), f"=SUM(C{first}:C{r - 1})", "grey", NUM, bold=True)
        subs[act] = r
        r += 2
    put(ws, (r, 1), "Transferências internas / não classificados", "label")
    put(ws, (r, 2), f'=SUMIFS(J_DC,J_Fluxo,"Excluído",{YTD})+SUMIFS(J_DC,J_Fluxo,"NÃO CLASSIFICADO",{YTD})', "calc", NUM)
    put(ws, (r, 3), '=SUMIFS(J_DC,J_Fluxo,"Excluído",J_Ano,CFG_Ano)+SUMIFS(J_DC,J_Fluxo,"NÃO CLASSIFICADO",J_Ano,CFG_Ano)', "calc", NUM)
    rx = r
    r += 1
    put(ws, (r, 1), "VARIAÇÃO LÍQUIDA DE CAIXA E EQUIVALENTES", "label", bold=True)
    for c in "BC":
        put(ws, (r, "BC".index(c) + 2), f"={c}{subs['Operacional']}+{c}{subs['Investimento']}+{c}{subs['Financiamento']}+{c}{rx}", "grey", NUM, bold=True)
    rv = r
    r += 1
    put(ws, (r, 1), "Caixa e equivalentes no início do exercício", "label")
    put(ws, (r, 2), "=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_IsAbe,1)", "calc", NUM)
    put(ws, (r, 3), f"=B{r}", "calc", NUM)
    ri = r
    r += 1
    put(ws, (r, 1), "CAIXA E EQUIVALENTES NO FIM DO PERÍODO", "label", bold=True)
    put(ws, (r, 2), f"=B{ri}+B{rv}", "grey", NUM, bold=True)
    put(ws, (r, 3), f"=C{ri}+C{rv}", "grey", NUM, bold=True)
    rf = r
    r += 2
    put(ws, (r, 1), "Controlo: saldo das contas de meios monetários à data de reporte", "label")
    put(ws, (r, 2), '=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "grey", NUM)
    put(ws, (r, 3), f'=IF(ABS(B{rf}-B{r})<0.005,"🟢 DFC conciliada com o Balanço/Razão","🔴 DFC ≠ saldo de caixa")', "grey")
    status_cf(ws, f"C{r}")
    name(wb, "DFC_Estado", S23, f"$C${r}")
    put(ws, (r + 1, 1), "Nota: juros pagos classificados em financiamento (opção de política — alterável em 01A_TABELAS). Depósitos a prazo tratados como investimento, não como equivalente de caixa.", "note")
    name(wb, "DFC_Op", S23, f"$B${subs['Operacional']}")
    name(wb, "DFC_Inv", S23, f"$B${subs['Investimento']}")
    name(wb, "DFC_Fin", S23, f"$B${subs['Financiamento']}")
    name(wb, "DFC_Var", S23, f"$B${rv}")
    name(wb, "DFC_Fim", S23, f"$B${rf}")
    protect(ws)


def build_orcamento(wb):
    ws = wb.create_sheet(S16)
    title(ws, "16 — ORÇAMENTO EMPRESARIAL", "Orçamentos parciais (vendas, compras, pessoal, despesas, investimento, financiamento) → orçamento de resultados, de caixa e patrimonial. Valores de exemplo (fictícios).",
          "CAMADA 1 — INPUT (azul)  |  CAMADA 2 — consolidação")
    header(ws, 5, 1, ["Código", "Rubrica"] + MESES + ["Total"], [9, 46] + [12] * 13)
    ws.freeze_panes = "C6"
    r = 6

    def row_input(label, vals, fmt=NUM, code=""):
        nonlocal r
        put(ws, (r, 1), code, "note")
        put(ws, (r, 2), label, "label")
        for m in range(12):
            v = vals(m + 1) if callable(vals) else vals
            if D.PRODUCAO and fmt != PCT:
                v = 0
            put(ws, (r, 3 + m), v, "input", fmt)
        put(ws, (r, 15), f"=SUM(C{r}:N{r})" if fmt == NUM else "", "grey", fmt)
        r += 1
        return r - 1

    def row_calc(label, f, fmt=NUM, code="", bold=False):
        nonlocal r
        put(ws, (r, 1), code, "note")
        put(ws, (r, 2), label, "label", bold=bold)
        for m in range(12):
            put(ws, (r, 3 + m), f(CL(3 + m), m + 1), "grey", fmt, bold=bold)
        put(ws, (r, 15), f"=SUM(C{r}:N{r})" if fmt == NUM else "", "grey", fmt, bold=bold)
        r += 1
        return r - 1
    section(ws, r, 1, "1. ORÇAMENTO DE VENDAS", 15); r += 1
    vol = row_input("Volume de vendas de mercadorias (unidades)", 300, "#,##0")
    prc = row_input("Preço médio de venda (Kz/un)", 1000, NUM)
    ven = row_calc("Vendas de mercadorias", lambda c, m: f"={c}{vol}*{c}{prc}", code="DR01")
    hrs = row_input("Serviços: nº de projectos / unidades de serviço", 1, "#,##0")
    tar = row_input("Serviços: valor médio por unidade (Kz)", 1_000_000, NUM)
    srv = row_calc("Prestações de serviço", lambda c, m: f"={c}{hrs}*{c}{tar}", code="DR02")
    section(ws, r, 1, "2. ORÇAMENTO DE COMPRAS E CUSTO DAS VENDAS", 15); r += 1
    cmvp = row_input("Custo das mercadorias vendidas (% das vendas)", 0.6, PCT)
    cmv = row_calc("Custo das mercadorias vendidas", lambda c, m: f"={c}{ven}*{c}{cmvp}", code="DR04")
    dst = row_input("Variação de stock desejada (Kz, + aumento)", 0, NUM)
    cmp_ = row_calc("Compras de mercadorias", lambda c, m: f"={c}{cmv}+{c}{dst}")
    section(ws, r, 1, "3. ORÇAMENTO DE PESSOAL", 15); r += 1
    hc = row_input("Efectivo médio (nº trabalhadores)", 1, "#,##0")
    sal = row_input("Remuneração bruta média mensal (Kz)", 400_000, NUM)
    enc = row_input("Encargos sociais da entidade (%)", 0.08, PCT)
    pes = row_calc("Custos com o pessoal", lambda c, m: f"={c}{hc}*{c}{sal}*(1+{c}{enc})", code="DR06")
    section(ws, r, 1, "4. ORÇAMENTO DE DESPESAS E OUTRAS RUBRICAS", 15); r += 1
    fse = row_input("Fornecimentos e serviços de terceiros", 350_000, NUM, "DR05")
    amo_extra = row_input("Amortizações adicionais de novo investimento", 0, NUM)
    amo = row_calc("Amortizações (registo de activos + adicionais)", lambda c, m: f"=INDEX(AF_DepMesReg,{m})+{c}{amo_extra}", code="DR07")
    oco = row_input("Outros custos operacionais", 0, NUM, "DR08")
    opo = row_input("Outros proveitos operacionais", 0, NUM, "DR03")
    pfi = row_input("Proveitos financeiros", 0, NUM, "DR09")
    cfi = row_input("Custos financeiros", lambda m: 0 if m == 1 else 45_000, NUM, "DR10")
    pno = row_input("Proveitos não operacionais / extraordinários", 0, NUM, "DR11")
    cno = row_input("Custos não operacionais / extraordinários", 0, NUM, "DR12")
    section(ws, r, 1, "5. ORÇAMENTO DE RESULTADOS", 15); r += 1
    rec = row_calc("Proveitos operacionais", lambda c, m: f"={c}{ven}+{c}{srv}+{c}{opo}", bold=True)
    ro = row_calc("Resultado operacional orçado", lambda c, m: f"={c}{rec}-{c}{cmv}-{c}{fse}-{c}{pes}-{c}{amo}-{c}{oco}", bold=True)
    rai = row_calc("Resultado antes de impostos orçado", lambda c, m: f"={c}{ro}+{c}{pfi}-{c}{cfi}+{c}{pno}-{c}{cno}", bold=True)
    imp = row_calc("Imposto Industrial estimado", lambda c, m: f"=MAX(0,{c}{rai})*TX_II", code="DR13")
    rl = row_calc("Resultado líquido orçado", lambda c, m: f"={c}{rai}-{c}{imp}", bold=True)
    section(ws, r, 1, "6. ORÇAMENTO DE INVESTIMENTO", 15); r += 1
    capex = row_input("Investimento em activos fixos (capex)", lambda m: 1_200_000 if m == 2 else 0, NUM)
    section(ws, r, 1, "7. ORÇAMENTO DE FINANCIAMENTO", 15); r += 1
    nemp = row_input("Novos empréstimos", lambda m: 3_000_000 if m == 2 else 0, NUM)
    reem = row_input("Reembolsos de empréstimos", 0, NUM)
    acap = row_input("Aumentos de capital", lambda m: 2_000_000 if m == 3 else 0, NUM)
    divd = row_input("Dividendos a pagar", 0, NUM)
    section(ws, r, 1, "8. ORÇAMENTO DE CAIXA", 15); r += 1
    pcob = row_input("% da receita (c/ IVA) cobrada no próprio mês", 0.5, PCT)
    ppag = row_input("% das compras/FSE (c/ IVA) pagas no próprio mês", 0.5, PCT)
    iva = "(1+INDEX(TX_Taxa,MATCH(\"IVA_GER\",TX_Cod,0)))"
    ent = row_calc("Recebimentos de clientes", lambda c, m: f"=({c}{ven}+{c}{srv})*{iva}*{c}{pcob}" + (f"+({CL(1 + m)}{ven}+{CL(1 + m)}{srv})*{iva}*(1-{CL(1 + m)}{pcob})" if m > 1 else ""))
    sai = row_calc("Pagamentos a fornecedores", lambda c, m: f"=({c}{cmp_}+{c}{fse})*{iva}*{c}{ppag}" + (f"+({CL(1 + m)}{cmp_}+{CL(1 + m)}{fse})*{iva}*(1-{CL(1 + m)}{ppag})" if m > 1 else ""))
    spes = row_calc("Pagamentos ao pessoal", lambda c, m: f"={c}{pes}")
    sfin = row_calc("Encargos financeiros", lambda c, m: f"={c}{cfi}")
    cxent = row_calc("TOTAL DE ENTRADAS PREVISTAS", lambda c, m: f"={c}{ent}+{c}{nemp}+{c}{acap}+{c}{pfi}", bold=True)
    cxsai = row_calc("TOTAL DE SAÍDAS PREVISTAS", lambda c, m: f"={c}{sai}+{c}{spes}+{c}{sfin}+{c}{capex}+{c}{reem}+{c}{divd}", bold=True)
    cxsi = row_calc("Saldo inicial previsto", lambda c, m: "=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_IsAbe,1)" if m == 1 else f"={CL(1 + m)}{r + 1}")
    cxsf = row_calc("Saldo final previsto", lambda c, m: f"={c}{cxsi}+{c}{cxent}-{c}{cxsai}", bold=True)
    for m in range(12):
        ws.cell(row=cxsi, column=3 + m).number_format = NUM
    ws.cell(row=cxsi, column=15).value = ""
    ws.cell(row=cxsf, column=15).value = f"=N{cxsf}"
    section(ws, r, 1, "9. ORÇAMENTO PATRIMONIAL SIMPLIFICADO (fim do exercício)", 15); r += 1
    pat = [("Activo fixo líquido previsto", f"=BSA_BA_IMOB_CORP+BSA_BA_IMOB_INC+BSA_BA_INV_FIN+O{capex}-O{amo}"),
           ("Meios monetários previstos", f"=N{cxsf}"),
           ("Capital próprio previsto", f"=BSA_CP+O{rl}+O{acap}-O{divd}"),
           ("Dívida financeira prevista", f"=BSA_BPNC_EMP+BSA_BPC_EMP+O{nemp}-O{reem}"),
           ("Fundo de maneio implícito (fecho)", "=C{cp}+C{dv}-C{af}-C{mm}")]
    p0 = r
    for i, (lab, f) in enumerate(pat):
        put(ws, (r, 2), lab, "label")
        put(ws, (r, 3), f.format(cp=p0 + 2, dv=p0 + 3, af=p0, mm=p0 + 1), "grey", NUM)
        r += 1
    # mapeamento por rubrica DR (para Budget vs Actual)
    section(ws, r + 1, 1, "MAPA ORÇAMENTO POR RUBRICA DR (alimenta 30_BUDGET_VS_ACTUAL) — valores com sinal: proveitos +, custos −", 15)
    r += 2
    m0 = r
    src = {"DR01": ven, "DR02": srv, "DR03": opo, "DR04": cmv, "DR05": fse, "DR06": pes, "DR07": amo, "DR08": oco, "DR09": pfi, "DR10": cfi,
           "DR11": pno, "DR12": cno, "DR13": imp}
    for code, nome, nat, sig in D.RUBRICAS_DR:
        put(ws, (r, 1), code, "note")
        put(ws, (r, 2), nome, "label")
        s = "" if sig == "S" else "-"
        for m in range(12):
            c = CL(3 + m)
            put(ws, (r, 3 + m), f"={s}{c}{src[code]}", "grey", NUM)
        put(ws, (r, 15), f"=SUM(C{r}:N{r})", "grey", NUM)
        r += 1
    name(wb, "ORC_Codes", S16, f"$A${m0}:$A${r - 1}")
    name(wb, "ORC_M", S16, f"$C${m0}:$N${r - 1}")
    name(wb, "ORC_CxEnt", S16, f"$C${cxent}:$N${cxent}")
    name(wb, "ORC_CxSai", S16, f"$C${cxsai}:$N${cxsai}")
    name(wb, "ORC_Capex", S16, f"$C${capex}:$N${capex}")
    name(wb, "ORC_RL", S16, f"$C${rl}:$N${rl}")
    protect(ws)


def build_bva(wb):
    ws = wb.create_sheet(S30)
    title(ws, "30 — BUDGET VS ACTUAL (Orçamento × Real)", "Acumulado até ao mês de reporte. Desvio = Real − Orçamento (custos com sinal negativo: desvio positivo = favorável).", "CAMADA 3 — OUTPUT (causa, responsável e acção = input)")
    heads = ["Código", "Rubrica", "Orçamento", "Realizado", "Desvio absoluto", "Desvio %", "Avaliação", "Causa", "Responsável", "Acção correctiva",
             "Orç. mês", "Real mês", "Desvio mês"]
    header(ws, 5, 1, heads, [7, 44, 15, 15, 15, 10, 26, 30, 16, 34, 14, 14, 14])
    r = 6
    first = r
    for code, nome, nat, sig in D.RUBRICAS_DR:
        put(ws, (r, 1), code, "note")
        put(ws, (r, 2), nome, "label")
        put(ws, (r, 3), f"=SUMPRODUCT((COLUMN(ORC_M)-COLUMN(INDEX(ORC_M,1,1))+1<=CFG_MesRep)*(ORC_Codes=A{r})*ORC_M)", "calc", NUM)
        put(ws, (r, 4), f"=SUMIFS(DR_Acum,DR_Codes,A{r})", "calc", NUM)
        put(ws, (r, 5), f"=D{r}-C{r}", "calc", NUM)
        put(ws, (r, 6), f"=IF(C{r}=0,0,E{r}/ABS(C{r}))", "calc", PCT)
        if code == "DR13":
            put(ws, (r, 7), f'=IF(AND(C{r}=0,D{r}=0),"—",IF(E{r}<0,"🟡 Imposto acima do orçado (acompanha o resultado)","🟢 Em linha"))', "calc")
        else:
            put(ws, (r, 7), f'=IF(AND(C{r}=0,D{r}=0),"—",IF(F{r}<-CFG_OrcTol,"🔴 Desvio desfavorável"&IF(C{r}<0," — orçamento excedido",""),IF(E{r}<0,"🟡 Desfavorável dentro da tolerância","🟢 Favorável / em linha")))', "calc")
        for c in (8, 9, 10):
            put(ws, (r, c), None, "input")
        put(ws, (r, 11), f"=SUMPRODUCT((ORC_Codes=A{r})*INDEX(ORC_M,0,CFG_MesRep))", "calc", NUM)
        put(ws, (r, 12), f"=SUMPRODUCT((DR_Codes=A{r})*INDEX(DR_Months,0,CFG_MesRep))", "calc", NUM)
        put(ws, (r, 13), f"=L{r}-K{r}", "calc", NUM)
        r += 1
    last = r - 1
    for lab, codes in [("RESULTADO OPERACIONAL", "DR0[1-8]"), ("RESULTADO LÍQUIDO", "ALL")]:
        put(ws, (r, 2), lab, "label", bold=True)
        rng = f"{first}:{first + 7}" if codes != "ALL" else f"{first}:{last}"
        a, b = rng.split(":")
        for c in "CDEKLM":
            put(ws, (r, "ABCDEFGHIJKLM".index(c) + 1), f"=SUM({c}{a}:{c}{b})", "grey", NUM, bold=True)
        put(ws, (r, 6), f"=IF(C{r}=0,0,E{r}/ABS(C{r}))", "grey", PCT, bold=True)
        put(ws, (r, 7), f'=IF(F{r}<-CFG_OrcTol,"🔴 Abaixo do orçamento",IF(E{r}<0,"🟡 Ligeiramente abaixo","🟢 Em linha / acima"))', "grey", bold=True)
        r += 1
    status_cf(ws, f"G6:G{r}")
    put(ws, (r + 1, 2), "Taxa de execução dos custos operacionais (real / orçado)", "label")
    put(ws, (r + 1, 4), f"=IF(SUM(C{first + 3}:C{first + 7})=0,0,SUM(D{first + 3}:D{first + 7})/SUM(C{first + 3}:C{first + 7}))", "grey", PCT)
    name(wb, "BVA_ExecCustos", S30, f"$D${r + 1}")
    put(ws, (r + 2, 2), "Rubricas com desvio desfavorável acima da tolerância", "label")
    put(ws, (r + 2, 4), f'=COUNTIF(G{first}:G{last},"🔴*")', "grey", "0")
    name(wb, "BVA_Red", S30, f"$D${r + 2}")
    put(ws, (r + 3, 2), "Rubricas de custo com orçamento excedido", "label")
    put(ws, (r + 3, 4), f'=COUNTIF(G{first}:G{last},"*excedido*")', "grey", "0")
    name(wb, "BVA_Exced", S30, f"$D${r + 3}")
    put(ws, (r + 4, 2), "Resultado líquido: orçado vs real", "label")
    put(ws, (r + 4, 3), f"=C{r - 1}", "grey", NUM)
    put(ws, (r + 4, 4), f"=D{r - 1}", "grey", NUM)
    name(wb, "BVA_RLOrc", S30, f"$C${r + 4}")
    name(wb, "BVA_Tab", S30, f"$B${first}:$G${last}")
    protect(ws)


def build_tesouraria(wb):
    ws = wb.create_sheet(S18)
    title(ws, "18 — TESOURARIA: PREVISÃO ROLLING DE 13 SEMANAS E MAPA DE NECESSIDADES", "Parte da posição de caixa à data de reporte e das partidas em aberto (clientes/fornecedores) com vencimento em cada semana.",
          "CAMADA 3 — OUTPUT (outras entradas/saídas = input)")
    header(ws, 5, 1, ["Rubrica"] + [f"S{k}" for k in range(1, 14)] + ["Total"], [48] + [12] * 14)
    ws.freeze_panes = "B6"
    rows = ["Início da semana", "Fim da semana", "SALDO INICIAL", "Recebimentos de clientes (vencimento na semana)", "Recuperação de saldos vencidos (probabilidade)",
            "Outras entradas previstas (input)", "TOTAL ENTRADAS", "Pagamentos a fornecedores (vencimento na semana; vencidos na S1)", "Salários líquidos",
            "IVA a pagar", "IRT, INSS e outras retenções", "Serviço da dívida", "Outras saídas previstas (input)", "TOTAL SAÍDAS", "SALDO FINAL PROJECTADO",
            "Saldo mínimo exigido", "Excedente / (défice) face ao mínimo", "Linha de crédito disponível", "Estado"]
    for i, lab in enumerate(rows):
        put(ws, (6 + i, 1), lab, "label", bold=lab.isupper())
    for k in range(1, 14):
        c = CL(1 + k)
        p = CL(k)
        f = {
            6: f"=CFG_DataRef+1+7*({k}-1)", 7: f"={c}6+6",
            8: '=SUMIFS(J_DC,J_Caixa,1,J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)' if k == 1 else f"={p}20",
            9: f'=SUMIFS(CLI_Saldo,CLI_DtVenc,">="&{c}6,CLI_DtVenc,"<="&{c}7)',
            10: f"=IF({k}<=4,CLI_Venc_T*CFG_ProbCob/4,0)",
            11: None, 12: f"=SUM({c}9:{c}11)",
            13: f'=SUMIFS(FRN_Saldo,FRN_DtVenc,">="&{c}6,FRN_DtVenc,"<="&{c}7)' + (f'+SUMIFS(FRN_Saldo,FRN_DtVenc,"<"&{c}6)' if k == 1 else ""),
            14: f"=IF(OR(AND(DATE(YEAR({c}6),MONTH({c}6),CFG_DiaSal)>={c}6,DATE(YEAR({c}6),MONTH({c}6),CFG_DiaSal)<={c}7),AND(MONTH({c}7)<>MONTH({c}6),DATE(YEAR({c}7),MONTH({c}7),CFG_DiaSal)<={c}7)),CFG_SalPrev,0)",
            15: f"=IF(AND(EOMONTH(CFG_DataRef,1)>={c}6,EOMONTH(CFG_DataRef,1)<={c}7),IVA_PagarMes,0)",
            16: f'=IF(AND(EOMONTH(CFG_DataRef,1)>={c}6,EOMONTH(CFG_DataRef,1)<={c}7),-(SUMIFS(J_DC,J_Conta,"34.3.*",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)+SUMIFS(J_DC,J_Conta,"34.9.1",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)+SUMIFS(J_DC,J_Conta,"34.6",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)),0)',
            17: f"=IF(OR(AND(EOMONTH({c}6,0)>={c}6,EOMONTH({c}6,0)<={c}7)),CFG_DividaPrev,0)",
            18: None, 19: f"=SUM({c}13:{c}18)", 20: f"={c}8+{c}12-{c}19", 21: "=CFG_SaldoMin", 22: f"={c}20-{c}21", 23: "=CFG_LimTes",
            24: f'=IF({c}20<0,IF(-{c}20<={c}23,"🔴 Défice — usar linha de crédito","🔴 Défice acima da linha"),IF({c}22<0,"🟡 Abaixo do saldo mínimo","🟢 Normal"))',
        }
        for rr, v in f.items():
            put(ws, (rr, 1 + k), v, "input" if v is None else "calc", DATE if rr in (6, 7) else (None if rr == 24 else NUM))
    for rr in (9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19):
        put(ws, (rr, 15), f"=SUM(B{rr}:N{rr})", "grey", NUM)
    status_cf(ws, "B24:N24")
    put(ws, "A4", "Vencido de clientes considerado (Kz):", "label")
    put(ws, "C4", '=SUMIFS(CLI_Saldo,CLI_DtVenc,"<="&CFG_DataRef)', "grey", NUM)
    name(wb, "CLI_Venc_T", S18, "$C$4")
    section(ws, 26, 1, "MAPA DE NECESSIDADE DE TESOURARIA", 8)
    nec = [("Saldo actual de meios monetários", "=B8", NUM, "TES_Saldo"), ("Saldo mínimo projectado (13 semanas)", "=MIN(B20:N20)", NUM, "TES_Min"),
           ("Semana crítica", '="S"&MATCH(MIN(B20:N20),B20:N20,0)&" ("&TEXT(INDEX(B6:N6,MATCH(MIN(B20:N20),B20:N20,0)),"dd/mm/yyyy")&")"', None, "TES_Semana"),
           ("Necessidade de financiamento para manter o mínimo", "=MAX(0,CFG_SaldoMin-MIN(B20:N20))", NUM, "TES_Necess"),
           ("Excedente disponível para aplicação/investimento", "=MAX(0,MIN(B20:N20)-CFG_SaldoMin)", NUM, "TES_Excedente"),
           ("Estado", '=IF(MIN(B20:N20)<0,IF(-MIN(B20:N20)<=CFG_LimTes,"🔴 Défice coberto pela linha de crédito","🔴 CRÍTICO — défice acima da linha"),IF(MIN(B20:N20)<CFG_SaldoMin,"🟡 Liquidez abaixo do mínimo parametrizado","🟢 Liquidez adequada"))', None, "TES_Estado")]
    for i, (lab, f, fmt, nm) in enumerate(nec):
        put(ws, (27 + i, 1), lab, "label")
        put(ws, (27 + i, 2), f, "grey", fmt)
        name(wb, nm, S18, f"$B${27 + i}")
    status_cf(ws, "B32")
    section(ws, 35, 1, "MAPA MENSAL — REALIZADO × PREVISTO (orçamento de caixa)", 14)
    header(ws, 36, 1, ["Rubrica"] + MESES + ["Acumulado"], None)
    lines = [("Entradas realizadas", '=SUMIFS(J_Deb,J_Caixa,1,J_Ano,CFG_Ano,J_Mes,{m},J_IsAbe,0)-SUMIFS(J_Deb,J_Caixa,1,J_Ano,CFG_Ano,J_Mes,{m},J_Fluxo,"Excluído")'),
             ("Saídas realizadas", '=SUMIFS(J_Cred,J_Caixa,1,J_Ano,CFG_Ano,J_Mes,{m},J_IsAbe,0)-SUMIFS(J_Cred,J_Caixa,1,J_Ano,CFG_Ano,J_Mes,{m},J_Fluxo,"Excluído")'),
             ("Fluxo líquido realizado", "={c}37-{c}38"), ("Entradas previstas (orçamento)", "=INDEX(ORC_CxEnt,{m})"),
             ("Saídas previstas (orçamento)", "=INDEX(ORC_CxSai,{m})"), ("Fluxo líquido previsto", "={c}40-{c}41"), ("Desvio (realizado − previsto)", "={c}39-{c}42"),
             ("Saldo realizado (fim do mês)", "=INDEX(FC_SF,{m})")]
    for i, (lab, f) in enumerate(lines):
        rr = 37 + i
        put(ws, (rr, 1), lab, "label")
        for m in range(1, 13):
            c = CL(1 + m)
            put(ws, (rr, 1 + m), f"=IF({m}>CFG_MesRep,\"\",{f.format(m=m, c=c)[1:]})" if i in (0, 1, 2, 6, 7) else f.format(m=m, c=c), "calc", NUM)
        if i in (0, 1, 2, 6):
            put(ws, (rr, 14), f"=SUM(B{rr}:M{rr})", "grey", NUM)
        elif i in (3, 4, 5):
            put(ws, (rr, 14), f"=SUMPRODUCT((COLUMN(B{rr}:M{rr})-1<=CFG_MesRep)*B{rr}:M{rr})", "grey", NUM)
    protect(ws)


def build_all(wb):
    pass
