"""Livros e mapas auxiliares gerados a partir do Diário (camada de OUTPUT/PROCESSAMENTO)."""
from datetime import date
from openpyxl.utils import get_column_letter as CL
from core import *
import data as D

MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
CX = '((J_Ano=CFG_Ano)*1)'


def listing(ws, r0, n, seq, cols, kcol=1):
    """Lista k-ésimas ocorrências de J_<seq>. cols: [(header, width, fmt, template)] ; template usa {r} e {L} (célula com linha)."""
    heads = ["#", "Linha"] + [c[0] for c in cols]
    widths = [5, 6] + [c[1] for c in cols]
    header(ws, r0 - 1, kcol, heads, widths)
    kc, lc = CL(kcol), CL(kcol + 1)
    put(ws, f"{kc}{r0 - 2}", f"=MAX(J_{seq})", "grey", "0")
    for i in range(n):
        r = r0 + i
        put(ws, f"{kc}{r}", f'=IF({i + 1}<=${kc}${r0 - 2},{i + 1},"")', "grey", "0")
        put(ws, f"{lc}{r}", f'=IF({kc}{r}="","",MATCH({kc}{r},J_{seq},0))', "grey", "0")
        for j, (_, _, fmt, tpl) in enumerate(cols):
            f = tpl.replace("{r}", str(r)).replace("{L}", f"${lc}{r}").replace("{K}", f"${kc}{r}")
            put(ws, (r, kcol + 2 + j), f, "calc", fmt)
    return r0 + n - 1


def J(col, text=False):
    return f'=IF({{L}}="","",INDEX(J_{col},{{L}}){"&" + chr(34) * 2 if text else ""})'


# ---------------------------------------------------------------------------
def build_diario_geral(wb):
    S = "03_DIÁRIO_GERAL"
    ws = wb.create_sheet(S)
    title(ws, "03 — DIÁRIO GERAL (gerado automaticamente)", "Filtros em azul; vazio = todos. Mostra até 500 linhas.", "CAMADA 3 — OUTPUT")
    filtros = [("Mês inicial", 1, "DG_MesIni", "0"), ("Mês final", "=CFG_MesRep", "DG_MesFim", "0"), ("Conta (prefixo)", None, "DG_Conta", "@"),
               ("Nº documento", None, "DG_Doc", "@"), ("Tipo de documento", None, "DG_TipoDoc", "@"), ("NIF terceiro", None, "DG_NIF", "@"),
               ("Centro de custo", None, "DG_CC", "@"), ("Projecto", None, "DG_Proj", "@"), ("Natureza", None, "DG_Nat", "@"), ("Utilizador", None, "DG_User", "@")]
    for i, (lab, v, nm, fmt) in enumerate(filtros):
        r = 4 + (i % 5)
        c = 1 + 3 * (i // 5)
        put(ws, (r, c + 2), lab, "label")
        put(ws, (r, c + 3), v, "input", fmt)
        name(wb, nm, S, f"${CL(c + 3)}${r}")
    put(ws, "J4", "Linhas filtradas:", "label")
    put(ws, "K4", "=MAX(J_SeqDG)", "grey", "0")
    put(ws, "J5", "Σ Débito (filtro):", "label")
    put(ws, "K5", "=SUM(K11:K510)", "grey", NUM)
    put(ws, "J6", "Σ Crédito (filtro):", "label")
    put(ws, "K6", "=SUM(L11:L510)", "grey", NUM)
    put(ws, "J7", "Equilíbrio:", "label")
    put(ws, "K7", '=IF(ABS(K5-K6)<0.005,"🟢 Equilibrado","🟡 Filtro parcial (D≠C)")', "grey")
    status_cf(ws, "K7")
    cols = [("ID", 7, "0", J("ID")), ("Data", 11, DATE, J("Data")), ("Tipo", 6, None, J("TipoDoc", True)), ("Série", 6, None, J("Serie", True)),
            ("Nº", 7, None, J("NumDoc", True)), ("Conta", 9, None, J("Conta", True)), ("Designação da conta", 28, None, J("NomeConta", True)),
            ("Descrição", 36, None, J("Desc", True)), ("Débito", 14, NUM, J("Deb")), ("Crédito", 14, NUM, J("Cred")),
            ("NIF", 12, None, J("NIF", True)), ("Terceiro", 22, None, J("Terceiro", True)), ("CC", 6, None, J("CC", True)),
            ("Projecto", 7, None, J("Proj", True)), ("Natureza", 20, None, J("Nat", True)), ("Utilizador", 9, None, J("User", True)),
            ("Estado", 13, None, J("EstVal", True))]
    listing(ws, 11, 500, "SeqDG", cols)
    status_cf(ws, "S11:S510")
    ws.freeze_panes = "C11"
    protect(ws)


def build_razao(wb):
    S = "04_RAZÃO"
    ws = wb.create_sheet(S)
    title(ws, "04 — RAZÃO (extracto de conta)", "Conta, subconta ou classe. Filtros opcionais por terceiro, projecto e centro de custo.", "CAMADA 3 — OUTPUT")
    sel = [("Conta / subconta / classe", "43.1.1", "RZ_Conta", "@"), ("Mês inicial", 1, "RZ_MesIni", "0"), ("Mês final", "=CFG_MesRep", "RZ_MesFim", "0"),
           ("NIF terceiro (opcional)", None, "RZ_NIF", "@"), ("Projecto (opcional)", None, "RZ_Proj", "@"), ("Centro de custo (opcional)", None, "RZ_CC", "@")]
    for i, (lab, v, nm, fmt) in enumerate(sel):
        put(ws, (4 + i, 3), lab, "label")
        put(ws, (4 + i, 5), v, "input", fmt)
        name(wb, nm, S, f"$E${4 + i}")
    dv_list(ws, "E4", "=PC_Cod")
    ws.column_dimensions["C"].width = 26
    put(ws, "G4", "Designação:", "label")
    put(ws, "H4", '=IFERROR(INDEX(PC_Desig,MATCH(RZ_Conta,PC_Cod,0)),"—")', "grey")
    acc = '(((J_Conta=RZ_Conta&"")+(LEFT(J_Conta,LEN(RZ_Conta)+1)=RZ_Conta&".")+((LEN(RZ_Conta)=1)*(J_Classe=RZ_Conta&"")))>0)'
    flt = '(((RZ_NIF="")+(J_NIF=RZ_NIF&""))>0)*(((RZ_Proj="")+(J_Proj=RZ_Proj&""))>0)*(((RZ_CC="")+(J_CC=RZ_CC&""))>0)'
    base = f"(J_Ano=CFG_Ano)*{acc}*{flt}"
    put(ws, "G5", "Saldo inicial", "label")
    put(ws, "H5", f'=SUMPRODUCT({base}*(((J_IsAbe=1)+(J_Mes<RZ_MesIni)*(J_IsAbe=0))>0)*J_DC)', "grey", NUM)
    put(ws, "G6", "+ Débitos do período", "label")
    per = f"{base}*(J_IsAbe=0)*(J_Mes>=RZ_MesIni)*(J_Mes<=RZ_MesFim)"
    put(ws, "H6", f'=SUMPRODUCT({per}*(J_DC>0)*J_DC)', "grey", NUM)
    put(ws, "G7", "− Créditos do período", "label")
    put(ws, "H7", f'=-SUMPRODUCT({per}*(J_DC<0)*J_DC)', "grey", NUM)
    put(ws, "G8", "Saldo final", "label", bold=True)
    put(ws, "H8", "=H5+H6-H7", "grey", NUM, bold=True)
    put(ws, "I8", '=IF(ABS(H8-IFERROR(IF(AND(RZ_NIF="",RZ_Proj="",RZ_CC="",RZ_MesFim=CFG_MesRep),INDEX(PC_Saldo,MATCH(RZ_Conta,PC_Cod,0)),H8),H8))<0.005,"🟢 Reconciliado com o Plano/Balancete","🔴 Diferença")', "grey")
    status_cf(ws, "I8")
    ws.column_dimensions["H"].width = 16
    cols = [("ID", 7, "0", J("ID")), ("Data", 11, DATE, J("Data")), ("Documento", 14, None, '=IF({L}="","",INDEX(J_TipoDoc,{L})&" "&INDEX(J_Serie,{L})&"/"&INDEX(J_NumDoc,{L}))'),
            ("Conta", 9, None, J("Conta", True)), ("Descrição", 36, None, J("Desc", True)), ("Débito", 14, NUM, J("Deb")), ("Crédito", 14, NUM, J("Cred")),
            ("Saldo acumulado", 16, NUM, '=IF({L}="","",IF({K}=1,$H$5,N(OFFSET(INDIRECT("J{r}"),-1,0)))+INDEX(J_DC,{L}))'),
            ("Terceiro", 22, None, J("Terceiro", True)), ("Projecto", 7, None, J("Proj", True))]
    last = listing(ws, 12, 500, "SeqRZ", cols)
    # saldo acumulado sem OFFSET/INDIRECT (mais robusto)
    for r in range(12, last + 1):
        ws[f"J{r}"] = f'=IF($B{r}="","",IF($A{r}=1,$H$5,J{r - 1})+INDEX(J_DC,$B{r}))' if r > 12 else f'=IF($B{r}="","",$H$5+INDEX(J_DC,$B{r}))'
        ws[f"J{r}"].number_format = NUM
    ws.freeze_panes = "C12"
    protect(ws)


def build_caixa(wb):
    S = "05_CAIXA"
    ws = wb.create_sheet(S)
    title(ws, "05 — CAIXA (conta 45)", "Movimento mensal, contagem física, fundo fixo e reconciliação. Saldo diário acumulado na lista abaixo.", "CAMADA 3 — OUTPUT (contagem física = input)")
    header(ws, 5, 1, ["Rubrica"] + MESES, [34] + [13] * 12)
    labels = ["Saldo inicial", "Entradas", "Saídas", "Saldo final contabilístico", "Contagem física (input)", "Diferença (contab. − física)", "Estado"]
    for i, lab in enumerate(labels):
        put(ws, (6 + i, 1), lab, "label", bold=i in (3, 6))
    for m in range(1, 13):
        c = CL(1 + m)
        if m == 1:
            put(ws, f"{c}6", '=SUMIFS(J_DC,J_G2,"45",J_Ano,CFG_Ano,J_IsAbe,1)', "calc", NUM)
        else:
            put(ws, f"{c}6", f"={CL(m)}9", "calc", NUM)
        put(ws, f"{c}7", f'=SUMIFS(J_Deb,J_G2,"45",J_Ano,CFG_Ano,J_Mes,{m},J_IsAbe,0)', "calc", NUM)
        put(ws, f"{c}8", f'=SUMIFS(J_Cred,J_G2,"45",J_Ano,CFG_Ano,J_Mes,{m},J_IsAbe,0)', "calc", NUM)
        put(ws, f"{c}9", f"={c}6+{c}7-{c}8", "calc", NUM, bold=True)
        put(ws, f"{c}10", D.demo(257_000) if m <= 3 else None, "input", NUM)
        put(ws, f"{c}11", f'=IF({c}10="","",{c}9-{c}10)', "calc", NUM)
        put(ws, f"{c}12", f'=IF({m}>CFG_MesRep,"",IF({c}9<0,"🔴 SALDO DE CAIXA NEGATIVO",IF({c}10="","🟡 Sem contagem física",IF(ABS({c}11)>CFG_Tol,"🔴 DIFERENÇA ENTRE CAIXA CONTABILÍSTICA E CAIXA FÍSICA",IF({c}9>CFG_FundoFixo,"🟡 Acima do fundo fixo","🟢 Reconciliado")))))', "calc")
    status_cf(ws, "B12:M12")
    ws["B10"].comment = None
    put(ws, "A14", "Saldo à data de reporte", "label", bold=True)
    put(ws, "B14", "=INDEX(B9:M9,CFG_MesRep)", "grey", NUM, bold=True)
    name(wb, "CX_Saldo", S, "$B$14")
    put(ws, "A15", "Diferença de contagem à data de reporte", "label")
    put(ws, "B15", '=N(INDEX(B11:M11,CFG_MesRep))', "grey", NUM)
    name(wb, "CX_Dif", S, "$B$15")
    put(ws, "A16", "Estado à data de reporte", "label")
    put(ws, "B16", "=INDEX(B12:M12,CFG_MesRep)", "grey")
    name(wb, "CX_Estado", S, "$B$16")
    status_cf(ws, "B16")
    name(wb, "CX_SF", S, "$B$9:$M$9")
    name(wb, "CX_Est", S, "$B$12:$M$12")
    put(ws, "A17", "Contagem física: valores de teste (fictícios) — substituir pela folha de contagem assinada.", "note")
    section(ws, 19, 1, "MOVIMENTOS DE CAIXA — SALDO DIÁRIO ACUMULADO", 10)
    cols = [("Data", 11, DATE, J("Data")), ("ID", 7, "0", J("ID")), ("Documento", 14, None, '=IF({L}="","",INDEX(J_TipoDoc,{L})&" "&INDEX(J_Serie,{L})&"/"&INDEX(J_NumDoc,{L}))'),
            ("Descrição", 36, None, J("Desc", True)), ("Entrada", 13, NUM, J("Deb")), ("Saída", 13, NUM, J("Cred")), ("Saldo acumulado", 14, NUM, "")]
    last = listing(ws, 22, 300, "SeqCaixa", cols, kcol=1)
    for r in range(22, last + 1):
        prev = f"I{r - 1}" if r > 22 else "0"
        ws[f"I{r}"] = f'=IF($B{r}="","",{prev}+INDEX(J_DC,$B{r}))'
        ws[f"I{r}"].number_format = NUM
    value_cf(ws, f"I22:I{last}", "AND(ISNUMBER(I22),I22<0)", C_ERR)
    protect(ws)


def build_bancos(wb):
    S = "06_BANCOS"
    ws = wb.create_sheet(S)
    title(ws, "06 — BANCOS E MAPA DE RECONCILIAÇÃO BANCÁRIA", "Uma linha por conta bancária. Saldo contabilístico automático; extracto e itens em reconciliação = input.", "CAMADA 3 — OUTPUT (+ input do extracto)")
    heads = ["Conta PGC", "Banco", "IBAN / nº conta", "Moeda", "Saldo contabilístico (Kz)", "Saldo do extracto (Kz)", "Depósitos em trânsito (+)",
             "Cheques em circulação (−)", "Transferências pendentes (−)", "Despesas bancárias não contabilizadas (+)", "Juros creditados não contabilizados (−)",
             "Outros ajustamentos (±)", "Extracto ajustado", "Diferença de reconciliação", "Itens pendentes", "Estado"]
    header(ws, 5, 1, heads, [9, 14, 26, 6, 15, 15, 12, 12, 12, 14, 14, 12, 15, 13, 10, 30])
    contas = D.BANCOS if D.BANCOS else [("43.1.1", "Banco A", "AO06 0000 0000 0000 0000 0000 1 (fictício)", "AOA", D.demo(8_596_500, 0), 0, 0, 0, D.demo(2_500, 0), 0, 0),
              ("43.1.2", "Banco B", "AO06 0000 0000 0000 0000 0000 2 (fictício)", "AOA", 0, 0, 0, 0, 0, 0, 0),
              ("43.2.1", "Banco A — USD", "AO06 0000 0000 0000 0000 0000 3 (fictício)", "USD", 0, 0, 0, 0, 0, 0, 0)]
    for i in range(15):
        r = 6 + i
        v = contas[i] if i < len(contas) else (None,) * 11
        put(ws, (r, 1), v[0], "input", "@")
        put(ws, (r, 2), v[1], "input")
        put(ws, (r, 3), v[2], "input")
        put(ws, (r, 4), v[3], "input")
        put(ws, (r, 5), f'=IF(A{r}="","",SUMIFS(J_DC,J_Conta,A{r},J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef))', "calc", NUM)
        for j in range(7):
            put(ws, (r, 6 + j), v[4 + j], "input", NUM)
        put(ws, (r, 13), f'=IF(A{r}="","",N(F{r})+N(G{r})-N(H{r})-N(I{r})+N(J{r})-N(K{r})+N(L{r}))', "calc", NUM)
        put(ws, (r, 14), f'=IF(A{r}="","",E{r}-M{r})', "calc", NUM)
        put(ws, (r, 15), f'=IF(A{r}="","",(N(G{r})<>0)+(N(H{r})<>0)+(N(I{r})<>0)+(N(J{r})<>0)+(N(K{r})<>0)+(N(L{r})<>0))', "calc", "0")
        put(ws, (r, 16), f'=IF(A{r}="","",IF(E{r}<0,"🔴 SALDO BANCÁRIO NEGATIVO",IF(ABS(N{r})>CFG_Tol,"🔴 Divergência de reconciliação",IF(O{r}>0,"🟡 Reconciliado com itens pendentes","🟢 Reconciliado"))))', "calc")
    status_cf(ws, "P6:P20")
    put(ws, "A22", "TOTAL", "label", bold=True)
    for c in "EFGHIJKLMN":
        put(ws, f"{c}22", f"=SUM({c}6:{c}20)", "grey", NUM, bold=True)
    put(ws, "O22", "=SUM(O6:O20)", "grey", "0")
    put(ws, "P22", '=IF(COUNTIF(P6:P20,"🔴*")>0,"🔴 Bancos com divergências",IF(COUNTIF(P6:P20,"🟡*")>0,"🟡 Reconciliação com itens pendentes","🟢 Bancos reconciliados"))', "grey")
    status_cf(ws, "P22")
    name(wb, "BK_Dif", S, "$N$22")
    name(wb, "BK_Pend", S, "$O$22")
    name(wb, "BK_Estado", S, "$P$22")
    name(wb, "BK_Saldo", S, "$E$22")
    put(ws, "A24", "Controlo: soma das contas 43 no Plano vs contas listadas", "label")
    put(ws, "E24", '=SUMIFS(J_DC,J_G2,"43",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "grey", NUM)
    put(ws, "F24", '=IF(ABS(E24-E22)<0.005,"🟢 Todas as contas bancárias estão no mapa","🔴 Existem contas 43 fora do mapa")', "grey")
    status_cf(ws, "F24")
    put(ws, "A26", "Regra de sinal: Extracto ajustado = extracto + depósitos em trânsito − cheques em circulação − transferências pendentes + despesas não contabilizadas − juros não contabilizados ± outros. Itens do lado da entidade (despesas/juros) devem ser contabilizados no mês seguinte.", "note")
    protect(ws)


def _subledger(wb, S, tipo):
    """tipo = 'CLI' ou 'FORN'."""
    cli = tipo == "CLI"
    g2 = "31" if cli else "32"
    ws = wb.create_sheet(S)
    nome = "CLIENTES E CONTAS A RECEBER" if cli else "FORNECEDORES E CONTAS A PAGAR"
    title(ws, ("07 — " if cli else "08 — ") + nome, "Conta corrente por documento (partidas em aberto). Liquidações identificadas pelo campo Ref_Origem_ID do Diário.", "CAMADA 3 — OUTPUT")
    seq = "SeqCli" if cli else "SeqForn"
    val_side, liq_side = ("J_Deb", "J_Cred") if cli else ("J_Cred", "J_Deb")
    r0 = 26
    buckets = [("Não vencido", -99999, 0), ("0–30", 1, 30), ("31–60", 31, 60), ("61–90", 61, 90), ("91–180", 91, 180), ("181–360", 181, 360), ("> 360", 361, 999999)]
    cols = [("ID", 6, "0", J("ID")), ("Data", 10, DATE, J("Data")), ("Documento", 12, None, '=IF({L}="","",INDEX(J_TipoDoc,{L})&" "&INDEX(J_Serie,{L})&"/"&INDEX(J_NumDoc,{L}))'),
            ("NIF", 12, None, J("NIF", True)), ("Terceiro", 24, None, J("Terceiro", True)),
            ("Vencimento", 10, DATE, '=IF({L}="","",IF(N(INDEX(J_Venc,{L}))=0,INDEX(J_Data,{L}),INDEX(J_Venc,{L})))'),
            ("Valor original", 13, NUM, f'=IF({{L}}="","",INDEX({val_side},{{L}}))'),
            ("Liquidado / regularizado", 13, NUM, f'=IF({{L}}="","",SUMIFS({liq_side},J_Ref,C{{r}},J_G2,"{g2}",J_Data,"<="&CFG_DataRef)-SUMIFS({val_side},J_Ref,C{{r}},J_G2,"{g2}",J_Data,"<="&CFG_DataRef))'),
            ("Saldo em aberto", 13, NUM, '=IF({L}="","",I{r}-J{r})'),
            ("Dias em atraso", 7, "0", '=IF({L}="","",IF(K{r}<=0.005,0,MAX(0,CFG_DataRef-H{r})))'),
            ("Escalão", 10, None, '=IF({L}="","",IF(K{r}<=0.005,"Liquidado",IF(L{r}=0,"Não vencido",IF(L{r}<=30,"0–30",IF(L{r}<=60,"31–60",IF(L{r}<=90,"61–90",IF(L{r}<=180,"91–180",IF(L{r}<=360,"181–360","> 360"))))))))')]
    for b, lo, hi in buckets:
        cols.append((b, 11, NUM, f'=IF({{L}}="","",IF(M{{r}}="{b}",K{{r}},0))'))
    cols.append(("> 90 dias", 11, NUM, '=IF({L}="","",SUM(R{r}:T{r}))'))
    last = listing(ws, r0, 300, seq, cols)
    P = "CLI" if cli else "FRN"
    for nm, c in [("NIF", "F"), ("Saldo", "K"), ("Dias", "L"), ("Esc", "M"), ("DtVenc", "H"), ("V91", "U")]:
        name(wb, f"{P}_{nm}", S, f"${c}${r0}:${c}${last}")
    for j, (b, _, _) in enumerate(buckets):
        name(wb, f"{P}_B{j}", S, f"${CL(14 + j)}${r0}:${CL(14 + j)}${last}")
    # ---------------- resumo / aging ----------------
    section(ws, 4, 1, "AGING E RECONCILIAÇÃO", 12)
    header(ws, 5, 1, ["Escalão"] + [b for b, _, _ in buckets] + ["Total"], None)
    put(ws, "A6", "Saldo em aberto (Kz)", "label")
    for j in range(len(buckets)):
        put(ws, (6, 2 + j), f"=SUM({P}_B{j})", "grey", NUM)
    put(ws, "I6", "=SUM(B6:H6)", "grey", NUM, bold=True)
    put(ws, "A7", "% do total", "label")
    for j in range(len(buckets) + 1):
        put(ws, (7, 2 + j), f"=IF($I$6=0,0,{CL(2 + j)}6/$I$6)", "grey", PCT)
    put(ws, "A9", f"Saldo contabilístico conta {g2} (Plano)", "label")
    sign = "" if cli else "-"
    put(ws, "D9", f'={sign}SUMIFS(J_DC,J_G2,"{g2}",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "grey", NUM)
    put(ws, "A10", "Saldo do auxiliar (partidas em aberto)", "label")
    put(ws, "D10", "=I6", "grey", NUM)
    put(ws, "A11", "Diferença (liquidações sem Ref_Origem ou lançamentos directos)", "label")
    put(ws, "D11", "=D9-D10", "grey", NUM)
    put(ws, "E11", f'=IF(ABS(D11)<=CFG_Tol,"🟢 {"Clientes" if cli else "Fornecedores"} conciliados","🔴 Auxiliar ≠ Razão — verificar Ref_Origem")', "grey")
    status_cf(ws, "E11")
    name(wb, f"{P}_Dif", S, "$D$11")
    name(wb, f"{P}_Total", S, "$I$6")
    name(wb, f"{P}_Estado", S, "$E$11")
    put(ws, "A13", "INDICADORES", "label", bold=True)
    fact = "SUM(T_FactCli)" if cli else "SUM(T_FactForn)"
    ind = [
        ("PMR (dias) — saldo / facturação do período × dias" if cli else "PMP (dias) — saldo / compras do período × dias",
         f"=IF({fact}=0,0,I6/{fact}*CFG_Dias)", "0.0", f"{P}_PM"),
        ("Taxa de cobrança (liquidado / facturado)" if cli else "Taxa de pagamento (liquidado / facturado)",
         f"=IF({fact}=0,0,1-I6/{fact})", PCT, f"{P}_TxLiq"),
        ("Saldo vencido (Kz)" if cli else "Obrigações vencidas (Kz)", "=SUM(C6:H6)", NUM, f"{P}_Venc"),
        ("Saldo vencido > 90 dias", "=SUM(F6:H6)", NUM, f"{P}_V90"),
        ("Compromissos / créditos ainda não vencidos", "=B6", NUM, f"{P}_NaoVenc"),
        ("Concentração — maior " + ("cliente" if cli else "fornecedor") + " (% do saldo)", f"=MAX({'T_QuotaCli' if cli else 'T_QuotaForn'})", PCT, f"{P}_Conc1"),
        ("Concentração — 3 maiores (% do saldo)", "=IFERROR(LARGE({t},1),0)+IFERROR(LARGE({t},2),0)+IFERROR(LARGE({t},3),0)".replace("{t}", "T_QuotaCli" if cli else "T_QuotaForn"), PCT, f"{P}_Conc3"),
    ]
    if cli:
        ind.append(("Clientes em risco (vencido >90d ou acima do limite)", '=COUNTIF(T_V91,">0")+COUNTIF(T_ExcLim,"🔴*")', "0", "CLI_Risco"))
    else:
        ind.append(("Documentos de fornecedores vencidos", '=COUNTIF(FRN_Dias,">0")', "0", "FRN_Risco"))
    for i, (lab, f, fmt, nm) in enumerate(ind):
        r = 14 + i
        put(ws, f"A{r}", lab, "label")
        put(ws, f"D{r}", f, "grey", fmt)
        name(wb, nm, S, f"$D${r}")
    if cli:
        put(ws, "F13", "PROVISÕES / PERDAS ESPERADAS (política interna)", "label", bold=True)
        header(ws, 14, 6, ["Escalão", "% provisão (input)", "Provisão sugerida"], None)
        pcts = [0, 0, 0.1, 0.25, 0.5, 0.75, 1.0]
        for j, (b, _, _) in enumerate(buckets):
            r = 15 + j
            put(ws, f"F{r}", b, "label")
            put(ws, f"G{r}", pcts[j], "input", PCT)
            put(ws, f"H{r}", f"={CL(2 + j)}6*G{r}", "grey", NUM)
        put(ws, "F22", "Total sugerido vs saldo conta 38", "label")
        put(ws, "H22", "=SUM(H15:H21)", "grey", NUM)
        put(ws, "I22", '=-SUMIFS(J_DC,J_G2,"38",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "grey", NUM)
        put(ws, "F23", "Percentagens = política contabilística interna. O tratamento fiscal das provisões (limites dedutíveis) está POR VALIDAR no Código do Imposto Industrial.", "note")
    ws.freeze_panes = f"C{r0}"
    protect(ws)


def build_clientes(wb):
    _subledger(wb, "07_CLIENTES", "CLI")


def build_fornecedores(wb):
    _subledger(wb, "08_FORNECEDORES", "FORN")


def build_inventarios(wb):
    S = "09_INVENTÁRIOS"
    ws = wb.create_sheet(S)
    title(ws, "09 — INVENTÁRIOS (ficha de artigos + kardex a custo médio ponderado)",
          "Movimentos vêm do Diário (colunas Artigo e Quantidade nas linhas da classe 2). Custo médio ponderado móvel implementado; FIFO previsto para a camada Power Query.",
          "CAMADA 1 — INPUT (A:H)  |  CAMADA 2/3 — cálculo")
    heads = ["Código", "Descrição", "Unidade", "Conta", "Stock mínimo", "Stock máximo", "Prazo reposição (dias)", "Stock segurança",
             "Qtd inicial", "Entradas", "Saídas", "Transferências/perdas/devoluções (±)", "Qtd final", "Valor contabilístico", "Custo médio",
             "Consumo médio diário", "Ponto de reposição", "Cobertura (dias)", "Rotação (anualizada)", "Último movimento", "Alerta"]
    header(ws, 5, 1, heads, [8, 30, 6, 7, 9, 9, 9, 9, 10, 10, 10, 12, 10, 14, 11, 10, 10, 9, 9, 11, 26])
    cr = 'J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef'
    for i in range(30):
        r = 6 + i
        v = D.ARTIGOS[i] if i < len(D.ARTIGOS) else (None,) * 8
        for j in range(8):
            put(ws, (r, 1 + j), v[j], "input", "@" if j in (0, 3) else None)
        put(ws, f"I{r}", f'=IF(A{r}="","",SUMIFS(J_Qtd,J_Artigo,A{r},J_IsAbe,1,J_Ano,CFG_Ano))', "grey", "#,##0")
        put(ws, f"J{r}", f'=IF(A{r}="","",SUMIFS(J_Qtd,J_Artigo,A{r},J_IsAbe,0,J_Qtd,">0",J_Nat,"Compra",{cr}))', "grey", "#,##0")
        put(ws, f"K{r}", f'=IF(A{r}="","",-SUMIFS(J_Qtd,J_Artigo,A{r},J_IsAbe,0,J_Qtd,"<0",J_Nat,"Venda",{cr}))', "grey", "#,##0")
        put(ws, f"L{r}", f'=IF(A{r}="","",SUMIFS(J_Qtd,J_Artigo,A{r},J_IsAbe,0,{cr})-J{r}+K{r})', "grey", "#,##0")
        put(ws, f"M{r}", f'=IF(A{r}="","",I{r}+J{r}-K{r}+L{r})', "grey", "#,##0")
        put(ws, f"N{r}", f'=IF(A{r}="","",SUMIFS(J_DC,J_Artigo,A{r},J_Classe,"2",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef))', "grey", NUM)
        put(ws, f"O{r}", f'=IF(OR(A{r}="",N(M{r})=0),"",N{r}/M{r})', "grey", NUM)
        put(ws, f"P{r}", f'=IF(A{r}="","",K{r}/CFG_Dias)', "grey", "#,##0.00")
        put(ws, f"Q{r}", f'=IF(A{r}="","",P{r}*N(G{r})+N(H{r}))', "grey", "#,##0")
        put(ws, f"R{r}", f'=IF(OR(A{r}="",N(P{r})=0),"",M{r}/P{r})', "grey", "#,##0")
        put(ws, f"S{r}", f'=IF(A{r}="","",IF((N(I{r})+N(M{r}))=0,"",K{r}*CFG_Anual/((I{r}+M{r})/2)))', "grey", "0.00")
        put(ws, f"T{r}", f'=IF(A{r}="","",IF(COUNTIF(J_Artigo,A{r})=0,"",_xlfn.MAXIFS(J_Data,J_Artigo,A{r})))', "grey", DATE)
        put(ws, f"U{r}", f'=IF(A{r}="","",IF(M{r}<0,"🔴 Stock negativo",IF(M{r}<=N(E{r}),"🔴 Abaixo do mínimo — repor",IF(M{r}<=Q{r},"🟡 Ponto de reposição atingido",IF(AND(N(F{r})>0,M{r}>F{r}),"🟡 Acima do máximo",IF(AND(ISNUMBER(T{r}),CFG_DataRef-N(T{r})>180,M{r}>0),"🟡 Sem movimento >180d (obsolescência)","🟢 Normal"))))))', "grey")
    status_cf(ws, "U6:U35")
    for nm, c in [("INV_Cod", "A"), ("INV_Qtd", "M"), ("INV_Valor", "N"), ("INV_Alerta", "U")]:
        name(wb, nm, S, f"${c}$6:${c}$35")
    put(ws, "A37", "Σ valor dos artigos", "label")
    put(ws, "N37", "=SUM(N6:N35)", "grey", NUM)
    put(ws, "A38", "Saldo contabilístico classe 2 (Plano)", "label")
    put(ws, "N38", '=SUMIFS(J_DC,J_Classe,"2",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)', "grey", NUM)
    put(ws, "A39", "Diferença (movimentos de classe 2 sem artigo)", "label")
    put(ws, "N39", "=N38-N37", "grey", NUM)
    put(ws, "O39", '=IF(ABS(N39)<=CFG_Tol,"🟢 Inventário conciliado","🔴 Inventário ≠ contabilidade")', "grey")
    status_cf(ws, "O39")
    name(wb, "INV_Dif", S, "$N$39")
    name(wb, "INV_Estado", S, "$O$39")
    # ---------------- Kardex ----------------
    section(ws, 41, 1, "FICHA DE ARMAZÉM (KARDEX) — CUSTO MÉDIO PONDERADO MÓVEL", 14)
    put(ws, "A42", "Artigo:", "label")
    put(ws, "B42", "A001", "input", "@")
    name(wb, "INV_ArtSel", S, "$B$42")
    dv_list(ws, "B42", "=INV_Cod")
    put(ws, "C42", "Método:", "label")
    put(ws, "D42", "Custo médio ponderado", "input")
    dv_list(ws, "D42", '"Custo médio ponderado,FIFO (roadmap Power Query)"')
    put(ws, "F42", "Diferenças de custo detectadas:", "label")
    put(ws, "H42", '=COUNTIF(Q45:Q244,"🔴*")', "grey", "0")
    cols = [("Data", 10, DATE, J("Data")), ("ID", 6, "0", J("ID")), ("Natureza", 18, None, J("Nat", True)), ("Qtd (±)", 9, "#,##0", J("Qtd")),
            ("Valor lançado (D−C)", 13, NUM, J("DC")), ("Qtd em stock", 10, "#,##0", ""), ("Valor em stock", 13, NUM, ""),
            ("CMP após movimento", 11, NUM, ""), ("Custo esperado da saída", 13, NUM, ""), ("Diferença (lançado − esperado)", 13, NUM, ""),
            ("Controlo", 20, None, "")]
    last = listing(ws, 45, 200, "SeqArt", cols)
    for r in range(45, last + 1):
        pq = f"H{r - 1}" if r > 45 else "0"
        pv = f"I{r - 1}" if r > 45 else "0"
        pc = f"J{r - 1}" if r > 45 else "0"
        ws[f"H{r}"] = f'=IF($B{r}="","",{pq}+F{r})'
        ws[f"I{r}"] = f'=IF($B{r}="","",{pv}+G{r})'
        ws[f"J{r}"] = f'=IF($B{r}="","",IF(H{r}=0,0,I{r}/H{r}))'
        ws[f"K{r}"] = f'=IF($B{r}="","",IF(F{r}<0,ROUND(F{r}*N({pc}),2),""))'
        ws[f"L{r}"] = f'=IF(OR($B{r}="",K{r}=""),"",G{r}-K{r})'
        ws[f"M{r}"] = f'=IF($B{r}="","",IF(L{r}="","Entrada",IF(ABS(L{r})<=CFG_Tol,"🟢 Custo correcto","🔴 Custo ≠ CMP")))'
        for c, f in zip("HIJKL", ["#,##0", NUM, NUM, NUM, NUM]):
            ws[f"{c}{r}"].number_format = f
    status_cf(ws, f"M45:M{last}")
    ws["H42"] = f'=COUNTIF(M45:M{last},"🔴*")'
    protect(ws)


def build_activos(wb):
    S = "10_ACTIVOS_FIXOS"
    ws = wb.create_sheet(S)
    title(ws, "10 — ACTIVOS FIXOS E MAPA DE DEPRECIAÇÕES", "Registo de imobilizado (input) → depreciação mensal, acumulada e VLC automáticas. Taxas: política contabilística; taxas fiscais máximas POR VALIDAR (ver 35_BASE_LEGAL).",
          "CAMADA 1 — INPUT (A:Q)  |  CAMADA 2 — cálculo (R em diante)")
    heads = ["Código", "Descrição", "Conta activo", "Conta amort. acum.", "Conta custo", "Data aquisição", "Valor aquisição", "Vida útil (anos)",
             "Método", "Valor residual", "Localização", "Responsável", "Estado", "Data alienação", "Valor alienação", "C. custo", "Projecto",
             "Taxa anual", "Depreciação mensal", "Início depreciação", "Meses até data ref.", "Meses até fim N−1", "Depreciação acumulada",
             "Depreciação do exercício", "VLC", "Mais/menos-valia (alienação)"] + [f"Dep. {m}" for m in MESES]
    header(ws, 5, 1, heads, [7, 28, 8, 8, 7, 11, 13, 7, 14, 10, 16, 16, 9, 11, 12, 6, 7, 7, 12, 11, 7, 7, 14, 14, 14, 12] + [11] * 12)
    ws.freeze_panes = "C6"
    for i in range(100):
        r = 6 + i
        v = D.ACTIVOS[i] if i < len(D.ACTIVOS) else (None,) * 17
        fmts = ["@", None, "@", "@", "@", DATE, NUM, "0", None, NUM, None, None, None, DATE, NUM, "@", "@"]
        for j in range(17):
            put(ws, (r, 1 + j), v[j], "input", fmts[j])
        put(ws, f"R{r}", f'=IF(OR(A{r}="",N(H{r})=0),"",1/H{r})', "grey", PCT)
        put(ws, f"S{r}", f'=IF(OR(A{r}="",N(H{r})=0),"",(G{r}-N(J{r}))/(H{r}*12))', "grey", NUM)
        put(ws, f"T{r}", f'=IF(A{r}="","",IF(CFG_IniDep=1,DATE(YEAR(F{r}),MONTH(F{r}),1),EOMONTH(F{r},0)+1))', "grey", DATE)
        endm = f'IF(N(N{r})>0,MIN({{d}},N{r}),{{d}})'
        mref = f'MAX(0,MIN(H{r}*12,(YEAR({endm.format(d="CFG_DataRef")})-YEAR(T{r}))*12+MONTH({endm.format(d="CFG_DataRef")})-MONTH(T{r})+1))'
        mprev = f'MAX(0,MIN(H{r}*12,(YEAR({endm.format(d="DATE(CFG_Ano-1,12,31)")})-YEAR(T{r}))*12+MONTH({endm.format(d="DATE(CFG_Ano-1,12,31)")})-MONTH(T{r})+1))'
        put(ws, f"U{r}", f'=IF(A{r}="","",{mref})', "grey", "0")
        put(ws, f"V{r}", f'=IF(A{r}="","",{mprev})', "grey", "0")
        put(ws, f"W{r}", f'=IF(A{r}="","",ROUND(U{r}*S{r},2))', "grey", NUM)
        put(ws, f"X{r}", f'=IF(A{r}="","",W{r}-ROUND(V{r}*S{r},2))', "grey", NUM)
        put(ws, f"Y{r}", f'=IF(A{r}="","",IF(N(N{r})>0,0,G{r}-W{r}))', "grey", NUM)
        put(ws, f"Z{r}", f'=IF(OR(A{r}="",N(N{r})=0),"",N(O{r})-(G{r}-W{r}))', "grey", NUM)
        for m in range(1, 13):
            c = CL(26 + m)
            put(ws, f"{c}{r}", f'=IF(A{r}="","",IF(AND(DATE(CFG_Ano,{m},1)>=T{r},DATE(CFG_Ano,{m},1)<EDATE(T{r},H{r}*12),OR(N(N{r})=0,DATE(CFG_Ano,{m},1)<=N{r})),S{r},0))', "grey", NUM)
    # totais e reconciliação
    put(ws, "A107", "TOTAL", "label", bold=True)
    for c in ["G", "W", "X", "Y"] + [CL(26 + m) for m in range(1, 13)]:
        put(ws, f"{c}107", f"=SUM({c}6:{c}105)", "grey", NUM, bold=True)
    put(ws, "A108", "Depreciação contabilizada (conta 73) por mês", "label")
    for m in range(1, 13):
        c = CL(26 + m)
        put(ws, f"{c}108", f'=SUMIFS(J_DC,J_G2,"73",J_Ano,CFG_Ano,J_Mes,{m})', "grey", NUM)
        put(ws, f"{c}109", f'=IF({m}>CFG_MesRep,"",IF(ABS({c}107-{c}108)<=CFG_Tol,"🟢 Processada",IF({c}108=0,"🟡 Por processar","🔴 Diferença")))', "grey")
    put(ws, "A109", "Estado da depreciação mensal", "label")
    status_cf(ws, "AA109:AL109")
    name(wb, "AF_DepEst", S, "$AA$109:$AL$109")
    put(ws, "A111", "RECONCILIAÇÃO REGISTO × CONTABILIDADE", "label", bold=True)
    rec = [("Valor bruto (registo, não alienado) vs contas 11+12+14", '=SUMIFS(G6:G105,N6:N105,"")', '=SUMIFS(J_DC,J_G2,"11",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)+SUMIFS(J_DC,J_G2,"12",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)+SUMIFS(J_DC,J_G2,"14",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)'),
           ("Depreciação acumulada vs conta 18", '=SUMIFS(W6:W105,N6:N105,"")', '=-SUMIFS(J_DC,J_G2,"18",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)'),
           ("Depreciação do exercício vs conta 73", "=X107", '=SUMIFS(J_DC,J_G2,"73",J_Ano,CFG_Ano,J_Data,"<="&CFG_DataRef)')]
    header(ws, 112, 1, ["Controlo", "Registo", "Contabilidade", "Diferença", "Estado"], None)
    for i, (lab, a, b) in enumerate(rec):
        r = 113 + i
        put(ws, f"A{r}", lab, "label")
        put(ws, f"B{r}", a, "grey", NUM)
        put(ws, f"C{r}", b, "grey", NUM)
        put(ws, f"D{r}", f"=C{r}-B{r}", "grey", NUM)
        put(ws, f"E{r}", f'=IF(ABS(D{r})<=CFG_Tol,"🟢 Conciliado","🔴 Diferença")', "grey")
    status_cf(ws, "E113:E115")
    put(ws, "A116", "Estado global activos", "label", bold=True)
    put(ws, "E116", '=IF(COUNTIF(E113:E115,"🔴*")>0,"🔴 Activos não conciliados","🟢 Activos conciliados")', "grey")
    status_cf(ws, "E116")
    name(wb, "AF_Estado", S, "$E$116")
    name(wb, "AF_DepMesReg", S, "$AA$107:$AL$107")
    name(wb, "AF_Capex", S, "$G$107")
    put(ws, "A118", "Nota: a mensagem 'Por processar' indica que a depreciação do mês calculada no registo ainda não foi lançada no Diário (lançamento DI/AMT).", "note")
    protect(ws)
